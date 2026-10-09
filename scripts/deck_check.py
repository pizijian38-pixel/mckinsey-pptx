"""Check a generated deck against its source material.

    python scripts/deck_check.py output/deck.pptx --source inputs/outline.docx [--source more.xlsx ...]

Reports, per deck:
  1. Numbers on slides that never appear in the sources  -> possible fabrication
  2. Share of source numbers that made it into the deck  -> data coverage
  3. Text slides that say little (advisory, source mode)  -> add source detail
  4. Leftover template placeholders ([...], xx, Lorem)    -> unfinished
  5. Claims / quoted terms / vocabulary not in the sources -> possible invention
  6. Same template many slides in a row, no chart, or a parallel group
     (group=) whose slides use different templates          -> layout issues
  7. Body text below 12pt (diagram labels below 10pt)     -> hard to read
  8. Footer "Source:" misuse (caption as source, wrong language)
  9. Insight bullets that only restate a table row
 10. English default labels ("Key insight", "Weighted total" ...) left in a
     Chinese / Korean / Japanese deck
 11. Layout mix (advisory, never fails): text layouts on more than half of the
     content slides -> a reminder to check for relationships a diagram shows better

Sources: .docx .pptx .md .txt .csv .xlsx (.pdf if pypdf is installed).
Exit code 1 if anything is flagged, so a build script can gate on it.
"""
from __future__ import annotations

import argparse
import html
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
try:
    from mckinsey_pptx.labels import english_defaults
    _EN_LABELS = {t.lower() for t in english_defaults()}
except Exception:  # checker still runs without the package
    _EN_LABELS = set()

# A leading sign counts only when it isn't a range dash ("30%-50%", "12-18").
NUM = re.compile(r"(?:(?<![\w.%)])[-+−])?(?<![\w.])\d{1,3}(?:,\d{3})+(?:\.\d+)?"
                 r"|(?:(?<![\w.%)])[-+−])?(?<![\w.])\d+(?:\.\d+)?")
# Outline scaffolding ("Slide 13:", "Table 2") is not data.
SCAFFOLD = re.compile(r"\b(?:slide|table|chart|figure|page|section|第)\s*\d+\b", re.I)
PLACEHOLDER = re.compile(r"\[(?:insert|description|key takeaway|lorem)[^\]]*\]|\blorem ipsum\b|^xx$|\b1\. xx\b|Source: xx",
                         re.I)
CJK = re.compile(r"[⺀-鿿가-힯＀-￯]")


# ---------- text extraction ----------

def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf8")
    xml = re.sub(r"</w:p>|</w:tc>", "\n", xml)
    # &quot; &amp; &lt; ... are XML entities: decode them, or quoted terms and
    # words next to "&" never match the deck text
    return html.unescape(re.sub(r"<[^>]+>", "", xml))


def _pptx_text(path: Path) -> str:
    return "\n".join(t for _, t in _slide_texts(Presentation(str(path))))


def _xlsx_text(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.append(" ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _pdf_text(path: Path) -> str:
    try:
        import pypdf
    except ImportError:
        print(f"  ! skipped {path.name}: pip install pypdf to read PDFs")
        return ""
    return "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages)


# List / heading numbering ("3. Market", "2) Risks", "1.2. Scope", "4、").
# The marker must be followed by a space or the line end, so a table cell that
# holds only "24.4" or "8.0" keeps its number.
_NUMBERING = re.compile(r"(?m)^[ \t]*#{0,6}[ \t]*\d{1,2}(?:\.\d{1,2})*(?:[.)](?=[ \t]|$)|、)[ \t]*")


def source_text(path: Path) -> str:
    text = _NUMBERING.sub(" ", _raw_source_text(path))   # "## 3. Market" headings
    return SCAFFOLD.sub(" ", text)


def _raw_source_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        return _docx_text(path)
    if ext == ".pptx":
        return _pptx_text(path)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx_text(path)
    if ext == ".pdf":
        return _pdf_text(path)
    if ext in (".csv", ".tsv"):
        return _csv_text(path, "\t" if ext == ".tsv" else ",")
    return path.read_text(encoding="utf8", errors="ignore")


def _csv_text(path: Path, delim: str) -> str:
    """One line per row, cells joined by ' | ': a delimiter comma is never read as
    a thousands separator ("APAC,980,1010" is 980 and 1010, not 9801010)."""
    import csv
    with open(path, newline="", encoding="utf-8-sig", errors="ignore") as f:
        return "\n".join(" | ".join(c.strip() for c in row) for row in csv.reader(f, delimiter=delim))


def _iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _frames(slide):
    """(text_frame, width_emu) for every text-bearing shape and table cell."""
    for sh in _iter_shapes(slide.shapes):
        if sh.has_text_frame:
            yield sh.text_frame, sh.width or 0
        if getattr(sh, "has_table", False) and sh.has_table:
            cols = sh.table.columns
            for row in sh.table.rows:
                for ci, cell in enumerate(row.cells):
                    yield cell.text_frame, cols[ci].width


def _chart_text(slide) -> str:
    """Category labels and values of native charts (they hold no text frames)."""
    out = []
    for sh in _iter_shapes(slide.shapes):
        if getattr(sh, "has_chart", False) and sh.has_chart:
            for plot in sh.chart.plots:
                out += [str(c) for c in plot.categories]
                for series in plot.series:
                    if series.name.startswith("_"):     # layout helper (waterfall base)
                        continue
                    out += [str(v) for v in series.values if v not in (None, 0)]
    return "\n".join(out)


_CHROME = re.compile(r"^\s*(?:ⓒ|©|copyright\b).*$", re.I | re.M)


def _content_frames(slide):
    """Text frames minus sequence-number chrome (shapes named 'chrome:*')."""
    for sh in _iter_shapes(slide.shapes):
        if sh.has_text_frame and not (sh.name or "").startswith("chrome:"):
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def derived_numbers(slide) -> set[str]:
    """Numbers the template computed itself (shares, totals, weighted scores):
    shapes named 'derived:*', or values listed in a shape's description
    ('derived:0.40; 1.20'). They need no source; section [2] still counts them."""
    out: set[str] = set()
    for sh in _iter_shapes(slide.shapes):
        if (sh.name or "").startswith("derived:"):
            if sh.has_text_frame:
                out |= numbers(sh.text_frame.text)
        try:
            descr = sh._element.xpath(".//p:cNvPr")[0].get("descr") or ""
        except (IndexError, AttributeError):
            descr = ""
        if descr.startswith("derived:"):
            out |= numbers(descr[len("derived:"):])
    return out


def slide_text(slide) -> str:
    text = "\n".join([tf.text for tf in _content_frames(slide)] + [_chart_text(slide)])
    return _CHROME.sub(" ", text)         # footer "ⓒ 2026 Acme  3" is chrome


def _slide_texts(prs):
    for i, slide in enumerate(prs.slides, start=1):
        yield i, slide_text(slide)


# ---------- checks ----------

def _norm(n: str) -> str:
    # sign-insensitive: charts store decreases as positive heights
    n = n.replace(",", "").replace("−", "-").lstrip("+-")
    if "." in n:
        n = n.rstrip("0").rstrip(".")
    if n.isdigit():
        n = n.lstrip("0") or "0"           # "01" (item numbering) == "1"
    return n


def numbers(text: str) -> set[str]:
    return {_norm(m) for m in NUM.findall(text)}


def fill_ratio(slide, slide_w, slide_h) -> float:
    """Estimated share of the slide covered by text ink (0-1)."""
    ink = 0.0
    for tf, width in _frames(slide):
        for p in tf.paragraphs:
            for r in p.runs:
                size = r.font.size.pt if r.font.size else 12
                chars = sum(1.0 if CJK.match(c) else 0.55 for c in r.text)
                ink += chars * size * size * 1.25  # pt^2
    area = Emu(slide_w).pt * Emu(slide_h).pt
    return ink / area if area else 0.0


# ---------- invention, layout, readability, sources ----------

# Claim types the skill forbids adding, by category (regex, case-insensitive).
# A match is flagged only when the matched phrase is absent from the sources.
RISK_PATTERNS = {
    "superlative / ranking": r"#\s?1\b|\bno\.?\s?1\b|\bnumber[- ]one\b|\bfastest(?:-\w+)?"
                             r"|\blargest\b|\bmarket[- ]lead\w*|\bindustry[- ]lead\w*"
                             r"|\bbest[- ]in[- ]class\b|\bfirst[- ]ever\b|\bunmatched\b"
                             r"|\bunrivall?ed\b|\bworld[- ]class\b",
    "exclusivity / ownership": r"\bproprietary\b|\bpatent\w*|\bexclusive\w*|\btrademark\w*",
    "credential / regulatory": r"\bcertifi\w*|\bclinically[- ]proven\b|\bclinical[- ]trials?\b"
                               r"|\bfda\b|\baccredit\w*|\baward[- ]?\w*|\biso\s?\d{4,5}\b",
    "partnership / deal": r"\bpartnership\w*|\bpartner(?:s|ed)? with\b|\balliance\b"
                          r"|\bjoint[- ]venture\b|\bacqui(?:re|sition)\w*|\bmerger\b"
                          r"|\bofficial\b",
    "commitment / target": r"\bguarantee\w*|\bcommit(?:s|ted)? to\b|\bdominat\w*"
                           r"|\bwill (?:double|triple)\b|\blead(?:s|ing)? (?:the market|in)\b",
    "offer / packaging": r"\bsubscription\w*|\bbundl\w*|\bmulti[- ]?(?:pack|item)\w*"
                         r"|\bfree[- ]trial\b|\bloyalty (?:program|programme|scheme)\b"
                         r"|\bmembership\b|\bpop[- ]up\b|\bflagship store\b",
}
_RISK_RE = {k: re.compile(v, re.I) for k, v in RISK_PATTERNS.items()}
# Double / curly quotes, or single quotes not used as apostrophes (YNBY's).
_QUOTED = re.compile(r"[\"“]([^\"“”\n]{3,60})[\"”]"
                     r"|(?<![A-Za-z])[‘']([^‘’'\n]{3,60})[’'](?![A-Za-z])")
_WORD = re.compile(r"[A-Za-z][A-Za-z'-]{4,}")
_STOP = set("""
insight insights confidential presentation strengths weaknesses opportunities
threats target channel
about above across action active actual added after again against ahead allow
along already also although always among amount another around based basic
became because become becomes before behind being below best better between
beyond brand brands build building built business capture captures change
changes clear close common company complete content continue core create
creates current daily deeper deliver delivers despite develop direct directly
drive driven drives driving during early either enable enables ensure entire
establish every evolve expand expanding expected experience extend extends
fast focus focused follow force forces forward framework fully further gain
gains generate given going great greater growing growth heavy helps higher
highly immediate improve include including increase increasing inside instead
integrate intense issue issues key large largest later launch launches lead
leader leaders leading leads level leverage leveraging local long longer major
make makes making market markets means measure modern moment moments month more
most moving multiple needs never next north offer offers often online order
other others overall paired pairing performance place plan player players point
points position positioning positions power premium pressure primary priority
proof provide provides putting quality quickly range rapid rapidly reach ready
really reduce region remain replace requires result results right rising risk
rivals robust scale secure segment segments serve service shift shifts short
should shows signal since single slide small solid solution specific speed
stand standing start state still strategic strategy stream strong strongly
structural structure support sustained target targeted targets their there
these thing those three through throughout today toward towards track trend
trends under unique unlock until upgrade value values very visible where which
while whole wider winning within without works would years young youth
""".split())


def _stem(w: str) -> str:
    w = w.lower().strip("'-")
    for suf, rep in (("ies", "y"), ("ing", ""), ("ed", ""), ("es", ""),
                     ("s", ""), ("ly", "")):
        if len(w) > 5 and w.endswith(suf):
            return w[: -len(suf)] + rep
    return w


def _norm_text(t: str) -> str:
    return re.sub(r"[^a-z0-9#%\u2e80-\u9fff]+", " ", t.lower())


def new_content(text: str, src_norm: str, src_stems: set[str]):
    """(risk claims, quoted phrases, novel words) on a slide vs. the sources."""
    risk = set()
    for cat, rx in _RISK_RE.items():
        for m in rx.finditer(text):
            phrase = _norm_text(m.group(0)).strip()
            if phrase and phrase not in src_norm:
                risk.add(f"{m.group(0).strip()} ({cat})")
    risk = sorted(risk)
    def stemmed(t):
        return " ".join(_stem(w) for w in _norm_text(t).split())
    src_stemmed = stemmed(src_norm)
    found = [a or b for a, b in _QUOTED.findall(text)]
    quoted = sorted({q.strip() for q in found
                     if stemmed(q) and stemmed(q) not in src_stemmed})
    words = sorted({w.lower() for w in _WORD.findall(text)
                    if w.lower() not in _STOP and _stem(w) not in src_stems})
    return risk, quoted, words


def restated_rows(slide):
    """Text outside a table that just re-says one of its rows."""
    tables = [sh for sh in _iter_shapes(slide.shapes)
              if getattr(sh, "has_table", False) and sh.has_table]
    if not tables:
        return []
    rows = []
    for t in tables:
        for r in list(t.table.rows)[1:]:
            stems = {_stem(w) for c in r.cells for w in _WORD.findall(c.text)}
            stems |= set(NUM.findall(" ".join(c.text for c in r.cells)))
            if stems:
                rows.append(stems)
    hits = []
    # only text level with or below the table (skips title and subtitle)
    top_cut = min(t.top for t in tables) - Emu(45720)
    for sh in _iter_shapes(slide.shapes):
        if not sh.has_text_frame or sh.top is None or sh.top < top_cut:
            continue
        for p in sh.text_frame.paragraphs:
            t = p.text.strip()
            stems = {_stem(w) for w in _WORD.findall(t)} | set(NUM.findall(t))
            if len(stems) < 3:
                continue
            best = max((len(stems & r) / len(stems) for r in rows), default=0)
            if best >= 0.6:
                hits.append(t[:70])
    return hits


_CHART_TEMPLATES = {"chart", "waterfall", "column_comparison", "column_simple_growth",
                    "column_split_growth", "column_historic_forecast",
                    "stacked_column_chart", "grouped_column_chart", "line_chart",
                    "bubble_chart", "bubble_chart_takeaways", "growth_share"}
_FAMILY = {"swot": "card_grid"}
# Templates whose value is the text itself; visual ones (charts, roadmaps,
# matrices, tables, scorecards, summaries) are not judged by text fill.
_TEXT_TEMPLATES = {"card_grid", "card_rows", "swot", "executive_summary",
                   "executive_summary_takeaways", "three_trends_icons",
                   "three_trends_table", "three_trends_numbered", "five_key_areas",
                   "overview_areas", "two_column_compare", "pros_cons", "logic_grid"}


# [11] Layout mix. Text layouts (cards, rows, tables, lists) vs everything else,
# counted over content slides only (aliases included: the record keeps the name
# the build used).
_TEXT_LAYOUTS = (_TEXT_TEMPLATES - {"executive_summary", "executive_summary_takeaways"}) | {
    "cards", "card_rows", "before_after", "logic_chain", "data_table", "table",
    "phases_table_4", "process_activities", "three_trends_table"}
_NON_CONTENT = {"cover_slide", "cover", "section_divider", "agenda", "executive_summary",
                "executive_summary_paragraph", "executive_summary_takeaways",
                "storyline_summary", "dark_navy_summary", "quote_slide", "quote",
                "stat_hero", "big_number", "strategic_challenge", "key_question"}
TEXT_SHARE_REMINDER = 0.5

# [7] Diagrams whose text is short labels on marks (bubbles, nodes, blocks,
# axis ends) are held to a 10pt floor; body text everywhere else to 12pt.
_LABEL_TEMPLATES = {
    "bubble_chart", "bubble_chart_takeaways", "growth_share", "bcg_matrix",
    "prioritization_matrix", "matrix_2x2", "matrix", "issue_tree", "org_chart",
    "marimekko", "mekko", "treemap", "sankey", "slopegraph", "slope", "dumbbell",
    "heatmap", "heat_map", "radar", "spider", "venn", "bump", "rank_chart",
    "cycle", "flywheel", "risk_heatmap", "risk_matrix", "hub_spoke", "swimlane",
    "layer_stack", "layers", "journey", "customer_journey", "funnel",
    "process_flow", "process_flow_horizontal"}
LABEL_MIN_PT, BODY_MIN_PT = 10, 12


def layout_mix(templates):
    """(text-layout slide numbers, content slide count) — slides 1-based."""
    content = [(i, t) for i, t in enumerate(templates, start=1)
               if t and t not in _NON_CONTENT]
    return [i for i, t in content if t in _TEXT_LAYOUTS], len(content)


_PERIOD = re.compile(r"^\s*(?:(?:19|20)\d{2}\s*[A-Za-z]{0,4}|FY\s?\d{2,4}|[QH][1-4](?:\s?\d{2,4})?"
                     r"|\d{4}\s*(?:Est\.?|E|F|A))\s*$", re.I)


def period_table(slide) -> bool:
    """A table whose columns (or rows) are 3+ periods and whose cells are numbers:
    a time series laid out as a table."""
    for sh in _iter_shapes(slide.shapes):
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = list(sh.table.rows)
            head = [c.text for c in rows[0].cells] if rows else []
            first_col = [r.cells[0].text for r in rows]
            if (sum(bool(_PERIOD.match(h)) for h in head) >= 3
                    or sum(bool(_PERIOD.match(h)) for h in first_col) >= 3):
                return numeric_share(slide) >= 0.6
    return False


def _record(slide):
    name = slide._element.cSld.get("name") or ""
    if not name.startswith("mp:"):
        return None, None
    tpl, _, group = name[3:].partition("|")
    return tpl, (group or None)


def template_of(slide):
    return _record(slide)[0]


def group_of(slide):
    return _record(slide)[1]


def numeric_share(slide) -> float:
    """Share of non-header table cells that are numbers (0 if no table)."""
    cells = []
    for sh in _iter_shapes(slide.shapes):
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in list(sh.table.rows)[1:]:
                for c in list(r.cells)[1:]:
                    t = c.text.strip()
                    if t:
                        cells.append(bool(re.fullmatch(r"[\s<>~≈+\-−$€£¥]*[\d.,]+\s*[%xMBK]?\s*", t)))
    return sum(cells) / len(cells) if cells else 0.0


def has_chart(slide) -> bool:
    return any(getattr(sh, "has_chart", False) and sh.has_chart
               for sh in _iter_shapes(slide.shapes)) or \
        template_of(slide) in _CHART_TEMPLATES


def body_size(slide, slide_h) -> float | None:
    """Dominant font size of body text (ignores title, footer, short labels)."""
    from collections import Counter
    sizes = Counter()
    top_cut, bottom_cut = Emu(int(slide_h * 0.14)), Emu(int(slide_h * 0.93))
    for sh in _iter_shapes(slide.shapes):
        if not sh.has_text_frame or sh.top is None:
            continue
        if (sh.name or "").startswith("chrome:"):     # step numbers, captions
            continue
        if sh.top < top_cut or sh.top > bottom_cut:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size and len(r.text) > 15 and r.font.size.pt < 20:
                    sizes[r.font.size.pt] += len(r.text)
    return sizes.most_common(1)[0][0] if sizes else None


_SRC_LABEL = re.compile(r"^(Source:|Sources:|资料来源：|來源：|출처:|出典：)\s*(.*)$")


def english_labels(slide):
    """Known English default labels drawn on the slide (whole text or 'Label: ...')."""
    hits = set()
    for tf, _ in _frames(slide):
        for p in tf.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            low = t.lower().rstrip(":")
            if low in _EN_LABELS:
                hits.add(t.rstrip(":"))
            elif ":" in t and t.split(":", 1)[0].lower() in _EN_LABELS:
                hits.add(t.split(":", 1)[0])
    return hits


def source_issues(slide, deck_is_cjk: bool):
    out = []
    texts = [tf.text.strip() for tf, _ in _frames(slide) if tf.text.strip()]
    for t in texts:
        m = _SRC_LABEL.match(t)
        if not m:
            continue
        label, value = m.group(1), m.group(2).strip()
        if CJK.search(label) and not deck_is_cjk:
            out.append(f"label {label!r} in a non-CJK deck (use make_theme(lang='en'))")
        if value:
            others = [_norm_text(x) for x in texts if x != t]
            v = _norm_text(value).strip()
            if v and any(v in o for o in others):
                out.append(f"source {value!r} repeats a caption/title on the slide — "
                           "a table name is not a source; use source=\"\"")
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("deck")
    ap.add_argument("--source", action="append", default=[],
                    help="source file the deck was built from (repeatable)")
    ap.add_argument("--sparse", type=float, default=0.12,
                    help="flag slides whose text fill ratio is below this")
    ap.add_argument("-v", "--verbose", action="store_true",
                    help="print every slide's text fill ratio")
    args = ap.parse_args(argv)

    prs = Presentation(args.deck)
    sw, sh = prs.slide_width, prs.slide_height
    flagged = False

    src_nums: set[str] = set()
    src_all = ""
    for s in args.source:
        t = source_text(Path(s))
        src_all += "\n" + t
        src_nums |= numbers(t)
    src_norm = _norm_text(src_all)
    src_stems = {_stem(w) for w in _WORD.findall(src_all)}
    all_text = "\n".join(slide_text(s) for s in prs.slides)
    deck_is_cjk = len(CJK.findall(all_text)) > 0.2 * max(len(re.findall(r"\w", all_text)), 1)

    print(f"Deck: {args.deck}  ({len(prs.slides)} slides)")
    deck_nums_all: set[str] = set()
    sparse, leftovers, invented = [], [], []
    risky, quoted_new, novel, small, src_bad, restates = [], [], [], [], [], []
    en_labels = []
    templates, groups = [], []
    for i, slide in enumerate(prs.slides, start=1):
        text = slide_text(slide)
        nums = numbers(_NUMBERING.sub(" ", text))      # "2. Risk" is a list number
        # page number and axis zero are chrome, not content
        nums.discard(str(i))
        nums.discard("0")
        deck_nums_all |= nums
        if args.source:
            derived = derived_numbers(slide)
            new = sorted(n for n in nums if n not in src_nums and n not in derived)
            if new:
                invented.append((i, new))
        ratio = fill_ratio(slide, sw, sh)
        if args.verbose:
            print(f"  slide {i:>2}: fill {ratio:.1%}")
        if (ratio < args.sparse and i > 1 and args.source
                and template_of(slide) in _TEXT_TEMPLATES):
            sparse.append((i, ratio))
        if args.source:
            risk, quoted, words = new_content(text, src_norm, src_stems)
            if risk:
                risky.append((i, risk))
            if quoted:
                quoted_new.append((i, quoted))
            if words:
                novel.append((i, words))
        rr = restated_rows(slide)
        if rr:
            restates.append((i, rr))
        bs = body_size(slide, sh)
        floor = LABEL_MIN_PT if template_of(slide) in _LABEL_TEMPLATES else BODY_MIN_PT
        if bs is not None and bs < floor and i > 1:
            small.append((i, bs, floor))
        if deck_is_cjk and _EN_LABELS:
            hits = sorted(english_labels(slide))
            if hits:
                en_labels.append((i, hits))
        issues = source_issues(slide, deck_is_cjk)
        if issues:
            src_bad.append((i, issues))
        templates.append(template_of(slide))
        groups.append(group_of(slide))
        hits = {m.group(0) for line in text.splitlines()
                for m in [PLACEHOLDER.search(line.strip())] if m}
        if hits:
            leftovers.append((i, sorted(hits)))

    if args.source:
        print("\n[1] Numbers not found in the sources (verify or remove):")
        if invented:
            flagged = True
            for i, ns in invented:
                print(f"  slide {i}: {', '.join(ns)}")
        else:
            print("  none")
        used = src_nums & deck_nums_all
        cov = len(used) / len(src_nums) if src_nums else 1.0
        print(f"\n[2] Source numbers used in the deck: {len(used)}/{len(src_nums)} ({cov:.0%})")
        missing = sorted(src_nums - deck_nums_all, key=lambda x: (len(x), x))
        if missing:
            print(f"  not used: {', '.join(missing[:60])}{' ...' if len(missing) > 60 else ''}")

    print(f"\n[3] Thin text slides (advisory; source mode, text templates, fill < {args.sparse:.0%}):")
    if sparse:
        for i, r in sparse:
            print(f"  slide {i}: {r:.1%} — add supporting detail *from the source* if it has any; "
                  "never pad. Fine for live talks.")
    else:
        print("  none")

    print("\n[4] Leftover placeholders:")
    if leftovers:
        flagged = True
        for i, hits in leftovers:
            print(f"  slide {i}: {hits}")
    else:
        print("  none")

    if args.source:
        print("\n[5] Content not in the sources:")
        print("  a) forbidden claim types (new targets / partners / formats ...) — remove or trace:")
        if risky:
            flagged = True
            for i, r in risky:
                print(f"     slide {i}: {', '.join(r)}")
        else:
            print("     none")
        print("  b) quoted terms not in the sources — remove or trace:")
        if quoted_new:
            flagged = True
            for i, q in quoted_new:
                print(f"     slide {i}: {'; '.join(q)}")
        else:
            print("     none")
        n_novel = sum(len(w) for _, w in novel)
        print(f"  c) vocabulary not in the sources ({n_novel} words) — review each slide;"
              " delete points that add facts, keep plain explanation:")
        for i, ws in novel:
            shown = ", ".join(ws[:14]) + (" ..." if len(ws) > 14 else "")
            print(f"     slide {i} ({len(ws)}): {shown}")

    print("\n[6] Layout variety:")
    known = [t for t in templates if t]
    if not known:
        print("  (deck has no template records — built outside PresentationBuilder)")
    else:
        # a run = consecutive slides of one template family that are NOT one
        # declared parallel group (groups are meant to look alike)
        runs = []
        run_key, run_n, start = None, 0, 0
        for i, (t, g) in enumerate(list(zip(templates, groups)) + [(None, None)], start=1):
            fam = _FAMILY.get(t, t)
            key = None if fam is None else (fam, g)
            if key is not None and key == run_key:
                run_n += 1
            else:
                if run_key and run_n > 3 and run_key[1] is None:
                    runs.append((run_key[0], start, start + run_n - 1))
                run_key, run_n, start = key, 1, i
        for t, a, b in runs:
            flagged = True
            print(f"  slides {a}-{b}: {b - a + 1} x {t} in a row (max 3) — re-express one, "
                  "or mark parallel slides with group=")
        mixed = {}
        for i, (t, g) in enumerate(zip(templates, groups), start=1):
            if g:
                mixed.setdefault(g, []).append((i, t))
        for g, members in mixed.items():
            if len({_FAMILY.get(t, t) for _, t in members}) > 1:
                flagged = True
                desc = ", ".join(f"{i} ({t})" for i, t in members)
                print(f"  group '{g}' mixes templates: slides {desc} — parallel slides "
                      "should share one layout")
        n_charts = sum(has_chart(s) for s in prs.slides)
        numeric_tables = [i for i, s_ in enumerate(prs.slides, start=1)
                          if template_of(s_) in ("data_table", None) and numeric_share(s_) >= 0.5]
        if n_charts == 0 and numeric_tables:
            flagged = True
            print(f"  no chart, but slide(s) {', '.join(map(str, numeric_tables))} show a mostly "
                  "numeric table — chart the series (scores / ratings may stay tables)")
        bad_groups = [g for g, m in mixed.items() if len({_FAMILY.get(t, t) for _, t in m}) > 1]
        if not runs and not bad_groups and not (n_charts == 0 and numeric_tables):
            print(f"  ok ({n_charts} chart slide(s))")

    print("\n[7] Small text (body < 12pt; diagram labels < 10pt):")
    if small:
        flagged = True
        for i, bs, floor in small:
            print(f"  slide {i}: {bs:g}pt (min {floor}pt) — shorten bullets, drop "
                  "subtitle/insight, or split")
    else:
        print("  none")

    print("\n[8] Footer source line:")
    if src_bad:
        flagged = True
        for i, iss in src_bad:
            for x in iss:
                print(f"  slide {i}: {x}")
    else:
        print("  ok")
    print("\n[9] Insight text that only restates a table row (synthesise instead):")
    if restates:
        flagged = True
        for i, rr in restates:
            for t in rr:
                print(f"  slide {i}: {t}")
    else:
        print("  none")

    print("\n[10] English default labels in a CJK deck (build with make_theme(lang=...)"
          " or pass translated labels):")
    if en_labels:
        flagged = True
        for i, hits in en_labels:
            print(f"  slide {i}: {', '.join(hits)}")
    else:
        print("  none")

    # Advisory only: never sets `flagged`. Text layouts are right for lists; the
    # reminder is to look again at slides whose content is a relationship.
    print("\n[11] Layout mix (advisory — a reminder, never a failure):")
    text_slides, n_content = layout_mix(templates)
    if n_content and len(text_slides) / n_content > TEXT_SHARE_REMINDER:
        print(f"  reminder: text layouts on {len(text_slides)} of {n_content} content slides "
              f"({len(text_slides) / n_content:.0%}) — slides {', '.join(map(str, text_slides))}.")
        print("  Check each: if it shows a flow, a share, a change, a ranking, causes or an "
              "overlap, a diagram says it faster (python scripts/catalog.py for the index). "
              "Keep text where the content really is a list.")
    else:
        print(f"  ok ({len(text_slides)} of {n_content} content slides use text layouts)")
    series_tables = [i for i, s_ in enumerate(prs.slides, start=1) if period_table(s_)]
    if series_tables:
        print(f"  reminder: slide(s) {', '.join(map(str, series_tables))} show numbers by period "
              "as a table — a line chart (every series, the focal one highlighted) shows the "
              "trend; keep the table only if exact values per cell are the point.")

    return 1 if flagged else 0


if __name__ == "__main__":
    sys.exit(main())
