"""Print a source file as plain text, tables kept as Markdown tables.

    python scripts/read_source.py <file> [<file> ...]
    python scripts/read_source.py --attachment     # the newest file the user attached
                                                   # in Google Antigravity

Reads .docx .xlsx .pptx .pdf .md .txt .csv (PDF needs pypdf or PyMuPDF).
Use it to read an attached outline instead of writing python -c snippets: every
paragraph in order, every table with all rows and columns, so no number is lost.
"""
from __future__ import annotations

import csv
import html
import re
import sys
import zipfile
from pathlib import Path

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _md_table(rows):
    rows = [[re.sub(r"\s+", " ", str(c or "")).strip().replace("|", "/") for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * n]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def read_docx(path: Path) -> str:
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = root.find(f"{_W}body")
    out, ntab = [], 0

    def para_text(p):
        return "".join(t.text or "" for t in p.iter(f"{_W}t"))

    for el in list(body) if body is not None else []:
        if el.tag == f"{_W}p":
            t = para_text(el).strip()
            if t:
                out.append(t)
        elif el.tag == f"{_W}tbl":
            rows = [[" ".join(para_text(p) for p in tc.iter(f"{_W}p")).strip()
                     for tc in tr.findall(f"{_W}tc")]
                    for tr in el.findall(f"{_W}tr")]
            ntab += 1
            out.append(f"\n[Table {ntab}: {len(rows)} rows]\n" + _md_table(rows) + "\n")
    return "\n".join(out)


def read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook  # optional; read_xlsx_raw is the fallback
    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        out.append(f"\n[Sheet: {ws.title}]\n" + _md_table(rows))
    return "\n".join(out)


def read_xlsx_raw(path: Path) -> str:
    """Stdlib fallback: shared strings + cell values, one table per sheet."""
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        shared = []
        if "xl/sharedStrings.xml" in names:
            xml = z.read("xl/sharedStrings.xml").decode("utf8")
            for si in re.findall(r"<si>(.*?)</si>", xml, re.S):
                shared.append(html.unescape("".join(re.findall(r"<t[^>]*>(.*?)</t>", si, re.S))))
        out = []
        for sheet in sorted(n for n in names if re.match(r"xl/worksheets/sheet\d+\.xml", n)):
            xml = z.read(sheet).decode("utf8")
            rows = []
            for row in re.findall(r"<row[^>]*>(.*?)</row>", xml, re.S):
                cells = []
                for attrs, inner in re.findall(r"<c([^>]*)>(.*?)</c>", row, re.S):
                    v = re.search(r"<v>(.*?)</v>", inner, re.S)
                    t = re.search(r't="(\w+)"', attrs)
                    val = v.group(1) if v else "".join(re.findall(r"<t[^>]*>(.*?)</t>", inner))
                    if t and t.group(1) == "s" and v:
                        val = shared[int(val)]
                    cells.append(html.unescape(val))
                rows.append(cells)
            out.append(f"\n[{Path(sheet).stem}]\n" + _md_table(rows))
        return "\n".join(out)


def read_pptx(path: Path) -> str:
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(str(path)).slides, start=1):
        out.append(f"\n[Slide {i}]")
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                out.append(_md_table([[c.text for c in r.cells] for r in sh.table.rows]))
    return "\n".join(out)


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(str(path)).pages)
    except ImportError:
        import fitz  # PyMuPDF
        return "\n".join(p.get_text() for p in fitz.open(str(path)))


def read_any(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        return read_docx(path)
    if ext in (".xlsx", ".xlsm"):
        try:
            return read_xlsx(path)
        except ImportError:
            return read_xlsx_raw(path)
    if ext == ".pptx":
        return read_pptx(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return _md_table(list(csv.reader(f)))
    return path.read_text(encoding="utf-8", errors="replace")


# Some hosts (Google Antigravity) show only the last ~8,100 characters of a
# command's output. A longer source is saved to a file the agent opens instead.
BUDGET = 7000
UPLOADS = Path.home() / ".gemini" / "antigravity" / "brain"


def newest_attachment():
    """(path, minutes ago) of the newest file in Antigravity's upload folders."""
    files = [p for p in UPLOADS.glob("*/.user_uploaded/*") if p.is_file()]
    if not files:
        return None, None
    p = max(files, key=lambda f: f.stat().st_mtime)
    import time
    return p, (time.time() - p.stat().st_mtime) / 60


def main(argv) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass
    if not argv:
        print(__doc__)
        return 1
    found = None
    if "--attachment" in argv:
        found, age = newest_attachment()
        argv = [a for a in argv if a != "--attachment"]
        if found is None:
            print(f"[read_source] no attachment found under {UPLOADS}. Ask the user for the "
                  "file's path.")
            return 1
        argv = [str(found)] + argv
    status = 0
    texts = []
    for a in argv:
        p = Path(a)
        if not p.exists():
            print(f"[read_source] not found: {p}", file=sys.stderr)
            status = 1
            continue
        try:
            text = read_any(p)
        except ImportError as e:
            print(f"[read_source] {p.name}: missing library ({e.name}); "
                  "install it once with pip, or ask the user for a .docx / .md copy.", file=sys.stderr)
            status = 1
            continue
        texts.append((p, text))
    body = "".join(f"===== {p.name} =====\n{t}\n" for p, t in texts)
    if len(body) <= BUDGET:
        print(body, end="")
    else:
        out = Path("output")
        out.mkdir(exist_ok=True)
        for p, t in texts:
            dest = out / f"source_{p.stem}.md"
            dest.write_text(t, encoding="utf-8")
            n_tables = t.count("\n|---")
            print(f"{p.name}: {len(t):,} characters, {n_tables} table(s) — too long to print "
                  f"(the console shows only the end). Full text: {dest.resolve()} — open it with "
                  "your file viewer; it is exactly what this command would print.")
    for p, _ in texts:
        note = (f" (newest upload, {age:.0f} min ago — check it is the file the user means)"
                if found is not None and p == found else "")
        print(f"--source path: {p.resolve()}{note}")
    return status


if __name__ == "__main__":
    try:
        sys.exit(main(sys.argv[1:]))
    except BrokenPipeError:
        sys.exit(0)
