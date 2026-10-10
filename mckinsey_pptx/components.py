"""Region components for composite slides.

Every component draws itself inside a box (x, y, w, h, inches) and fits its
own text. They are the building blocks of `composite` and of the
option_profiles / risk_register templates.

    flow       value chain / business model: boxes with arrows + caption
    kv_table   label -> value rows (unit economics, key facts)
    metrics    big numbers with small captions
    cards      card grid (same as card_grid)
    pros_cons  ✓ advantages / ✗ disadvantages
    list       numbered items (title + note)
    bullets    bullet list
    text       paragraph(s)
    phases     compact roadmap: phase columns with period band + bullets
    chart      native chart (same arguments as `chart`)
    table      native table (same arguments as `data_table`)
    callout    highlighted statement (value proposition, the option in a line)
    pyramid    positioning / hierarchy pyramid with the chosen level highlighted
    sections   stacked titled groups (Feasibility / Pros / Cons) with ✓ / ✗ marks
"""
from __future__ import annotations

from typing import Dict, Sequence

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from .base import add_line, add_rect, add_textbox, write_paragraph
from .design import (add_icon, add_rich_runs, fit_one_line, fit_size, text_height_in,
                     text_width_pt, tone_rgb, warn_small, write_rich_paragraph)
from .theme import Theme
from .labels import loc

HEADING_H = 0.34
MIN_PT, MAX_PT = 10, 16


def letter_space(paragraph, hundredths_pt: int = 150):
    for r in paragraph.runs:
        r._r.get_or_add_rPr().set("spc", str(hundredths_pt))


def draw_heading(slide, theme: Theme, x, y, w, text, *, color=None, size=11):
    """Small-caps style label above a region ("BUSINESS MODEL")."""
    tb = add_textbox(slide, x, y, w, HEADING_H - 0.04, anchor=MSO_ANCHOR.BOTTOM)
    p = write_paragraph(tb.text_frame, str(loc(theme, text)).upper(), size=size, bold=True,
                        color=color or theme.palette.mid_blue,
                        family=theme.typography.family, first=True)
    letter_space(p, 120)
    return HEADING_H


# ---------- flow ----------

def _chevron(slide, x, y, w, h, rgb):
    s = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x), Inches(y),
                               Inches(w), Inches(h))
    s.rotation = 90
    s.fill.solid()
    s.fill.fore_color.rgb = rgb
    s.line.fill.background()
    s.shadow.inherit = False


def comp_flow(slide, theme, x, y, w, h, *, steps: Sequence[Dict], caption=None,
              highlight=None, where=""):
    """steps: [{title, body?, tone?}] left→right; highlight: index drawn dark."""
    pal, typo = theme.palette, theme.typography
    n = max(len(steps), 1)
    arrow_w = 0.22
    gap = 0.12
    box_w = (w - (n - 1) * (arrow_w + 2 * gap)) / n
    cap_h = 0.0
    if caption:
        cap_h = min(0.6, text_height_in([caption], w, 12) + 0.2)
    box_h = min(1.7, h - cap_h - (0.1 if caption else 0))
    paras = [[s.get("title", "")] + ([s["body"]] if s.get("body") else []) for s in steps]
    size = min((fit_size(p, box_w - 0.3, box_h - 0.3, max_size=14, min_size=MIN_PT,
                         para_gap_pt=4) for p in paras), default=13)
    warn_small("flow", where, size, "Shorten flow step text or use fewer steps.")
    for i, st in enumerate(steps):
        bx = x + i * (box_w + arrow_w + 2 * gap)
        dark = highlight == i
        fill = tone_rgb(theme, st.get("tone")) if dark else pal.soft_gray
        add_rect(slide, bx, y, box_w, box_h, fill=fill)
        tb = add_textbox(slide, bx + 0.15, y + 0.12, box_w - 0.3, box_h - 0.24,
                         anchor=MSO_ANCHOR.MIDDLE)
        p = write_paragraph(tb.text_frame, str(st.get("title", "")).upper(), size=size - 1,
                            bold=True, color=pal.white if dark else pal.text_dark,
                            family=typo.family, align=PP_ALIGN.CENTER, first=True,
                            space_after=4)
        letter_space(p, 80)
        if st.get("body"):
            write_rich_paragraph(tb.text_frame, st["body"], size=size - 1, theme=theme,
                                 color=pal.white if dark else pal.text_dark,
                                 align=PP_ALIGN.CENTER)
        if i < n - 1:
            ax = bx + box_w + gap
            _chevron(slide, ax, y + box_h / 2 - arrow_w / 2, arrow_w, arrow_w,
                     pal.bright_blue)
    if caption:
        cy = y + box_h + 0.12
        add_line(slide, x, cy, x + w, cy, color=pal.rule_gray, width_pt=0.75)
        tb = add_textbox(slide, x, cy + 0.06, w, cap_h)
        write_rich_paragraph(tb.text_frame, caption, size=max(MIN_PT, min(12, size)),
                             theme=theme, color=pal.footer_gray,
                             align=PP_ALIGN.CENTER, first=True)


# ---------- key-value table ----------

def _kv(row):
    """-> (label, value, tone, total). Accepts dicts, [label, value],
    [label, value, tone] and [label, (value, tone)]."""
    if isinstance(row, dict):
        return (row.get("label", ""), row.get("value", ""), row.get("tone"),
                bool(row.get("total")))
    label, value = row[0], row[1]
    tone = row[2] if len(row) > 2 else None
    if isinstance(value, (tuple, list)):
        value, tone = value[0], (value[1] if len(value) > 1 else tone)
    return label, value, tone, False


def comp_kv_table(slide, theme, x, y, w, h, *, rows: Sequence, where=""):
    """rows: [[label, value], ...] or [{label, value, tone?, total?}]; values
    right-aligned bold; total=True draws a rule above and bolds the label."""
    pal = theme.palette
    n = max(len(rows), 1)
    row_h = min(0.55, h / n)
    lab_w = w * 0.6
    size = min((fit_size([str(_kv(r)[0])], lab_w - 0.1, row_h - 0.06, max_size=MAX_PT,
                         min_size=MIN_PT) for r in rows), default=13)
    size = min([size] + [fit_one_line(str(_kv(r)[1]), w - lab_w - 0.1, size, MIN_PT, True)
                         for r in rows])
    warn_small("kv_table", where, size, "Shorten labels or values.")
    for i, r in enumerate(rows):
        label, value, tone, total = _kv(r)
        ry = y + i * row_h
        if total:
            add_line(slide, x, ry, x + w, ry, color=pal.text_dark, width_pt=1.25)
        tb = add_textbox(slide, x, ry, lab_w, row_h, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, str(label), size=size, theme=theme,
                             bold=total, first=True)
        tb = add_textbox(slide, x + lab_w, ry, w - lab_w, row_h, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, str(value), size=size, theme=theme, bold=True,
                             color=tone_rgb(theme, tone) if tone else None,
                             align=PP_ALIGN.RIGHT, first=True)
        add_line(slide, x, ry + row_h, x + w, ry + row_h, color=pal.grid_gray,
                 width_pt=0.5)


# ---------- metrics ----------

def comp_metrics(slide, theme, x, y, w, h, *, items: Sequence[Dict], where=""):
    """items: [{value, label, tone?}] side by side: big number, small caption."""
    pal = theme.palette
    n = max(len(items), 1)
    cw = w / n
    val_h = min(h * 0.6, 0.8)
    vsize = min((fit_one_line(str(it.get("value", "")), cw - 0.2, 32, 14, True)
                 for it in items), default=28)
    vsize = min(vsize, int(val_h * 72 / 1.15))
    lsize = min((fit_size([it.get("label", "")], cw - 0.2, h - val_h, max_size=13,
                          min_size=MIN_PT) for it in items), default=12)
    for i, it in enumerate(items):
        cx = x + i * cw
        tb = add_textbox(slide, cx, y, cw - 0.2, val_h, anchor=MSO_ANCHOR.BOTTOM)
        write_rich_paragraph(tb.text_frame, str(it.get("value", "")), size=vsize,
                             theme=theme, bold=True,
                             color=tone_rgb(theme, it.get("tone")) if it.get("tone")
                             else pal.text_dark, first=True)
        tb = add_textbox(slide, cx, y + val_h + 0.04, cw - 0.2, h - val_h - 0.04)
        write_rich_paragraph(tb.text_frame, it.get("label", ""), size=lsize, theme=theme,
                             color=pal.footer_gray, first=True)


# ---------- pros / cons ----------

def comp_pros_cons(slide, theme, x, y, w, h, *, pros: Sequence[str] = (),
                   cons: Sequence[str] = (), labels=("Advantages", "Disadvantages"),
                   layout="stacked", size=None, where=""):
    """✓ / ✗ lists; layout 'stacked' (one above the other) or 'columns'."""
    groups = [(loc(theme, labels[0]), list(pros), "check", "green"),
              (loc(theme, labels[1]), list(cons), "cross", "red")]
    groups = [g for g in groups if g[1]]
    if not groups:
        return
    d, hh, item_gap, group_gap = 0.22, 0.28, 0.08, 0.12
    tw = (w / len(groups) - 0.2 if layout == "columns" else w) - d - 0.12

    def group_h(items, sz):
        return hh + 0.06 + sum(max(text_height_in([it], tw, sz), d) + item_gap for it in items)

    if layout == "columns":
        boxes = [(x + i * (w / len(groups)), y, w / len(groups) - 0.2, h)
                 for i in range(len(groups))]
        size = size or next((sz for sz in range(MAX_PT, MIN_PT - 1, -1)
                             if all(group_h(g[1], sz) <= h for g in groups)), MIN_PT)
    else:
        size = size or next((sz for sz in range(MAX_PT, MIN_PT - 1, -1)
                             if sum(group_h(g[1], sz) for g in groups)
                             + group_gap * (len(groups) - 1) <= h), MIN_PT)
        boxes, cy = [], y
        for g in groups:
            gh = group_h(g[1], size)
            boxes.append((x, cy, w, gh))
            cy += gh + group_gap
    warn_small("pros_cons", where, size, "Shorten advantages / disadvantages.")
    for (label, items, icon, tone), (bx, by, bw, bh) in zip(groups, boxes):
        tb = add_textbox(slide, bx, by, bw, hh, anchor=MSO_ANCHOR.BOTTOM)
        p = write_paragraph(tb.text_frame, str(label).upper(), size=10, bold=True,
                            color=tone_rgb(theme, tone), family=theme.typography.family,
                            first=True)
        letter_space(p, 120)
        cy = by + hh + 0.06
        for it in items:
            ih = max(text_height_in([it], tw, size), d)
            add_icon(slide, icon, bx, cy + max(0, (size / 72 * 1.2 - d) / 2), d, theme, tone)
            tb = add_textbox(slide, bx + d + 0.12, cy, tw, ih)
            write_rich_paragraph(tb.text_frame, it, size=size, theme=theme, first=True)
            cy += ih + item_gap


def pros_cons_height(pros, cons, w, size, layout="stacked"):
    """Height comp_pros_cons needs at a given size — for planners."""
    d, hh, item_gap, group_gap = 0.22, 0.28, 0.08, 0.12
    groups = [g for g in (pros, cons) if g]
    if not groups:
        return 0
    tw = (w / len(groups) - 0.2 if layout == "columns" else w) - d - 0.12
    hs = [hh + 0.06 + sum(max(text_height_in([it], tw, size), d) + item_gap for it in g)
          for g in groups]
    if layout == "columns":
        return max(hs)
    return sum(hs) + group_gap * (len(groups) - 1)


# ---------- numbered list / bullets / text ----------

def comp_list(slide, theme, x, y, w, h, *, items: Sequence, where=""):
    """items: [{title, body?, tone?}] or [str]; numbered circles."""
    pal = theme.palette
    items = [it if isinstance(it, dict) else {"title": it} for it in items]
    n = max(len(items), 1)
    row_h = min(0.75, h / n)
    d = min(0.36, row_h - 0.1)
    texts = [it["title"] + (" — " + it["body"] if it.get("body") else "") for it in items]
    size = min((fit_size([t], w - d - 0.2, row_h - 0.08, max_size=MAX_PT, min_size=MIN_PT)
                for t in texts), default=13)
    warn_small("list", where, size, "Shorten list items.")
    for i, it in enumerate(items):
        ry = y + i * row_h
        add_icon(slide, str(i + 1), x, ry + (row_h - d) / 2, d, theme, it.get("tone", "gray"))
        tb = add_textbox(slide, x + d + 0.15, ry, w - d - 0.15, row_h, anchor=MSO_ANCHOR.MIDDLE)
        p = write_rich_paragraph(tb.text_frame, f"**{it['title']}**", size=size, theme=theme,
                                 first=True)
        if it.get("body"):
            add_rich_runs(p, "  " + it["body"], size=size, color=pal.footer_gray,
                          family=theme.typography.family, theme=theme)


def comp_bullets(slide, theme, x, y, w, h, *, items: Sequence[str], where=""):
    size = fit_size(list(items) or [" "], w, h, max_size=MAX_PT, min_size=MIN_PT,
                    para_gap_pt=5, indent_in=0.25)
    warn_small("bullets", where, size, "Shorten bullets.")
    tb = add_textbox(slide, x, y, w, h)
    for i, it in enumerate(items):
        write_rich_paragraph(tb.text_frame, it, size=size, theme=theme, bullet=True,
                             first=i == 0, space_before=0 if i == 0 else 5)


def comp_text(slide, theme, x, y, w, h, *, text, size=None, color=None, where=""):
    paras = [text] if isinstance(text, str) else list(text)
    size = size or fit_size(paras, w, h, max_size=MAX_PT, min_size=MIN_PT, para_gap_pt=6)
    warn_small("text", where, size, "Shorten the text.")
    tb = add_textbox(slide, x, y, w, h)
    for i, t in enumerate(paras):
        write_rich_paragraph(tb.text_frame, t, size=size, theme=theme, first=i == 0,
                             color=color, space_before=0 if i == 0 else 6)


# ---------- phases ----------

def comp_phases(slide, theme, x, y, w, h, *, phases: Sequence[Dict], where=""):
    """phases: [{name, period?, goal?, bullets?: [str], tone?}] as columns with
    a header band (period) — a compact roadmap inside a region."""
    pal, typo = theme.palette, theme.typography
    n = max(len(phases), 1)
    gap = 0.15
    cw = (w - gap * (n - 1)) / n
    band_h = 0.38
    head_h = 0.62
    paras = [list(p.get("bullets", [])) or [" "] for p in phases]
    size = min((fit_size(pp, cw - 0.4, h - band_h - head_h - 0.25, max_size=MAX_PT,
                         min_size=MIN_PT, para_gap_pt=4, indent_in=0.25) for pp in paras),
               default=13)
    warn_small("phases", where, size, "Shorten phase bullets.")
    for i, ph in enumerate(phases):
        px = x + i * (cw + gap)
        tone = ph.get("tone") or ("navy", "mid_blue", "blue", "light_blue")[min(i, 3)]
        add_rect(slide, px, y, cw, band_h, fill=tone_rgb(theme, tone))
        tb = add_textbox(slide, px, y, cw, band_h, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, ph.get("period", ph.get("name", "")), size=12, bold=True,
                        color=pal.text_dark if tone == "light_blue" else pal.white,
                        family=typo.family, align=PP_ALIGN.CENTER, first=True)
        add_rect(slide, px, y + band_h, cw, h - band_h, fill=pal.soft_gray)
        tb = add_textbox(slide, px + 0.2, y + band_h + 0.12, cw - 0.4, head_h - 0.12)
        p = write_paragraph(tb.text_frame, str(ph.get("name", "")).upper(), size=12, bold=True,
                            color=pal.text_dark, family=typo.family, first=True)
        letter_space(p, 100)
        if ph.get("goal"):
            write_rich_paragraph(tb.text_frame, ph["goal"], size=11, theme=theme,
                                 color=pal.footer_gray)
        tb = add_textbox(slide, px + 0.2, y + band_h + head_h + 0.05, cw - 0.4,
                         h - band_h - head_h - 0.15)
        for k, bl in enumerate(ph.get("bullets", [])):
            write_rich_paragraph(tb.text_frame, bl, size=size, theme=theme, bullet=True,
                                 first=k == 0, space_before=0 if k == 0 else 4)


# ---------- callout ----------

def comp_callout(slide, theme, x, y, w, h, *, text, label=None, style="outline", tone=None,
                 where=""):
    """Highlighted statement: a new value proposition, the option in one line,
    a 'so what'. style: outline (label sits on the border) | solid | tint."""
    from .design import tint
    pal, typo = theme.palette, theme.typography
    label = loc(theme, label)
    color = tone_rgb(theme, tone) if tone else pal.mid_blue
    lab_h = 0.3 if label else 0.0
    inner_w = w - 0.4
    size = fit_size([text], inner_w, max(h - lab_h - 0.3, 0.3), max_size=16, min_size=MIN_PT,
                    bold=style == "solid")
    warn_small("callout", where, size, "Shorten the callout text.")
    need = text_height_in([text], inner_w, size, bold=style == "solid") + 0.36
    bh = min(h - lab_h / 2, need + (lab_h / 2 if style == "outline" else lab_h))
    by = y + (lab_h / 2 if style == "outline" else 0)
    if style == "solid":
        add_rect(slide, x, by, w, bh, fill=color)
        txt_color = pal.white
    elif style == "tint":
        add_rect(slide, x, by, w, bh, fill=tint(color, 0.85))
        txt_color = pal.deep_navy
    else:
        add_rect(slide, x, by, w, bh, fill=pal.white, line=color, line_width=1.5)
        txt_color = pal.text_dark
    ty = by
    if label:
        lsize = fit_one_line(label, w - 0.6, 12, 9, True)
        lw = min(w - 0.4, text_width_pt(label, lsize, True) / 72 / 0.92 + 0.3)
        if style == "outline":
            tb = add_textbox(slide, x + (w - lw) / 2, y, lw, lab_h, fill=pal.white,
                             anchor=MSO_ANCHOR.MIDDLE)
            ty = y + lab_h
        else:
            tb = add_textbox(slide, x + 0.2, by + 0.08, w - 0.4, lab_h, anchor=MSO_ANCHOR.MIDDLE)
            ty = by + lab_h
        write_paragraph(tb.text_frame, label, size=lsize, bold=True,
                        color=pal.white if style == "solid" else color,
                        family=typo.family,
                        align=PP_ALIGN.CENTER if style == "outline" else PP_ALIGN.LEFT,
                        first=True)
    tb = add_textbox(slide, x + 0.2, ty, inner_w, by + bh - ty - 0.06,
                     anchor=MSO_ANCHOR.MIDDLE)
    write_rich_paragraph(tb.text_frame, text, size=size, theme=theme, color=txt_color,
                         bold=style == "solid",
                         align=PP_ALIGN.CENTER if style == "outline" else PP_ALIGN.LEFT,
                         first=True)


# ---------- pyramid ----------

def _polygon(slide, pts, fill):
    from pptx.util import Emu
    emu = [(int(Inches(px)), int(Inches(py))) for px, py in pts]
    ff = slide.shapes.build_freeform(emu[0][0], emu[0][1], scale=1.0)
    ff.add_line_segments(emu[1:], close=True)
    shp = ff.convert_to_shape()
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def comp_pyramid(slide, theme, x, y, w, h, *, levels: Sequence[str], highlight=None,
                 notes: Sequence[str] = (), note=None, tone=None, where=""):
    """Positioning / hierarchy pyramid, top level first. `highlight`: index or
    list of indices (the chosen position); `notes`: one note per level, or a
    single `note` placed beside the highlighted level."""
    pal, typo = theme.palette, theme.typography
    n = max(len(levels), 1)
    hl = set([highlight] if isinstance(highlight, int) else (highlight or []))
    has_notes = bool(notes) or bool(note)
    pw = min(w * (0.55 if has_notes else 1.0), h * 1.35)
    ph = min(h, pw * 0.85)
    px = x
    py = y + (h - ph) / 2
    gap = 0.05
    band = (ph - gap * (n - 1)) / n
    cx = px + pw / 2

    def half(yy):  # half-width of the triangle at height yy
        return (yy - py) / ph * pw / 2

    # each label must fit on one line at the narrowest point it sits on: the top
    # level's label sits at the base of the apex triangle
    def text_w(i):
        y1 = py + i * (band + gap)
        return 2 * half(y1 + 0.72 * band if i == 0 else y1) * 0.85
    lsize = min(fit_one_line(lv, text_w(i), 15, 8, True)
                for i, lv in enumerate(levels)) if levels else 12
    lsize = min(lsize, int(band * 72 / 1.6))
    centers = []
    for i, lv in enumerate(levels):
        y1 = py + i * (band + gap)
        y2 = y1 + band
        fill = tone_rgb(theme, tone or "mid_blue") if i in hl else pal.footer_gray
        if i == 0:
            pts = [(cx, y1), (cx + half(y2), y2), (cx - half(y2), y2)]
        else:
            pts = [(cx - half(y1), y1), (cx + half(y1), y1), (cx + half(y2), y2),
                   (cx - half(y2), y2)]
        _polygon(slide, pts, fill)
        ty = y1 + band * (0.4 if i == 0 else 0.0)
        tb = add_textbox(slide, cx - half(y2), ty, 2 * half(y2), y2 - ty - (0.04 if i == 0 else 0),
                         anchor=MSO_ANCHOR.MIDDLE if i else MSO_ANCHOR.BOTTOM)
        tb.text_frame.word_wrap = False
        write_paragraph(tb.text_frame, lv, size=lsize, bold=True, color=pal.white,
                        family=typo.family, align=PP_ALIGN.CENTER, first=True)
        centers.append((y1 + y2) / 2 + (band * 0.15 if i == 0 else 0))
    nx = px + pw + 0.25
    nw = x + w - nx
    if notes and nw > 0.8:
        nsize = min(fit_size([t], nw, band, max_size=14, min_size=MIN_PT) for t in notes)
        for i, t in enumerate(notes[:n]):
            if not t:
                continue
            add_line(slide, cx + half(centers[i]) + 0.05, centers[i], nx - 0.05, centers[i],
                     color=pal.rule_gray, width_pt=0.75)
            tb = add_textbox(slide, nx, centers[i] - band / 2, nw, band, anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, t, size=nsize, theme=theme, first=True)
    elif note and nw > 0.8:
        i = min(hl) if hl else n // 2
        nsize = fit_size([note], nw, h * 0.6, max_size=15, min_size=MIN_PT)
        nh = text_height_in([note], nw, nsize) + 0.1
        ny = min(max(centers[i] - nh / 2, y), y + h - nh)
        add_line(slide, cx + half(centers[i]) + 0.05, centers[i], nx - 0.05, centers[i],
                 color=pal.rule_gray, width_pt=0.75)
        tb = add_textbox(slide, nx, ny, nw, nh, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, note, size=nsize, theme=theme, first=True)


# ---------- sections ----------

_MARKS = {"check": ("✓", "green"), "cross": ("✗", "red"), "dot": ("•", None)}


def _set_mark(paragraph, char, rgb):
    from lxml import etree
    from pptx.oxml.ns import qn
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if rgb is not None:
        clr = etree.SubElement(pPr, qn("a:buClr"))
        srgb = etree.SubElement(clr, qn("a:srgbClr"))
        srgb.set("val", str(rgb))
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)
    pPr.set("indent", "-228600")
    pPr.set("marL", "228600")


def comp_sections(slide, theme, x, y, w, h, *, sections: Sequence[Dict], boxed=True,
                  where=""):
    """Stacked titled groups — e.g. Feasibility / Pros / Cons of one option.
    sections: [{"title", "bullets"? | "body"?, "tone"?, "mark"?: check|cross|dot}].
    Pros-like sections default to ✓ (tone green), cons-like to ✗ (red).
    boxed=False drops the outline boxes (use inside a `panel`)."""
    from .design import add_stage_banners
    pal = theme.palette
    secs = [dict(s) for s in sections]
    head_h, gap, pad = 0.32, 0.12, 0.12
    inner = w - 2 * pad
    m = max(len(secs), 1)

    def items(s):
        return ([s["body"]] if s.get("body") else []) + list(s.get("bullets", []))

    def need(s, sz):
        return head_h + 0.06 + text_height_in(items(s) or [" "], inner, sz, para_gap_pt=4,
                                              indent_in=0.25) + 2 * pad
    size = MIN_PT
    for sz in range(MAX_PT - 1, MIN_PT - 1, -1):
        if sum(need(s, sz) for s in secs) + gap * (m - 1) <= h:
            size = sz
            break
    warn_small("sections", where, size, "Shorten the section bullets.")
    needs = [need(s, size) for s in secs]
    extra = h - sum(needs) - gap * (m - 1)
    add = min(extra / m, 0.4) if extra > 0 else 0.0
    if extra < 0:  # even the minimum size does not fit: stay inside the region
        k = (h - gap * (m - 1)) / sum(needs)
        needs = [max(nh * k, head_h + 0.3) for nh in needs]
    cy = y
    for s, nh in zip(secs, needs):
        sh = nh + add
        tone = s.get("tone")
        color = tone_rgb(theme, tone) if tone else pal.deep_navy
        add_stage_banners(slide, theme, [(x, w, s.get("title", ""), "result", color)], cy,
                          head_h, style="rule", size=13)
        by = cy + head_h + 0.06
        if boxed:
            add_rect(slide, x, by, w, sh - head_h - 0.06, fill=pal.white, line=color,
                     line_width=1.0)
        tb = add_textbox(slide, x + pad, by + pad, inner, sh - head_h - 0.06 - 2 * pad,
                         anchor=MSO_ANCHOR.MIDDLE)
        mark = s.get("mark") or {"green": "check", "red": "cross"}.get(tone)
        ch, mtone = _MARKS.get(mark, (None, None))
        for k, t in enumerate(items(s)):
            p = write_rich_paragraph(tb.text_frame, t, size=size, theme=theme,
                                     bullet=ch is None, first=k == 0,
                                     space_before=0 if k == 0 else 4)
            if ch:
                _set_mark(p, ch, tone_rgb(theme, mtone) if mtone else color)
        cy += sh + gap


# ---------- wrappers around full templates' drawers ----------

def comp_cards(slide, theme, x, y, w, h, *, cards, columns=None, where=""):
    from .slides.card_slides import draw_cards
    draw_cards(slide, theme, x, y, w, h, cards, columns=columns, where=where, max_size=MAX_PT)


def comp_chart(slide, theme, x, y, w, h, *, where="", **kw):
    from .slides.native_chart import draw_chart
    draw_chart(slide, theme, x, y, w, h, **kw)


def comp_table(slide, theme, x, y, w, h, *, where="", **kw):
    from .slides.table_slides import draw_table
    kw.setdefault("font_size", 12)
    draw_table(slide, theme, x, y, w, h, **kw)


# Diagram templates that draw cleanly into a region: template -> smallest region
# (width, height in inches). Measured with the CATALOG example of each template:
# every shape inside the region, text >= 9 pt, no word broken mid-word
# (tests/test_regressions.py re-checks this). Denser content needs more room.
# Not listed: cycle, swimlane, hub_spoke, risk_heatmap and venn draw labels outside
# their plot or overflow a short region - use them as full slides.
DIAGRAM_MIN = {
    "bump": (6.0, 3.8), "dumbbell": (4.4, 3.4), "fishbone": (8.0, 4.5),
    "heatmap": (4.4, 3.4), "journey": (7.0, 4.4), "layer_stack": (5.2, 3.8),
    "marimekko": (6.0, 4.0), "positioning_scale": (6.0, 4.0), "radar": (6.0, 3.8),
    "sankey": (4.4, 3.4), "slopegraph": (4.4, 3.4), "timeline": (6.0, 3.4),
    "treemap": (6.0, 3.8), "value_chain": (7.0, 4.4),
}


def diagram_templates(columns) -> list:
    """Names of the diagram templates used in a composite's columns."""
    found = []
    for col in columns or []:
        for region in (col if isinstance(col, (list, tuple)) else [col]):
            if isinstance(region, dict) and region.get("type") == "diagram":
                found.append(str(region.get("template")))
    return found


def comp_diagram(slide, theme, x, y, w, h, *, template, where="", **kw):
    """Draw a diagram template (`template`: its name) into the region. The
    template's own arguments go alongside; title / subtitle / insight / source
    belong to the composite, not here."""
    from . import builder, validate
    from .design import draw_into, warn_small
    fn = builder._REGISTRY.get(template)
    if fn is None:
        raise ValueError(validate.unknown_template(template, sorted(DIAGRAM_MIN)))
    name = builder.template_name(fn)
    if name not in DIAGRAM_MIN:
        raise ValueError(f"diagram: {template!r} cannot be drawn into a region. "
                         f"Use one of: {', '.join(sorted(DIAGRAM_MIN))}")
    chrome = [k for k in ("title", "subtitle", "insight", "insight_label", "source",
                          "footnote", "page_number", "section_marker") if k in kw]
    if chrome:
        raise ValueError(f"diagram {name}: {', '.join(chrome)} belong on the composite, "
                         "not on the diagram region")
    errs = [m for lv, m in validate.signature_problems(fn, kw) if lv == "error"]
    if errs:
        raise ValueError(f"diagram {name}: " + "; ".join(errs))
    min_w, min_h = DIAGRAM_MIN[name]
    if w < min_w or h < min_h:
        warn_small("diagram", where or name, 0,
                   f"{name} needs a region of at least {min_w:.1f} x {min_h:.1f} in "
                   f"(got {w:.1f} x {h:.1f}); give it a wider column or more height.")
    prs = slide.part.package.presentation_part.presentation
    with draw_into(slide, x, y, w, h):
        fn(prs, title=where or name, theme=theme, **kw)


COMPONENTS = {
    "flow": comp_flow, "kv_table": comp_kv_table, "metrics": comp_metrics,
    "cards": comp_cards, "pros_cons": comp_pros_cons, "list": comp_list,
    "bullets": comp_bullets, "text": comp_text, "chart": comp_chart,
    "table": comp_table, "phases": comp_phases,
    "callout": comp_callout, "pyramid": comp_pyramid, "sections": comp_sections,
    "diagram": comp_diagram,
}


def draw_region(slide, theme: Theme, x, y, w, h, region: Dict, *, where=""):
    """Draw one composite region: optional heading, optional panel, component."""
    region = dict(region)
    kind = region.pop("type")
    heading = region.pop("heading", None)
    panel = region.pop("panel", False)
    region.pop("weight", None)
    fn = COMPONENTS.get(kind)
    if fn is None:
        raise ValueError(f"Unknown region type {kind!r}. Available: {sorted(COMPONENTS)}")
    region.pop("arrow", None)  # drawn by the composite (arrow from the region above)
    if panel:
        from .design import tint
        pal = theme.palette
        if panel == "outline":
            add_rect(slide, x, y, w, h, fill=pal.white, line=pal.mid_blue, line_width=1.0)
        elif panel == "tint":
            add_rect(slide, x, y, w, h, fill=tint(pal.light_blue, 0.85))
        else:
            add_rect(slide, x, y, w, h, fill=pal.soft_gray)
        x, y, w, h = x + 0.2, y + 0.12, w - 0.4, h - 0.24
    if heading:
        y += draw_heading(slide, theme, x, y, w, heading)
        h -= HEADING_H
        y += 0.06
        h -= 0.06
    fn(slide, theme, x, y, w, h, where=where, **region)
