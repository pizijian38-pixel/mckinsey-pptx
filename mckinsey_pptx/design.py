"""Shared design primitives for the richer templates.

- Tones: semantic colours ("red" = problem/risk, "green" = opportunity ...).
- Rich text: `**bold**` and `{tone|coloured bold}` spans inside any string.
- Icons: a bundled Lucide icon set drawn white on a tone-coloured circle.
- Fitting: pick the largest font size whose estimated height fits a box.
- Panels: the "Key insight" side panel and the highlighted callout bar.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Iterable, Optional

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches

from .base import add_oval, add_rect, add_textbox, set_run, write_paragraph
from .theme import Theme
from .labels import loc
from .metrics import (TONE_NAMES, fit_one_line, fit_size, plain,  # noqa: F401
                      text_height_in, text_width_pt)

ICON_DIR = Path(__file__).resolve().parent / "assets" / "icons"
ICONS = sorted(p.stem for p in ICON_DIR.glob("*.png"))



def tone_rgb(theme: Theme, tone: Optional[str]):
    pal = theme.palette
    return {
        None: pal.deep_navy, "navy": pal.deep_navy, "blue": pal.bright_blue,
        "mid_blue": pal.mid_blue, "light_blue": pal.light_blue,
        "red": pal.status_red, "green": pal.status_green,
        "amber": pal.status_amber, "gray": pal.footer_gray,
        # neutral emphasis (best score, highlight without good/bad meaning)
        "gold": RGBColor(0xA6, 0x7C, 0x1B),
    }.get(tone, pal.deep_navy)


# ---------- rich text ----------

_SPAN = re.compile(r"(\*\*.+?\*\*|\{(?:%s)\|.+?\})" % "|".join(TONE_NAMES))


def add_rich_runs(paragraph, text, *, size, color, family, theme: Theme,
                  bold=False):
    """Append runs to `paragraph`, honouring **bold** and {tone|text} spans."""
    for part in _SPAN.split(str(text)):
        if not part:
            continue
        run_color, run_bold, s = color, bold, part
        if part.startswith("**") and part.endswith("**"):
            s, run_bold = part[2:-2], True
        elif part.startswith("{") and part.endswith("}") and "|" in part:
            tone, s = part[1:-1].split("|", 1)
            run_color, run_bold = tone_rgb(theme, tone), True
        set_run(paragraph.add_run(), s, size=size, bold=run_bold,
                color=run_color, family=family)


def write_rich_paragraph(tf, text, *, size, theme: Theme, color=None,
                         bold=False, align=PP_ALIGN.LEFT, bullet=False,
                         first=False, space_before=None, space_after=None):
    color = color or theme.palette.text_dark
    p = write_paragraph(tf, "", size=size, bold=bold, color=color,
                        family=theme.typography.family, align=align,
                        bullet=bullet, first=first, space_before=space_before,
                        space_after=space_after)
    r = p.runs[0]._r
    r.getparent().remove(r)
    add_rich_runs(p, text, size=size, color=color,
                  family=theme.typography.family, theme=theme, bold=bold)
    return p


# ---------- fitting ----------

# ---------- warnings ----------

MIN_READABLE_PT = 12


def warn_small(where: str, title: str, size: int, hint: str):
    """Tell the build log when fitting had to go below readable size."""
    if size == 0:  # layout warning, not a size
        print(f"[mckinsey_pptx] WARNING {where} \"{str(title)[:50]}\": {hint}",
              file=sys.stderr)
    elif size < MIN_READABLE_PT:
        print(f"[mckinsey_pptx] WARNING {where} \"{str(title)[:50]}\": text fitted "
              f"at {size}pt (< {MIN_READABLE_PT}pt). {hint}", file=sys.stderr)


# ---------- icons ----------

def add_icon(slide, name: Optional[str], x_in, y_in, d_in, theme: Theme,
             tone: Optional[str] = None):
    """Tone-coloured circle with a white icon (or 1-2 characters) inside."""
    add_oval(slide, x_in, y_in, d_in, d_in, fill=tone_rgb(theme, tone))
    if not name:
        return
    path = ICON_DIR / f"{name}.png"
    if path.exists():
        pad = d_in * 0.22
        slide.shapes.add_picture(str(path), Inches(x_in + pad), Inches(y_in + pad),
                                 Inches(d_in - 2 * pad), Inches(d_in - 2 * pad))
        return
    label = str(name)
    if len(label) > 2:
        print(f"[mckinsey_pptx] unknown icon {label!r}; available: "
              f"{', '.join(ICONS)}", file=sys.stderr)
        label = label[:1].upper()
    tb = add_textbox(slide, x_in, y_in, d_in, d_in, anchor=MSO_ANCHOR.MIDDLE)
    write_paragraph(tb.text_frame, label, size=max(10, int(d_in * 30)),
                    bold=True, color=theme.palette.white,
                    family=theme.typography.family, align=PP_ALIGN.CENTER,
                    first=True)


# ---------- panels ----------

def add_insight_panel(slide, theme: Theme, x, y, w, h, *, title="Key insight",
                      text: Optional[str] = None, bullets: Iterable[str] = (),
                      size: Optional[int] = None):
    """Tinted side panel: small blue title, bold insight, then bullets."""
    pal = theme.palette
    title = loc(theme, title)
    bullets = list(bullets)
    inner_w = w - 0.5
    paras = [title] + ([text] if text else []) + bullets
    # Shrink the panel to its content when the text is short (no empty box).
    need = text_height_in(paras, inner_w, 14, para_gap_pt=8, indent_in=0.3,
                          bold=True) + 0.6
    if need < h * 0.7:
        h = max(need, min(h, 1.8))
    add_rect(slide, x, y, w, h, fill=pal.soft_gray)
    inner_h = h - 0.5
    if not size:
        size = fit_size(paras, inner_w, inner_h, max_size=14, min_size=10,
                        para_gap_pt=8, indent_in=0.3)
        warn_small("insight panel", text or title, size,
                   "Keep the insight to one sentence + at most 3 short bullets.")
    tb = add_textbox(slide, x + 0.25, y + 0.25, inner_w, inner_h)
    write_paragraph(tb.text_frame, title, size=size, bold=True,
                    color=pal.mid_blue, family=theme.typography.family,
                    first=True, space_after=8)
    if text:
        write_rich_paragraph(tb.text_frame, text, size=size, theme=theme,
                             bold=True, space_after=8)
    for b in bullets:
        write_rich_paragraph(tb.text_frame, b, size=size, theme=theme,
                             bullet=True, space_before=6)


def tint(color: RGBColor, t: float) -> RGBColor:
    """Mix `color` with white (t=0 -> color, t=1 -> white)."""
    return RGBColor(*(round(c + (255 - c) * t) for c in color))


# ---------- connectors and logic shapes ----------

def add_arrow(slide, x1, y1, x2, y2, *, color, width_pt=1.25, dashed=False,
              head=True):
    """Straight line with a triangle arrow head at (x2, y2)."""
    from pptx.oxml.ns import qn
    from lxml import etree
    from .base import add_line
    ln = add_line(slide, x1, y1, x2, y2, color=color, width_pt=width_pt)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        el = ln.line._get_or_add_ln()
        tail = etree.SubElement(el, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return ln


def set_dashed(shape, color, width_pt=1.0):
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.util import Pt
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def add_wedge(slide, x, y, w, h, color, *, direction="right"):
    """Solid triangle pointing right / down inside the visual box (x, y, w, h)
    — the flow marker between logic stages."""
    from pptx.enum.shapes import MSO_SHAPE
    if direction == "down":
        s = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
        s.rotation = 180
    else:  # right: draw a triangle h wide and w tall, then turn it 90 degrees
        cx, cy = x + w / 2, y + h / 2
        s = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(cx - h / 2),
                                   Inches(cy - w / 2), Inches(h), Inches(w))
        s.rotation = 90
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def add_fade_wedge(slide, theme: Theme, x, y, w, h):
    """Wide, flat downward wedge fading from light to navy — 'therefore'."""
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x), Inches(y),
                               Inches(w), Inches(h))
    s.rotation = 180
    s.shadow.inherit = False
    s.line.fill.background()
    s.fill.gradient()
    s.fill.gradient_angle = 90
    stops = s.fill.gradient_stops
    stops[0].color.rgb = theme.palette.mid_blue
    stops[0].position = 0.0
    stops[1].color.rgb = tint(theme.palette.mid_blue, 0.85)
    stops[1].position = 1.0
    return s


def add_stage_banners(slide, theme: Theme, spans, y, h, *, style="chevron",
                      size=None, where=""):
    """Column headers that read as one flow.

    spans: [(x, w, label, kind[, color])] left to right; kind "stage" (light)
    or "result" (dark); color overrides the rule colour. style: "chevron" (arrow banners), "rule" (●—— label ——●),
    "bar" (solid bands).
    """
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Pt
    from .base import add_line
    pal, typo = theme.palette, theme.typography
    spans = [(s[0], s[1], loc(theme, s[2]), *s[3:]) for s in spans]
    if size is None:
        size = 14
        for (x, w, label, _k, *_c) in spans:
            tw = w - (0.5 if style == "rule" else 0.4)
            # one line down to 11pt; otherwise two lines, never breaking a word
            one = fit_one_line(label, tw, 14, 11, bold=True)
            if text_width_pt(label, one, True) <= tw * 72 * 0.95:
                size = min(size, one)
                continue
            words = plain(label).split() or [""]
            size = min(size, fit_size([label], tw, h - 0.04, max_size=14,
                                      min_size=10, line_spacing=1.05, bold=True))
            size = min(size, fit_one_line(max(words, key=len), tw, 14, 10, bold=True))
    warn_small("stage headers", where, size, "Shorten the column headers.")
    for i, (x, w, label, kind, *custom) in enumerate(spans):
        dark = kind == "result"
        if style == "rule":
            color = custom[0] if custom else (pal.deep_navy if dark else pal.mid_blue)
            tw = min(text_width_pt(label, size, True) / 72 / 0.92 + 0.3, w - 0.3)
            tb = add_textbox(slide, x + (w - tw) / 2, y, tw, h, anchor=MSO_ANCHOR.MIDDLE)
            write_paragraph(tb.text_frame, label, size=size, bold=True, color=color,
                            family=typo.family, align=PP_ALIGN.CENTER, first=True)
            ly = y + h / 2
            d = 0.07
            for (a, b) in ((x + 0.05, x + (w - tw) / 2 - 0.05),
                           (x + (w + tw) / 2 + 0.05, x + w - 0.05)):
                if b - a > 0.15:
                    add_line(slide, a, ly, b, ly, color=color, width_pt=1.25)
            add_oval(slide, x, ly - d / 2, d, d, fill=color)
            add_oval(slide, x + w - d, ly - d / 2, d, d, fill=color)
            continue
        if style == "bar":
            shp = add_rect(slide, x, y, w, h, fill=pal.deep_navy if dark else pal.mid_blue)
            inset = 0.1
        else:
            kind_shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
            ext = 0.12 if i < len(spans) - 1 else 0.0  # last banner stays inside the margin
            shp = slide.shapes.add_shape(kind_shape, Inches(x), Inches(y), Inches(w + ext),
                                         Inches(h))
            shp.adjustments[0] = 0.35
            shp.shadow.inherit = False
            shp.line.fill.background()
            shp.fill.solid()
            shp.fill.fore_color.rgb = pal.deep_navy if dark else tint(pal.mid_blue, 0.82)
            inset = 0.25
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(inset)
        tf.margin_top = tf.margin_bottom = Emu(0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        color = pal.white if (dark or style == "bar") else pal.deep_navy
        write_paragraph(tf, label, size=size, bold=True, color=color, family=typo.family,
                        align=PP_ALIGN.CENTER, first=True)
    return size


def add_callout_bar(slide, theme: Theme, x, y, w, h, text, *,
                    label: Optional[str] = "Key insight", icon="lightbulb",
                    tone="blue", size: Optional[int] = None):
    """Full-width highlighted takeaway bar with an icon on the left."""
    pal = theme.palette
    label = loc(theme, label)
    add_rect(slide, x, y, w, h, fill=pal.soft_gray)
    d = min(0.42, h - 0.16)
    add_icon(slide, icon, x + 0.18, y + (h - d) / 2, d, theme, tone)
    tx = x + 0.18 + d + 0.2
    tw = w - (tx - x) - 0.2
    full = f"{label}: {text}" if label else text
    size = size or fit_size([full], tw, h - 0.12, max_size=15, min_size=10)
    tb = add_textbox(slide, tx, y, tw, h, anchor=MSO_ANCHOR.MIDDLE)
    p = write_paragraph(tb.text_frame, "", size=size, family=theme.typography.family,
                        first=True)
    p.runs[0]._r.getparent().remove(p.runs[0]._r)
    if label:
        set_run(p.add_run(), f"{label}: ", size=size, bold=True,
                color=tone_rgb(theme, tone), family=theme.typography.family)
    add_rich_runs(p, text, size=size, color=pal.text_dark,
                  family=theme.typography.family, theme=theme)


# ---------- focus, eyebrows and legend (shared by the diagram templates) ----------
#
# Rules borrowed from editorial diagram practice and applied in the house
# palette: one focal accent per slide (the item the action title is about),
# every other node neutral; structure drawn as hairlines rather than filled
# blocks; small tracked caps ("eyebrows") for axis, quadrant and column
# labels; the legend as one strip under the diagram.

HAIRLINE_PT = 0.75


def _focus_keys(focus):
    if focus is None or focus is False:
        return ()
    if isinstance(focus, (str, int)):
        return (focus,)
    return tuple(focus)


def has_focus(focus) -> bool:
    return bool(_focus_keys(focus))


def is_focus(focus, index: int, label=None) -> bool:
    """True when `focus` names this item by position (int) or label (str,
    case-insensitive, markup ignored)."""
    lab = plain(label).strip().lower() if label is not None else None
    for k in _focus_keys(focus):
        if isinstance(k, bool):
            continue
        if isinstance(k, int) and k == index:
            return True
        if isinstance(k, str) and lab is not None and plain(k).strip().lower() == lab:
            return True
    return False


def unmatched_focus(focus, labels) -> list:
    """Focus keys that name no item (a typo would silently drop the accent)."""
    labs = [plain(l).strip().lower() for l in labels]
    out = []
    for k in _focus_keys(focus):
        if isinstance(k, int) and not isinstance(k, bool):
            if not 0 <= k < len(labs):
                out.append(k)
        elif isinstance(k, str) and plain(k).strip().lower() not in labs:
            out.append(k)
    return out


def check_focus(where: str, title: str, focus, labels):
    bad = unmatched_focus(focus, labels)
    if bad:
        warn_small(where, title, 0, f"focus {bad!r} matches no item; "
                   f"use an item label or a 0-based index.")


def muted_fill(theme: Theme):
    """Fill for non-focal items when one item carries the accent."""
    return tint(theme.palette.dark_navy, 0.62)


def eyebrow(slide, theme: Theme, x, y, w, text, *, color=None,
            align=PP_ALIGN.LEFT, size=10, h=0.24, anchor=MSO_ANCHOR.MIDDLE):
    """Small tracked caps label (axis, quadrant, column, step number).
    CJK text has no caps and loses legibility when tracked: drawn plain, 1pt larger."""
    from .metrics import _CJK
    text = str(loc(theme, text))
    cjk = bool(_CJK.search(text))
    tb = add_textbox(slide, x, y, w, h, anchor=anchor)
    p = write_paragraph(tb.text_frame, text if cjk else text.upper(),
                        size=size + 1 if cjk else size, bold=True,
                        color=color or theme.palette.footer_gray,
                        family=theme.typography.family, align=align, first=True)
    if not cjk:
        for r in p.runs:
            r._r.get_or_add_rPr().set("spc", "120")
    return tb


def eyebrow_width(theme: Theme, text, size=10) -> float:
    from .metrics import _CJK
    text = str(loc(theme, text))
    if _CJK.search(text):
        return text_width_pt(text, size + 1, True) / 72
    return (text_width_pt(text.upper(), size, True) + 1.2 * len(text)) / 72


def focus_tag(slide, theme: Theme, x, y, text, *, h=0.24, size=9, anchor="left"):
    """Bright-blue tag with a white eyebrow ("WE ARE HERE"). Returns its width.
    anchor="right" puts the tag's right edge at x."""
    w = eyebrow_width(theme, text, size) + 0.2
    if anchor == "right":
        x -= w
    add_rect(slide, x, y, w, h, fill=theme.palette.bright_blue)
    eyebrow(slide, theme, x, y, w, text, color=theme.palette.white,
            align=PP_ALIGN.CENTER, size=size, h=h)
    return w


def legend_strip(slide, theme: Theme, items, y, *, x=None, w=None, note=None):
    """One horizontal legend under the diagram, above the footer.
    items: [(kind, color, label)] with kind "dot" | "ring" | "box" | "line" | "dash"."""
    from .base import add_line
    pal, lay = theme.palette, theme.layout
    x = lay.margin_left_in if x is None else x
    w = (lay.slide_width_in - lay.margin_right_in - x) if w is None else w
    add_line(slide, x, y, x + w, y, color=pal.grid_gray, width_pt=HAIRLINE_PT)
    cx = x
    for kind, color, label in items:
        label = str(loc(theme, label))
        if kind == "dot":
            add_oval(slide, cx, y + 0.12, 0.16, 0.16, fill=color)
        elif kind == "ring":
            add_oval(slide, cx, y + 0.12, 0.16, 0.16, line=color, line_width=1.25)
        elif kind == "box":
            add_rect(slide, cx, y + 0.12, 0.22, 0.16, fill=color)
        elif kind in ("line", "dash"):
            ln = add_line(slide, cx, y + 0.2, cx + 0.3, y + 0.2, color=color, width_pt=1.75)
            if kind == "dash":
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        lw = text_width_pt(label, 10) / 72 + 0.1
        tb = add_textbox(slide, cx + 0.38, y + 0.08, lw, 0.24, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, label, size=10, color=pal.text_dark,
                        family=theme.typography.family, first=True)
        cx += 0.38 + lw + 0.4
    if note:
        tb = add_textbox(slide, cx, y + 0.08, max(x + w - cx, 0.5), 0.24,
                         anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, str(loc(theme, note)), size=10, italic=True,
                        color=pal.footer_gray, family=theme.typography.family,
                        align=PP_ALIGN.RIGHT, first=True)
    return 0.4


def set_alpha(shape, opacity: float):
    """Make a solid-filled shape partly transparent (0 = clear, 1 = opaque)."""
    from pptx.oxml.ns import qn
    from lxml import etree
    clr = shape.fill._xPr.find(qn("a:solidFill"))
    if clr is None or not len(clr):
        return
    c = clr[0]
    for old in c.findall(qn("a:alpha")):
        c.remove(old)
    etree.SubElement(c, qn("a:alpha")).set("val", str(int(round(opacity * 100000))))


def fmt_num(v, fmt: Optional[str] = None) -> str:
    """Format a data value: `fmt` is a format string ("${:,.1f}B"), else a plain
    number with thousands separators and at most one decimal."""
    if fmt:
        return fmt.format(v)
    v = float(v)
    if abs(v - round(v)) < 1e-9 or abs(v) >= 100:
        return f"{v:,.0f}"
    return f"{v:,.1f}"


def pct(v: float) -> str:
    return f"{v * 100:.0f}%"


def nice_bounds(lo: float, hi: float, n: int = 5, include_zero=False):
    """Round axis bounds and a step that contain [lo, hi]."""
    import math as _m
    if include_zero:
        lo, hi = min(lo, 0.0), max(hi, 0.0)
    if hi <= lo:
        hi = lo + 1
    raw = (hi - lo) / n
    mag = 10 ** _m.floor(_m.log10(raw))
    step = next(s * mag for s in (1, 2, 2.5, 5, 10) if s * mag >= raw)
    return _m.floor(lo / step) * step, _m.ceil(hi / step) * step, step


# Sequence numbers (01, Step 02, L3 ...) are layout, not data: name them so the
# deck checker does not read them as numbers that need a source.
INDEX_NAME = "chrome:index"


def mark_index(shape):
    """Tag a step / sequence-number shape as layout chrome; returns the shape."""
    shape.name = INDEX_NAME
    return shape
