"""Card grids: the workhorse layout for structured content.

- card_grid: 2-8 cards (icon + header + body/bullets, or a big value),
  optional intro banner on top and key-insight bar at the bottom.
  Text size is fitted to the space, so short content gets larger type
  instead of empty cards.
- swot: a 2x2 card grid with fixed S / W / O / T colours and icons.
"""
from __future__ import annotations

from typing import Dict, Optional, Sequence

from pptx.enum.text import MSO_ANCHOR

from ..base import add_chrome, add_rect, add_textbox, blank_slide, write_paragraph
from ..design import (plain, warn_small, add_callout_bar, add_icon, fit_one_line, fit_size, text_height_in,
                      tone_rgb, write_rich_paragraph)
from ..theme import Theme, DEFAULT_THEME

GAP = 0.25
PAD = 0.22


def _auto_columns(n: int, has_values: bool) -> int:
    if n <= 3:
        return n
    if n == 4:
        return 4 if has_values else 2
    if n <= 6:
        return 3
    return 4


def _card_paras(card: Dict):
    paras = []
    if card.get("body"):
        paras.append(card["body"])
    paras += list(card.get("bullets", []))
    return paras


def draw_cards(slide, theme: Theme, left, top, width, height, cards: Sequence[Dict], *,
               columns: Optional[int] = None, where: str = "", max_size: int = 18):
    """Card grid inside a box (the body of card_grid; also a composite region)."""
    pal = theme.palette
    n = max(len(cards), 1)
    has_values = any(c.get("value") for c in cards)
    cols = columns or _auto_columns(n, has_values)
    rows = -(-n // cols)
    card_w = (width - GAP * (cols - 1)) / cols
    card_h = (height - GAP * (rows - 1)) / rows

    icon_d = 0.5 if card_h >= 1.6 else 0.38
    has_icons = any(c.get("icon") for c in cards)
    if (has_icons and not has_values and card_h < 1.6 and card_w >= 2.8 * card_h):
        # Wide, short cards (e.g. a stack in a composite column): stacking icon,
        # header and body leaves the body ~0.2" and forces 10pt. Put the icon
        # on the left and header + body in one text column instead.
        return _draw_side_cards(slide, theme, left, top, card_w, card_h, cols, cards,
                                where=where, max_size=max_size)
    header_w = card_w - 2 * PAD - ((icon_d + 0.15) if has_icons and not has_values else 0)
    value_h = min(1.3, card_h * 0.36) if has_values else 0

    # One body size for the whole grid, fitted to the tightest card.
    body_w = card_w - 2 * PAD
    sizes = []
    for c in cards:
        head_size_guess = 15
        head_h = text_height_in([c.get("title", "")], header_w, head_size_guess)
        head_h = max(head_h, icon_d if has_icons and not has_values else 0)
        avail = card_h - 2 * PAD - head_h - value_h - 0.12
        sizes.append(fit_size(_card_paras(c) or [" "], body_w, max(avail, 0.3),
                              max_size=max_size, min_size=10, para_gap_pt=5,
                              indent_in=0.25))
    body_size = min(sizes) if sizes else 14
    warn_small("card_grid", where, body_size,
               "Shorten the longest card's bullets, drop the subtitle or the "
               "insight bar, or split the slide.")
    head_size = min(body_size + 2, 20)
    # Headers never break inside a word: shrink until the longest word fits.
    for c in cards:
        hw_ = header_w if (has_icons and not c.get("value")) else body_w
        words = plain(c.get("title", "")).split() or [""]
        longest = max(words, key=len)
        head_size = min(head_size, fit_one_line(longest, hw_, head_size, 11, bold=True))
    if columns and card_w < 2.6 and not has_values:
        warn_small("card_grid", where, 0,
                   f"columns={columns} makes cards {card_w:.1f}\" wide — text cards "
                   "read better in the default layout (4 cards -> 2x2); drop columns=.")

    # If the text tops out at the size cap, shrink the cards to their content
    # and centre the grid, rather than leaving the space inside each card.
    def needed(c):
        hw_ = header_w if (has_icons and not c.get("value")) else body_w
        h = 2 * PAD + text_height_in([c.get("title", "")], hw_, head_size, bold=True)
        if has_icons and not c.get("value"):
            h = max(h, 2 * PAD + icon_d)
        h += 0.12 + (value_h + 0.05 if c.get("value") else 0)
        paras = _card_paras(c)
        if paras:
            h += text_height_in(paras, body_w, body_size, para_gap_pt=5,
                                indent_in=0.25)
        return h + 0.1
    natural = max((needed(c) for c in cards), default=card_h)
    if natural < card_h * 0.85:
        grid_h = rows * card_h + GAP * (rows - 1)
        card_h = max(natural, card_h * 0.55)
        top += (grid_h - (rows * card_h + GAP * (rows - 1))) / 2

    for i, c in enumerate(cards):
        r, k = divmod(i, cols)
        x = left + k * (card_w + GAP)
        y = top + r * (card_h + GAP)
        tone = c.get("tone")
        add_rect(slide, x, y, card_w, card_h, fill=pal.soft_gray)
        cy = y + PAD

        if c.get("value"):
            # one line, never broken mid-word
            vsize = min(fit_one_line(c["value"], body_w, 54, 16, bold=True),
                        int(value_h * 72 / 1.15))
            tb = add_textbox(slide, x + PAD, cy, body_w, value_h,
                             anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, c["value"], size=vsize,
                                 theme=theme, color=tone_rgb(theme, tone),
                                 bold=True, first=True)
            cy += value_h + 0.05
            hx, hw = x + PAD, body_w
        elif has_icons:
            add_icon(slide, c.get("icon"), x + PAD, cy, icon_d, theme, tone)
            hx, hw = x + PAD + icon_d + 0.15, header_w
        else:
            hx, hw = x + PAD, body_w

        head_h = text_height_in([c.get("title", "")], hw, head_size, bold=True)
        if has_icons and not c.get("value"):
            head_h = max(head_h, icon_d)
        tb = add_textbox(slide, hx, cy, hw, head_h,
                         anchor=MSO_ANCHOR.MIDDLE if has_icons and not c.get("value")
                         else MSO_ANCHOR.TOP)
        write_rich_paragraph(tb.text_frame, c.get("title", ""), size=head_size,
                             theme=theme, color=pal.text_dark, bold=True,
                             first=True)
        cy += head_h + 0.12

        paras = _card_paras(c)
        if paras:
            tb = add_textbox(slide, x + PAD, cy, body_w, y + card_h - PAD - cy)
            first = True
            if c.get("body"):
                write_rich_paragraph(tb.text_frame, c["body"], size=body_size,
                                     theme=theme, first=True, space_after=4)
                first = False
            for b in c.get("bullets", []):
                write_rich_paragraph(tb.text_frame, b, size=body_size,
                                     theme=theme, bullet=True, first=first,
                                     space_before=0 if first else 5)
                first = False


def _draw_side_cards(slide, theme: Theme, left, top, card_w, card_h, cols, cards, *,
                     where: str = "", max_size: int = 18):
    """Wide-card layout: icon at the left, header + body beside it."""
    pal = theme.palette
    pad = 0.16
    icon_d = min(0.46, card_h - 2 * pad)
    tx_off = pad + icon_d + 0.18
    tw = card_w - tx_off - pad
    avail = card_h - 2 * pad

    def need(c, size):
        h = text_height_in([c.get("title", "")], tw, size + 1, bold=True) + 0.06
        paras = _card_paras(c)
        if paras:
            h += text_height_in(paras, tw, size, para_gap_pt=4, indent_in=0.25)
        return h

    size = 10
    for s in range(max_size, 9, -1):
        if all(need(c, s) <= avail for c in cards):
            size = s
            break
    warn_small("card_grid", where, size,
               "Shorten the longest card's text, drop the insight bar, or split the slide.")
    for i, c in enumerate(cards):
        r, k = divmod(i, cols)
        x = left + k * (card_w + GAP)
        y = top + r * (card_h + GAP)
        add_rect(slide, x, y, card_w, card_h, fill=pal.soft_gray)
        add_icon(slide, c.get("icon"), x + pad, y + (card_h - icon_d) / 2, icon_d, theme,
                 c.get("tone"))
        tb = add_textbox(slide, x + tx_off, y + pad, tw, avail, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, c.get("title", ""), size=size + 1, theme=theme,
                             color=pal.text_dark, bold=True, first=True, space_after=3)
        if c.get("body"):
            write_rich_paragraph(tb.text_frame, c["body"], size=size, theme=theme,
                                 space_after=2)
        for b_ in c.get("bullets", []):
            write_rich_paragraph(tb.text_frame, b_, size=size, theme=theme, bullet=True,
                                 space_before=3)


def add_card_grid(prs, *,
                  title: str = "[Card grid / Insert action title]",
                  cards: Sequence[Dict],
                  subtitle: Optional[str] = None,
                  intro: Optional[str] = None,
                  insight: Optional[str] = None,
                  insight_label: Optional[str] = "Key insight",
                  columns: Optional[int] = None,
                  page_number=None, section_marker=None,
                  source=None, footnote=None,
                  theme: Theme = DEFAULT_THEME):
    """cards: [{"title": str, "body"?: str, "bullets"?: [str], "icon"?: str,
                "tone"?: str, "value"?: str}]
    - icon: a bundled icon name (see CATALOG) or 1-2 characters ("1", "A").
    - tone: navy | blue | mid_blue | light_blue | red | green | amber | gray.
    - value: big number shown at the top of the card (KPI / target cards).
    Strings accept **bold** and {tone|coloured} markup.
    """
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo, layout = theme.palette, theme.typography, theme.layout
    left = layout.margin_left_in
    width = layout.slide_width_in - layout.margin_left_in - layout.margin_right_in
    top = layout.body_top_in + 0.05
    bottom = layout.footer_top_in - 0.25

    if subtitle:
        tb = add_textbox(slide, left, top, width, 0.35)
        write_paragraph(tb.text_frame, subtitle, size=typo.section_title_size,
                        bold=True, color=pal.text_dark, family=typo.family,
                        first=True)
        top += 0.45
    if intro:
        h = 0.75
        add_rect(slide, left, top, width, h, fill=pal.deep_navy)
        size = fit_size([intro], width - 0.6, h - 0.15, max_size=16, min_size=11)
        tb = add_textbox(slide, left + 0.3, top, width - 0.6, h,
                         anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, intro, size=size, theme=theme,
                             color=pal.white, first=True)
        top += h + GAP
    if insight:
        h = 0.7
        add_callout_bar(slide, theme, left, bottom - h, width, h, insight,
                        label=insight_label)
        bottom -= h + GAP

    draw_cards(slide, theme, left, top, width, bottom - top, cards,
               columns=columns, where=title)
    return slide


def add_swot(prs, *,
             title: str = "[SWOT / Insert action title]",
             strengths: Sequence[str] = (),
             weaknesses: Sequence[str] = (),
             opportunities: Sequence[str] = (),
             threats: Sequence[str] = (),
             labels: Sequence[str] = ("Strengths", "Weaknesses",
                                      "Opportunities", "Threats"),
             insight: Optional[str] = None,
             subtitle: Optional[str] = None,
             page_number=None, section_marker=None,
             source=None, footnote=None,
             theme: Theme = DEFAULT_THEME):
    """Classic 2x2 SWOT. Empty quadrants are kept (shown as "—") so the frame
    stays readable; pass only what the source supports."""
    from ..labels import loc
    labels = [loc(theme, t) for t in labels]
    quads = [
        (labels[0], strengths, "shield", "blue"),
        (labels[1], weaknesses, "alert", "red"),
        (labels[2], opportunities, "trend_up", "green"),
        (labels[3], threats, "zap", "amber"),
    ]
    cards = [{"title": t, "bullets": list(items) or ["—"], "icon": ic, "tone": tone}
             for t, items, ic, tone in quads]
    return add_card_grid(prs, title=title, cards=cards, subtitle=subtitle,
                         insight=insight, columns=2, page_number=page_number,
                         section_marker=section_marker, source=source,
                         footnote=footnote, theme=theme)


def add_card_rows(prs, *,
                  title: str = "[Card rows / Insert action title]",
                  rows: Sequence[Dict],
                  subtitle: Optional[str] = None,
                  insight: Optional[str] = None,
                  insight_label: Optional[str] = "Key insight",
                  label_width: float = 3.4,
                  page_number=None, section_marker=None,
                  source=None, footnote=None,
                  theme: Theme = DEFAULT_THEME):
    """Horizontal list: one band per item — icon + header on the left,
    body / bullets on the right. Same item shape as card_grid:
    rows: [{"title", "body"?, "bullets"?, "icon"?, "tone"?, "value"?}]
    Good for 2-6 items with a sentence or two each (moves, risks, principles,
    decisions), and as a visual change from card_grid.
    """
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo, layout = theme.palette, theme.typography, theme.layout
    left = layout.margin_left_in
    width = layout.slide_width_in - layout.margin_left_in - layout.margin_right_in
    top = layout.body_top_in + 0.05
    bottom = layout.footer_top_in - 0.25
    if subtitle:
        tb = add_textbox(slide, left, top, width, 0.35)
        write_paragraph(tb.text_frame, subtitle, size=typo.section_title_size,
                        bold=True, color=pal.text_dark, family=typo.family,
                        first=True)
        top += 0.45
    if insight:
        h = 0.7
        add_callout_bar(slide, theme, left, bottom - h, width, h, insight,
                        label=insight_label)
        bottom -= h + GAP

    n = max(len(rows), 1)
    gap = 0.15
    row_h = min((bottom - top - gap * (n - 1)) / n, 1.5)
    icon_d = min(0.5, row_h - 0.3)
    has_icons = any(r.get("icon") for r in rows)
    lab_text_w = label_width - PAD - ((icon_d + 0.15) if has_icons else 0) - 0.1
    body_x = left + label_width + 0.2
    body_w = width - label_width - 0.2 - PAD

    sizes = [fit_size(_card_paras(r) or [" "], body_w, row_h - 2 * 0.14,
                      max_size=17, min_size=10, para_gap_pt=4, indent_in=0.25)
             for r in rows]
    body_size = min(sizes) if sizes else 14
    warn_small("card_rows", title, body_size,
               "Shorten the longest row or split the slide.")
    head_size = min(body_size + 2, 20)
    for r in rows:
        words = plain(r.get("title", "")).split() or [""]
        head_size = min(head_size, fit_one_line(max(words, key=len), lab_text_w,
                                                head_size, 11, bold=True))
        head_size = min(head_size, fit_size([r.get("title", "")], lab_text_w,
                                            row_h - 0.2, max_size=head_size,
                                            min_size=11, line_spacing=1.15))

    block_h = n * row_h + gap * (n - 1)
    top += max(0, (bottom - top - block_h) / 2)
    for i, r in enumerate(rows):
        y = top + i * (row_h + gap)
        tone = r.get("tone")
        add_rect(slide, left, y, width, row_h, fill=pal.soft_gray)
        add_rect(slide, left, y, label_width, row_h, fill=pal.light_gray)
        x = left + PAD
        if has_icons:
            add_icon(slide, r.get("icon"), x, y + (row_h - icon_d) / 2, icon_d,
                     theme, tone)
            x += icon_d + 0.15
        tb = add_textbox(slide, x, y, lab_text_w, row_h, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, r.get("title", ""), size=head_size,
                             theme=theme, bold=True, first=True)
        if r.get("value"):
            write_rich_paragraph(tb.text_frame, r["value"], size=head_size + 4,
                                 theme=theme, color=tone_rgb(theme, tone),
                                 bold=True)
        tb = add_textbox(slide, body_x, y + 0.08, body_w, row_h - 0.16,
                         anchor=MSO_ANCHOR.MIDDLE)
        first = True
        if r.get("body"):
            write_rich_paragraph(tb.text_frame, r["body"], size=body_size,
                                 theme=theme, first=True, space_after=3)
            first = False
        for b_ in r.get("bullets", []):
            write_rich_paragraph(tb.text_frame, b_, size=body_size, theme=theme,
                                 bullet=True, first=first,
                                 space_before=0 if first else 4)
            first = False
    return slide
