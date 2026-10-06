"""Generic data table: any header × rows grid, as a native PowerPoint table.

Use it for source tables that don't fit a specialised template (market share
by year, strengths × impact, tier × price × channel, legacy vs. new ...).
Cells accept light markup: **bold** spans are rendered bold.
"""
from __future__ import annotations
import re
from typing import Optional, Sequence

from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from ..base import (add_chrome, add_rect, add_textbox, blank_slide, set_run,
                    write_paragraph, write_rich)
from ..theme import Theme, DEFAULT_THEME

_BOLD = re.compile(r"(\*\*.+?\*\*)")


def _write_cell(cell, text, *, size, color, family, bold=False,
                align=PP_ALIGN.LEFT):
    tf = cell.text_frame
    tf.word_wrap = True
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = str(text).split("\n")
    for li, line in enumerate(lines):
        p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
        p.alignment = align
        for part in _BOLD.split(line):
            if not part:
                continue
            is_bold = part.startswith("**") and part.endswith("**")
            run = p.add_run()
            set_run(run, part[2:-2] if is_bold else part, size=size,
                    bold=bold or is_bold, color=color, family=family)


def _fill(cell, rgb):
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb


def add_data_table(prs, *,
                   title: str = "[Data table / Insert action title]",
                   columns: Sequence[str],
                   rows: Sequence[Sequence],
                   subtitle: Optional[str] = None,
                   col_widths: Optional[Sequence[float]] = None,
                   highlight_rows: Sequence[int] = (),
                   highlight_col: Optional[int] = None,
                   first_col_bold: bool = True,
                   numeric_align: str = "center",
                   insight: Optional[str] = None,
                   insight_title: str = "Key insight",
                   insight_bullets: Sequence[str] = (),
                   font_size: Optional[int] = None,
                   page_number=None, section_marker=None,
                   source=None, footnote=None,
                   theme: Theme = DEFAULT_THEME):
    """columns: header labels. rows: list of row cell lists (str/number).
    col_widths: relative widths per column (default: first column wider).
    highlight_rows: row indices (0-based, body rows) tinted + bold — e.g. "us".
    highlight_col: column index tinted — e.g. the recommended option.
    insight / insight_bullets: optional Key-insight panel on the right.
    """
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo, layout = theme.palette, theme.typography, theme.layout
    # Fewer rows -> larger type, so short tables still read at a distance.
    size = font_size or (typo.body_size + 2 if len(rows) <= 7 else typo.body_size)

    left = layout.margin_left_in
    top = layout.body_top_in + 0.05
    width = layout.slide_width_in - layout.margin_left_in - layout.margin_right_in
    bottom = layout.footer_top_in - 0.3

    if subtitle:
        tb = add_textbox(slide, left, top, width, 0.35)
        write_paragraph(tb.text_frame, subtitle, size=typo.section_title_size,
                        bold=True, color=pal.text_dark, family=typo.family,
                        first=True)
        top += 0.5

    has_panel = bool(insight or insight_bullets)
    panel_w = 3.6 if has_panel else 0
    gap = 0.35 if has_panel else 0
    table_w = width - panel_w - gap

    n_cols = max(len(columns), 1)
    n_rows = len(rows) + 1
    if col_widths is None:
        col_widths = [1.6] + [1.0] * (n_cols - 1) if n_cols > 1 else [1.0]
    total = float(sum(col_widths))
    widths_in = [table_w * w / total for w in col_widths]

    # Row height: share the body height so the table fills the slide.
    avail = bottom - top
    row_h = min(0.9, max(0.32, avail / n_rows))
    table_h = row_h * n_rows
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top),
                                Inches(table_w), Inches(table_h))
    table = gf.table
    # Drop the default PowerPoint table style banding; we colour explicitly.
    tbl_pr = gf._element.graphic.graphicData.tbl.tblPr
    tbl_pr.set("firstRow", "1")
    tbl_pr.set("bandRow", "0")
    for ci, w in enumerate(widths_in):
        table.columns[ci].width = Emu(int(Inches(w)))
    for ri in range(n_rows):
        table.rows[ri].height = Emu(int(Inches(row_h)))

    def is_num(v):
        return bool(re.fullmatch(r"[\s<>≈~+\-−]*[\d.,]+\s*%?\s*", str(v)))

    for ci, label in enumerate(columns):
        cell = table.cell(0, ci)
        _fill(cell, pal.bright_blue if highlight_col == ci else pal.deep_navy)
        _write_cell(cell, label, size=size, color=pal.white, family=typo.family,
                    bold=True, align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER)

    hl = set(highlight_rows)
    for ri, row in enumerate(rows, start=1):
        for ci in range(n_cols):
            val = row[ci] if ci < len(row) else ""
            cell = table.cell(ri, ci)
            if (ri - 1) in hl:
                _fill(cell, theme.palette.light_gray)
            elif highlight_col == ci:
                _fill(cell, theme.palette.soft_gray)
            else:
                _fill(cell, pal.white if ri % 2 else pal.soft_gray)
            align = PP_ALIGN.LEFT
            if ci > 0 and numeric_align == "center" and is_num(val):
                align = PP_ALIGN.CENTER
            _write_cell(cell, val, size=size, color=pal.text_dark,
                        family=typo.family,
                        bold=(ci == 0 and first_col_bold) or (ri - 1) in hl,
                        align=align)

    if has_panel:
        px = left + table_w + gap
        add_rect(slide, px, top, panel_w, table_h, fill=pal.soft_gray)
        tb = add_textbox(slide, px + 0.25, top + 0.25, panel_w - 0.5,
                         table_h - 0.5)
        write_paragraph(tb.text_frame, insight_title, size=size,
                        bold=True, color=pal.mid_blue, family=typo.family,
                        first=True, space_after=8)
        if insight:
            write_rich(tb.text_frame, insight, size=size, bold=True,
                       color=pal.text_dark, family=typo.family, space_after=8)
        for b in insight_bullets:
            write_rich(tb.text_frame, b, size=size, color=pal.text_dark,
                       family=typo.family, bullet=True, space_before=6)
    return slide
