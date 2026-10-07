"""Analysis templates from strategy-case decks.

- cycle             : flywheel / causal chain. "branch": start -> parallel
                      paths of effects -> end, dashed feedback to the start
                      (virtuous cycle). "loop": 3-6 steps around a centre
                      (vicious or virtuous loop).
- risk_heatmap      : risks placed on probability x impact with graded
                      zones, mitigations listed by number on the right.
- positioning_scale : indicators as rows (what it means) with players
                      placed on a low -> high scale per row.
- value_chain       : chevron stages with KSFs / activities per stage; the
                      company's own stages highlighted.
"""
from __future__ import annotations

import math
from typing import Dict, List, Optional, Sequence, Union

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from ..base import add_line, add_oval, add_rect, add_textbox, write_paragraph
from ..components import _polygon, letter_space
from ..design import (add_arrow, add_fade_wedge, add_icon, add_stage_banners, fit_one_line,
                      fit_size, plain, set_dashed, text_height_in, text_width_pt, tint,
                      tone_rgb, warn_small, write_rich_paragraph)
from ..labels import loc
from ..theme import Theme, DEFAULT_THEME
from .evaluation_slides import _frame
from .logic_slides import _cell, _cell_h, _write_cell

MIN_PT = 10


def _as_node(n):
    return n if isinstance(n, dict) else {"title": str(n)}


def _node(slide, theme, x, y, w, h, node, size, *, fill, line=None, color=None, dashed=False):
    pal = theme.palette
    box = add_rect(slide, x, y, w, h, fill=fill, line=line, line_width=1.25 if line else None)
    if dashed:
        set_dashed(box, line or pal.mid_blue, 1.0)
    tx, tw = x + 0.12, w - 0.24
    if node.get("icon"):
        d = min(0.46, h - 0.2)
        add_icon(slide, node["icon"], x + 0.14, y + (h - d) / 2, d, theme, node.get("tone", "blue"))
        tx, tw = x + 0.14 + d + 0.1, w - d - 0.36
    tb = add_textbox(slide, tx, y + 0.04, tw, h - 0.08, anchor=MSO_ANCHOR.MIDDLE)
    write_rich_paragraph(tb.text_frame, node.get("title", ""), size=size, theme=theme,
                         color=color or pal.text_dark, bold=True, align=PP_ALIGN.CENTER,
                         first=True)
    if node.get("body"):
        write_rich_paragraph(tb.text_frame, node["body"], size=max(size - 2, MIN_PT), theme=theme,
                             color=color or pal.footer_gray, align=PP_ALIGN.CENTER)


def _node_size(nodes, w, h, has_icon=False, max_size=16):
    tw = w - 0.24 - (0.56 if has_icon else 0)
    size = max_size
    for n in nodes:
        paras = [n.get("title", "")] + ([n["body"]] if n.get("body") else [])
        size = min(size, fit_size(paras, tw, h - 0.12, max_size=max_size, min_size=MIN_PT,
                                  bold=True, line_spacing=1.1))
        words = plain(n.get("title", "")).split() or [""]
        size = min(size, fit_one_line(max(words, key=len), tw, max_size, MIN_PT, True))
    return size


def _label(slide, theme, x, y, w, text, size=11):
    tb = add_textbox(slide, x, y, w, 0.26, anchor=MSO_ANCHOR.BOTTOM)
    write_paragraph(tb.text_frame, text, size=size, italic=True, bold=True,
                    color=theme.palette.mid_blue, family=theme.typography.family,
                    align=PP_ALIGN.CENTER, first=True)


def _conclusion(slide, theme, left, width, y, h, text):
    pal = theme.palette
    add_fade_wedge(slide, theme, left + width * 0.34, y, width * 0.32, 0.26)
    by = y + 0.36
    bh = h - 0.36
    add_rect(slide, left, by, width, bh, fill=pal.white, line=pal.mid_blue, line_width=1.75)
    size = fit_size([text], width - 0.6, bh - 0.1, max_size=17, min_size=12, bold=True)
    tb = add_textbox(slide, left + 0.3, by, width - 0.6, bh, anchor=MSO_ANCHOR.MIDDLE)
    write_rich_paragraph(tb.text_frame, text, size=size, theme=theme, color=pal.deep_navy,
                         bold=True, align=PP_ALIGN.CENTER, first=True)


# ---------- cycle ----------

def add_cycle(prs, *,
              title: str = "[Cycle / Insert action title]",
              start: Union[str, Dict, None] = None,
              paths: Optional[Sequence[Sequence[Union[str, Dict]]]] = None,
              end: Union[str, Dict, None] = None,
              steps: Optional[Sequence[Union[str, Dict]]] = None,
              center: Optional[str] = None,
              kind: str = "virtuous",
              feedback: Union[bool, str] = True,
              connector_label: Optional[str] = None,
              conclusion: Optional[str] = None,
              subtitle: Optional[str] = None,
              insight: Optional[str] = None,
              insight_label: Optional[str] = "Key insight",
              page_number=None, section_marker=None,
              source=None, footnote=None,
              theme: Theme = DEFAULT_THEME):
    """Branch flywheel: start -> paths (each a list of steps) -> end, with a
    dashed feedback arrow from end back to start.
        start / end: str | {title, body?, icon?}; paths: [[step, step], ...]
        connector_label: text on the arrows ("leads to ...").
    Loop: steps (3-6) around a centre; kind "vicious" draws it in red.
    conclusion: bold statement under the diagram (the "so what").
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    if conclusion:
        ch = 0.95
        _conclusion(slide, theme, left, width, bottom - ch, ch, conclusion)
        bottom -= ch + 0.12
    if steps:
        _loop(slide, theme, left, top, width, bottom, [_as_node(s) for s in steps], center,
              kind, title)
    else:
        _branch(slide, theme, left, top, width, bottom, _as_node(start or ""),
                [[_as_node(s) for s in p] for p in (paths or [])], _as_node(end or ""),
                feedback, connector_label, title)
    return slide


def _branch(slide, theme, left, top, width, bottom, start, paths, end, feedback, clabel, where):
    pal = theme.palette
    n = max(len(paths), 1)
    k = max((len(p) for p in paths), default=1)
    sw = ew = min(2.0, width * 0.15)
    conn = 1.15 if clabel else 0.7  # horizontal gap holding an arrow (+ label)
    fb_h = 0.45 if feedback else 0.0
    step_w = (width - sw - ew - conn * (k + 1)) / k
    band_top = top + (0.3 if clabel else 0.1)
    band_bot = bottom - fb_h
    gap = 0.32
    nh = min(1.15, (band_bot - band_top - gap * (n - 1)) / n)
    block = n * nh + gap * (n - 1)
    y0 = band_top + (band_bot - band_top - block) / 2
    ys = [y0 + i * (nh + gap) for i in range(n)]
    mid = y0 + block / 2
    all_steps = [s for p in paths for s in p]
    has_icon = any(s.get("icon") for s in all_steps)
    size = _node_size(all_steps or [{"title": " "}], step_w, nh, has_icon)
    se_h = min(1.4, max(nh, 1.0))
    se_size = _node_size([start, end], sw, se_h, False, 17)
    size = min(size, se_size + 1)
    warn_small("cycle", where, size, "Shorten the step titles or use fewer paths.")
    line_c = pal.mid_blue
    # start node
    _node(slide, theme, left, mid - se_h / 2, sw, se_h, start, se_size,
          fill=tint(pal.light_blue, 0.82), line=pal.mid_blue)
    bx = left + sw + 0.3  # branch line
    add_line(slide, left + sw, mid, bx, mid, color=line_c, width_pt=1.25)
    if n > 1:
        add_line(slide, bx, ys[0] + nh / 2, bx, ys[-1] + nh / 2, color=line_c, width_pt=1.25)
    ex = left + width - ew  # end node x
    mx = ex - 0.3  # merge line
    for i, path in enumerate(paths):
        cy = ys[i] + nh / 2
        x = left + sw + conn
        add_arrow(slide, bx, cy, x - 0.02, cy, color=line_c)
        if clabel:
            _label(slide, theme, bx - 0.25, cy - 0.3, x - bx + 0.25, clabel)
        for j, st in enumerate(path):
            _node(slide, theme, x, ys[i], step_w, nh, st, size, fill=pal.soft_gray,
                  line=pal.mid_blue)
            if j < len(path) - 1:
                nx = x + step_w + conn
                add_arrow(slide, x + step_w, cy, nx - 0.02, cy, color=line_c)
                if clabel:
                    _label(slide, theme, x + step_w, cy - 0.3, conn, clabel)
                x = nx
            else:
                x += step_w
        add_line(slide, x, cy, mx, cy, color=line_c, width_pt=1.25)
    if n > 1:
        add_line(slide, mx, ys[0] + nh / 2, mx, ys[-1] + nh / 2, color=line_c, width_pt=1.25)
    add_arrow(slide, mx, mid, ex - 0.02, mid, color=line_c)
    if clabel:
        _label(slide, theme, mx - 0.6, mid - 0.3, ex - mx + 0.6, clabel)
    _node(slide, theme, ex, mid - se_h / 2, ew, se_h, end, se_size,
          fill=tint(pal.light_blue, 0.82), line=pal.mid_blue)
    if feedback:
        fy = bottom - 0.12
        sx, exx = left + sw / 2, ex + ew / 2
        for (a, b, c, d) in ((exx, mid + se_h / 2, exx, fy), (exx, fy, sx, fy)):
            ln = add_line(slide, a, b, c, d, color=line_c, width_pt=1.25)
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        add_arrow(slide, sx, fy, sx, mid + se_h / 2 + 0.02, color=line_c, dashed=True)
        if isinstance(feedback, str):
            tb = add_textbox(slide, left + width * 0.3, fy - 0.3, width * 0.4, 0.28,
                             anchor=MSO_ANCHOR.BOTTOM)
            write_paragraph(tb.text_frame, feedback, size=11, italic=True, bold=True,
                            color=pal.mid_blue, family=theme.typography.family,
                            align=PP_ALIGN.CENTER, first=True)


def _rect_exit(cx, cy, a, b, dx, dy):
    """Distance from a rectangle's centre to its edge along (dx, dy)."""
    tx = a / abs(dx) if dx else float("inf")
    ty = b / abs(dy) if dy else float("inf")
    return min(tx, ty)


def _loop(slide, theme, left, top, width, bottom, steps, center, kind, where):
    pal = theme.palette
    n = max(len(steps), 3)
    accent = pal.status_red if kind == "vicious" else pal.mid_blue
    H = bottom - top
    cx, cy = left + width / 2, top + H / 2
    nw, nh = min(2.6, width / 4.2), min(1.0, H / 4.2)
    rx, ry = min(width / 2 - nw / 2 - 0.2, H * 1.15), H / 2 - nh / 2 - 0.05
    pos = []
    for i in range(len(steps)):
        ang = -math.pi / 2 + 2 * math.pi * i / len(steps)
        pos.append((cx + rx * math.cos(ang), cy + ry * math.sin(ang)))
    size = _node_size(steps, nw, nh, any(s.get("icon") for s in steps), 15)
    warn_small("cycle", where, size, "Shorten the loop steps.")
    for i, (px, py) in enumerate(pos):
        qx, qy = pos[(i + 1) % len(pos)]
        dx, dy = qx - px, qy - py
        L = math.hypot(dx, dy) or 1
        ux, uy = dx / L, dy / L
        s0 = _rect_exit(px, py, nw / 2, nh / 2, ux, uy) + 0.08
        s1 = _rect_exit(qx, qy, nw / 2, nh / 2, ux, uy) + 0.1
        if L > s0 + s1 + 0.1:
            add_arrow(slide, px + ux * s0, py + uy * s0, qx - ux * s1, qy - uy * s1,
                      color=accent, width_pt=2.0)
    for (px, py), st in zip(pos, steps):
        _node(slide, theme, px - nw / 2, py - nh / 2, nw, nh, st, size,
              fill=tint(accent, 0.88), line=accent)
    if center:
        d = min(2.0, H * 0.42, rx * 1.1)
        add_oval(slide, cx - d / 2, cy - d / 2, d, d, fill=pal.deep_navy if kind != "vicious"
                 else pal.status_red)
        csz = fit_size([center], d * 0.72, d * 0.6, max_size=16, min_size=10, bold=True)
        tb = add_textbox(slide, cx - d * 0.36, cy - d * 0.3, d * 0.72, d * 0.6,
                         anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, center, size=csz, theme=theme, color=pal.white,
                             bold=True, align=PP_ALIGN.CENTER, first=True)


# ---------- risk heat map ----------

_LEVEL = {"very low": 0.1, "low": 0.22, "medium": 0.5, "med": 0.5, "moderate": 0.5,
          "high": 0.8, "very high": 0.92, "critical": 0.92}


def _norm_levels(vals):
    nums = [v for v in vals if not isinstance(v, str)]
    hi = max(nums, default=1)
    out = []
    for v in vals:
        if isinstance(v, str):
            out.append(_LEVEL.get(v.strip().lower(), 0.5))
        elif hi <= 1:
            out.append(float(v))
        elif hi <= 3:
            out.append((float(v) - 0.5) / 3)
        elif hi <= 5:
            out.append((float(v) - 0.5) / 5)
        else:
            out.append(float(v) / hi)
    return [min(max(v, 0.04), 0.96) for v in out]


def _zone(slide, x0, y0, W, H, r, fill):
    """Region of the plot further than r (normalised) from the low-low corner."""
    pts = []
    steps = 24
    if r <= 1:
        for k in range(steps + 1):
            a = (math.pi / 2) * k / steps
            pts.append((r * math.cos(a), r * math.sin(a)))
        pts += [(0, 1), (1, 1), (1, 0)]
    else:
        a0, a1 = math.acos(1 / r), math.asin(1 / r)
        for k in range(steps + 1):
            a = a0 + (a1 - a0) * k / steps
            pts.append((r * math.cos(a), r * math.sin(a)))
        pts += [(1, 1)]
    pts = [(x0 + u * W, y0 + H - v * H) for u, v in pts]
    _polygon(slide, pts, fill)


def add_risk_heatmap(prs, *,
                     title: str = "[Risk heat map / Insert action title]",
                     risks: Sequence[Dict],
                     x_label: str = "Probability", y_label: str = "Impact",
                     ends: Sequence[str] = ("Low", "High"),
                     map_label: Optional[str] = "Risk mapping",
                     mitigation_label: Optional[str] = "Mitigations",
                     zones: str = "blue",
                     subtitle: Optional[str] = None,
                     insight: Optional[str] = None,
                     insight_label: Optional[str] = "Bottom line",
                     page_number=None, section_marker=None,
                     source=None, footnote=None,
                     theme: Theme = DEFAULT_THEME):
    """risks: [{"title", "probability", "impact", "mitigation"? (str | [str])}]
    probability / impact: "low" | "medium" | "high", 1-3, 1-5 or 0-1.
    zones: "blue" (shades of the brand colour) or "traffic" (green/amber/red).
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    has_mit = any(r.get("mitigation") for r in risks)
    rw = width * 0.37 if has_mit else 0.0
    gap = 0.35 if has_mit else 0.0
    lw = width - rw - gap
    head_h = 0.42
    spans = [(left, lw, map_label or "", "stage")] if map_label else []
    if has_mit:
        spans.append((left + lw + gap, rw, mitigation_label or "", "stage"))
    if spans:
        add_stage_banners(slide, theme, spans, top, head_h, style="bar", where=title)
        top += head_h + 0.15

    # plot box
    ax_w = 0.55  # y label + end labels
    px, py = left + ax_w, top + 0.28
    pw, ph = lw - ax_w - 0.1, bottom - 0.55 - py
    if zones == "traffic":
        shades = [tint(pal.status_green, 0.75), tint(pal.status_amber, 0.55),
                  tint(pal.status_red, 0.55)]
        base = tint(pal.status_green, 0.88)
    else:
        shades = [tint(pal.mid_blue, 0.78), tint(pal.mid_blue, 0.55), tint(pal.mid_blue, 0.3)]
        base = tint(pal.mid_blue, 0.93)
    add_rect(slide, px, py, pw, ph, fill=base)
    for r_, c in zip((0.45, 0.8, 1.08), shades):
        _zone(slide, px, py, pw, ph, r_, c)
    add_rect(slide, px, py, pw, ph, fill=None, line=pal.deep_navy, line_width=1.0)
    add_arrow(slide, px - 0.12, py + ph + 0.12, px + pw + 0.05, py + ph + 0.12,
              color=pal.deep_navy)
    add_arrow(slide, px - 0.12, py + ph + 0.12, px - 0.12, py - 0.05, color=pal.deep_navy)
    lo, hi = loc(theme, ends[0]), loc(theme, ends[1])
    for txt, x, al in ((lo, px - 0.5, PP_ALIGN.LEFT), (hi, px + pw - 1.0, PP_ALIGN.RIGHT)):
        tb = add_textbox(slide, x, py + ph + 0.18, 1.0, 0.28)
        write_paragraph(tb.text_frame, txt.upper(), size=10, bold=True, color=pal.footer_gray,
                        family=typo.family, align=al, first=True)
    tb = add_textbox(slide, px - 0.5, py - 0.3, 1.0, 0.26)
    write_paragraph(tb.text_frame, hi.upper(), size=10, bold=True, color=pal.footer_gray,
                    family=typo.family, first=True)
    tb = add_textbox(slide, px + pw / 2 - 1.5, py + ph + 0.18, 3.0, 0.3)
    write_paragraph(tb.text_frame, loc(theme, x_label), size=13, bold=True,
                    color=pal.deep_navy, family=typo.family, align=PP_ALIGN.CENTER, first=True)
    tb = add_textbox(slide, left - 1.2 + 0.25, py + ph / 2 - 0.16, 2.4, 0.32,
                     anchor=MSO_ANCHOR.MIDDLE)
    tb.rotation = 270
    write_paragraph(tb.text_frame, loc(theme, y_label), size=13, bold=True,
                    color=pal.deep_navy, family=typo.family, align=PP_ALIGN.CENTER, first=True)

    # risk boxes
    us = _norm_levels([r.get("probability", 0.5) for r in risks])
    vs = _norm_levels([r.get("impact", 0.5) for r in risks])
    bw = min(2.3, pw * 0.36)
    labels = [f"{i + 1}. {r.get('title', '')}" for i, r in enumerate(risks)]
    size = min([fit_size([t], bw - 0.2, 0.95, max_size=13, min_size=MIN_PT) for t in labels]
               or [12])
    warn_small("risk_heatmap", title, size, "Shorten risk titles.")
    placed = []
    for i, (u, v, t) in enumerate(zip(us, vs, labels)):
        bh = text_height_in([t], bw - 0.2, size) + 0.16
        bx = min(max(px + u * pw - bw / 2, px + 0.06), px + pw - bw - 0.06)
        by = min(max(py + (1 - v) * ph - bh / 2, py + 0.06), py + ph - bh - 0.06)
        for _ in range(12):  # nudge away from earlier boxes
            clash = next(((qx, qy, qw, qh) for qx, qy, qw, qh in placed
                          if bx < qx + qw + 0.05 and qx < bx + bw + 0.05
                          and by < qy + qh + 0.05 and qy < by + bh + 0.05), None)
            if not clash:
                break
            qx, qy, qw, qh = clash
            down = qy + qh + 0.08
            up = qy - bh - 0.08
            by = down if down + bh <= py + ph - 0.06 else max(up, py + 0.06)
            if by == py + 0.06:
                bx = min(bx + bw * 0.5, px + pw - bw - 0.06)
        placed.append((bx, by, bw, bh))
        add_rect(slide, bx, by, bw, bh, fill=pal.white, line=pal.deep_navy, line_width=1.25)
        tb = add_textbox(slide, bx + 0.1, by, bw - 0.2, bh, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, t, size=size, theme=theme, align=PP_ALIGN.CENTER,
                             first=True)

    if has_mit:
        mx = left + lw + gap
        mits = [r.get("mitigation") or "—" for r in risks]
        mits = [m if isinstance(m, str) else "; ".join(m) for m in mits]
        m = len(mits)
        g = 0.14
        avail = bottom - top
        msize = min(fit_size([t], rw - 0.75, (avail - g * (m - 1)) / m - 0.12, max_size=14,
                             min_size=MIN_PT) for t in mits)
        warn_small("risk_heatmap", title, msize, "Shorten the mitigations.")
        hs = [text_height_in([t], rw - 0.75, msize) + 0.24 for t in mits]
        extra = avail - sum(hs) - g * (m - 1)
        hs = [h + max(0, min(extra / m, 0.35)) for h in hs]
        y = top
        for i, (t, h) in enumerate(zip(mits, hs)):
            box = add_rect(slide, mx, y, rw, h, fill=pal.white)
            set_dashed(box, pal.deep_navy, 1.0)
            d = 0.34
            add_icon(slide, str(i + 1), mx + 0.14, y + (h - d) / 2, d, theme, "navy")
            tb = add_textbox(slide, mx + 0.6, y, rw - 0.75, h, anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, t, size=msize, theme=theme, first=True)
            y += h + g
    return slide


# ---------- positioning scale ----------

_PLAYER_TONES = ("navy", "blue", "amber", "gray", "green", "light_blue")


def add_positioning_scale(prs, *,
                          title: str = "[Positioning scale / Insert action title]",
                          players: Sequence[Union[str, Dict]],
                          indicators: Sequence[Dict],
                          focus: Optional[int] = 0,
                          scale: Optional[Sequence[float]] = None,
                          ends: Sequence[str] = ("Low", "High"),
                          headers: Sequence[str] = ("Indicator", "What it means", "Comparison"),
                          players_label: str = "Players",
                          subtitle: Optional[str] = None,
                          insight: Optional[str] = None,
                          insight_label: Optional[str] = "Key insight",
                          page_number=None, section_marker=None,
                          source=None, footnote=None,
                          theme: Theme = DEFAULT_THEME):
    """players: [str | {name, short?, tone?, keywords?: [str]}] (2-5);
    `short` is the 1-2 character marker label (default: first letter);
    keywords add a left panel with each player's profile.
    indicators: [{name, note?, ends?: (low, high), values: [num per player]}]
    scale: (min, max) of the values (default: min/max found, or 0-1).
    focus: index of the player to emphasise (its marker drawn largest).
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    pls = [p if isinstance(p, dict) else {"name": str(p)} for p in players]
    for i, p in enumerate(pls):
        p.setdefault("short", plain(p["name"])[:1].upper())
        p.setdefault("tone", _PLAYER_TONES[i % len(_PLAYER_TONES)])
    has_kw = any(p.get("keywords") for p in pls)
    pan_w = 2.45 if has_kw else 0.0
    pan_g = 0.3 if has_kw else 0.0
    x0 = left + pan_w + pan_g
    W = width - pan_w - pan_g
    head_h = 0.42
    num_w, name_w = 0.55, W * 0.2
    note_w = W * 0.3 if any(r.get("note") for r in indicators) else 0.0
    sc_x = x0 + num_w + name_w + note_w + 0.25
    sc_w = left + width - sc_x - 0.15
    add_stage_banners(slide, theme, [(x0, num_w + name_w, headers[0], "stage")]
                      + ([(x0 + num_w + name_w, note_w, headers[1], "stage")] if note_w else [])
                      + [(sc_x - 0.15, sc_w + 0.15, headers[2], "result")],
                      top, head_h, style="chevron", where=title)
    body_top = top + head_h + 0.15

    # legend (when no keyword panel)
    if not has_kw:
        lx = x0
        for p in pls:
            d = 0.26
            add_icon(slide, p["short"], lx, body_top + 0.02, d, theme, p["tone"])
            nw = text_width_pt(p["name"], 12, True) / 72 / 0.92 + 0.15
            tb = add_textbox(slide, lx + d + 0.08, body_top, nw, 0.3, anchor=MSO_ANCHOR.MIDDLE)
            write_paragraph(tb.text_frame, p["name"], size=12, bold=True, color=pal.text_dark,
                            family=typo.family, first=True)
            lx += d + 0.08 + nw + 0.3
        body_top += 0.45

    n = max(len(indicators), 1)
    g = 0.1
    rh = max(0.62, min(1.0, (bottom - body_top - g * (n - 1)) / n))
    vals = [v for r in indicators for v in r.get("values", [])]
    lo_v, hi_v = (scale if scale else ((0.0, 1.0) if vals and max(vals) <= 1 and min(vals) >= 0
                                       else (min(vals or [0]), max(vals or [1]))))
    span = (hi_v - lo_v) or 1
    names = [r.get("name", "") for r in indicators]
    nsize = min(fit_size([t], name_w - 0.1, rh - 0.1, max_size=14, min_size=MIN_PT, bold=True)
                for t in names)
    notes = [r.get("note", "") for r in indicators]
    tsize = min([fit_size([t], note_w - 0.15, rh - 0.12, max_size=13, min_size=MIN_PT)
                 for t in notes if t] or [12])
    warn_small("positioning_scale", title, min(nsize, tsize), "Shorten indicator notes.")
    for i, r in enumerate(indicators):
        y = body_top + i * (rh + g)
        add_rect(slide, x0, y, W, rh, fill=pal.soft_gray)
        d = min(0.4, rh - 0.2)
        add_icon(slide, str(i + 1), x0 + 0.08, y + (rh - d) / 2, d, theme, "mid_blue")
        tb = add_textbox(slide, x0 + num_w, y, name_w - 0.1, rh, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, r.get("name", ""), size=nsize, theme=theme,
                             bold=True, first=True)
        if note_w and r.get("note"):
            tb = add_textbox(slide, x0 + num_w + name_w, y, note_w - 0.15, rh,
                             anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, r["note"], size=tsize, theme=theme,
                                 color=pal.text_dark, first=True)
        ly = y + rh * 0.42
        add_arrow(slide, sc_x + 0.2, ly, sc_x - 0.02, ly, color=pal.footer_gray, width_pt=1.0)
        add_arrow(slide, sc_x + 0.2, ly, sc_x + sc_w, ly, color=pal.footer_gray, width_pt=1.0)
        e0, e1 = r.get("ends") or ends
        for txt, xx, al in ((loc(theme, e0), sc_x, PP_ALIGN.LEFT),
                            (loc(theme, e1), sc_x + sc_w - 1.2, PP_ALIGN.RIGHT)):
            tb = add_textbox(slide, xx, ly + 0.19, 1.2, 0.22)
            write_paragraph(tb.text_frame, txt, size=10, italic=True, color=pal.footer_gray,
                            family=typo.family, align=al, first=True)
        # markers: keep a minimum gap between neighbours, focus player on top
        vals_r = list(r.get("values", []))
        pts = []
        for k in range(len(pls)):
            v = vals_r[k] if k < len(vals_r) else None
            if v is None:
                continue
            md = 0.36 if k == focus else 0.3
            pts.append([sc_x + 0.25 + (float(v) - lo_v) / span * (sc_w - 0.5), md, k])
        pts.sort(key=lambda t: t[0])
        for a, b in zip(pts, pts[1:]):
            need = (a[1] + b[1]) / 2 + 0.04
            if b[0] - a[0] < need:
                b[0] = a[0] + need
        over = pts[-1][0] + pts[-1][1] / 2 - (sc_x + sc_w) if pts else 0
        if over > 0:
            for t in pts:
                t[0] -= over
        for cxm, md, k in sorted(pts, key=lambda t: t[2] == focus):
            add_icon(slide, pls[k]["short"], cxm - md / 2, ly - md / 2, md, theme,
                     pls[k]["tone"])

    if has_kw:
        ph = bottom - top
        add_stage_banners(slide, theme, [(left, pan_w, players_label, "result")], top, head_h,
                          style="bar", where=title)
        m = len(pls)
        g2 = 0.15
        ch = (ph - head_h - 0.15 - g2 * (m - 1)) / m
        for i, p in enumerate(pls):
            y = top + head_h + 0.15 + i * (ch + g2)
            add_rect(slide, left, y, pan_w, ch, fill=pal.white,
                     line=tone_rgb(theme, p["tone"]), line_width=1.5)
            d = 0.34
            add_icon(slide, p["short"], left + 0.12, y + 0.12, d, theme, p["tone"])
            tb = add_textbox(slide, left + 0.55, y + 0.1, pan_w - 0.65, 0.38,
                             anchor=MSO_ANCHOR.MIDDLE)
            write_paragraph(tb.text_frame, p["name"], size=13, bold=True, color=pal.text_dark,
                            family=typo.family, first=True)
            kws = p.get("keywords", [])
            if kws:
                ks = fit_size(kws, pan_w - 0.3, ch - 0.6, max_size=13, min_size=MIN_PT)
                tb = add_textbox(slide, left + 0.15, y + 0.52, pan_w - 0.3, ch - 0.6,
                                 anchor=MSO_ANCHOR.MIDDLE)
                for j, kw in enumerate(kws):
                    write_rich_paragraph(tb.text_frame, kw, size=ks, theme=theme,
                                         color=pal.deep_navy, bold=True,
                                         align=PP_ALIGN.CENTER, first=j == 0)
    return slide


# ---------- value chain ----------

def add_value_chain(prs, *,
                    title: str = "[Value chain / Insert action title]",
                    stages: Sequence[Dict],
                    groups: Sequence[Dict] = (),
                    own_label: Optional[str] = None,
                    body_label: Optional[str] = None,
                    subtitle: Optional[str] = None,
                    insight: Optional[str] = None,
                    insight_label: Optional[str] = "Key insight",
                    page_number=None, section_marker=None,
                    source=None, footnote=None,
                    theme: Theme = DEFAULT_THEME):
    """stages: [{name, value?: str (2nd line in the chevron, e.g. a price),
                 bullets?: [str] | body?: str, own?: bool}] left -> right (3-6).
    own: stages the company performs — dark chevron and outlined box;
         own_label adds a legend ("Activities of <company>").
    groups: [{label, start, end}] ruled headers spanning stages (indices,
         inclusive) between chevrons and boxes, e.g. Suppliers / Buyers.
    body_label: small caption above the boxes ("Key success factors").
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    n = max(len(stages), 1)
    sw = width / n
    ch = 0.78 if any(s.get("value") for s in stages) else 0.62
    cells = [{"body": s.get("body"), "bullets": s.get("bullets", [])} for s in stages]
    has_body = any(_cell(c) != ([], []) for c in cells)
    pre = ch + 0.2 + (0.46 if groups else 0) + (0.32 if body_label else 0)
    if has_body:
        inner0 = sw - 0.42
        size0 = 15
        for c in cells:
            for sz in range(size0, MIN_PT - 1, -1):
                if _cell_h(c, inner0, sz) <= bottom - top - pre - 0.3:
                    size0 = sz
                    break
            else:
                size0 = MIN_PT
        body_need = max(max(_cell_h(c, inner0, size0) for c in cells) + 0.36, 1.2)
    else:
        body_need = 0.0
    has_legend = bool(own_label and any(s_.get("own") for s_ in stages))
    spare = bottom - top - pre - body_need - (0.35 if has_legend else 0)
    if spare > 0.6:
        top += spare * 0.4
    if own_label and any(s.get("own") for s in stages):
        lw = text_width_pt(own_label, 11, False) / 72 / 0.92 + 0.6
        box = add_rect(slide, left + width - lw, top - 0.02, 0.4, 0.22, fill=None)
        set_dashed(box, pal.status_red, 1.25)
        tb = add_textbox(slide, left + width - lw + 0.5, top - 0.06, lw - 0.5, 0.3,
                         anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, own_label, size=11, color=pal.text_dark,
                        family=typo.family, first=True)
        top += 0.35
    names = [s.get("name", "") for s in stages]
    hsize = min(fit_size([t], sw - 0.55, ch / 2 if any(s.get("value") for s in stages)
                         else ch - 0.1, max_size=14, min_size=MIN_PT, bold=True,
                         line_spacing=1.05) for t in names)
    hsize = min([hsize] + [fit_one_line(max(plain(t).split() or [""], key=len), sw - 0.55, 14,
                                        MIN_PT, True) for t in names])
    for i, s in enumerate(stages):
        x = left + i * sw
        shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        shp = slide.shapes.add_shape(shape, Inches(x), Inches(top), Inches(sw + (0.14 if i < n - 1 else 0)),
                                     Inches(ch))
        shp.adjustments[0] = 0.28
        shp.shadow.inherit = False
        own = bool(s.get("own"))
        shp.fill.solid()
        shp.fill.fore_color.rgb = pal.deep_navy if own else tint(pal.mid_blue,
                                                                 0.35 + 0.45 * i / max(n - 1, 1))
        if own:
            shp.line.color.rgb = pal.status_red
            shp.line.width = Pt(1.75)
        else:
            shp.line.fill.background()
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3 if i else 0.12)
        tf.margin_right = Inches(0.2)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        dark = own or i < n / 2
        col = pal.white if dark else pal.deep_navy
        write_paragraph(tf, s.get("name", ""), size=hsize, bold=True, color=col,
                        family=typo.family, align=PP_ALIGN.CENTER, first=True)
        if s.get("value"):
            write_paragraph(tf, s["value"], size=max(hsize - 1, MIN_PT), bold=False, color=col,
                            family=typo.family, align=PP_ALIGN.CENTER)
    y = top + ch + 0.2
    if groups:
        spans = [(left + g["start"] * sw + 0.1, (g["end"] - g["start"] + 1) * sw - 0.2,
                  g["label"], "stage") for g in groups]
        add_stage_banners(slide, theme, spans, y, 0.34, style="rule", where=title)
        y += 0.46
    if body_label:
        tb = add_textbox(slide, left, y, width, 0.28)
        p = write_paragraph(tb.text_frame, loc(theme, body_label).upper(), size=10, bold=True,
                            color=pal.mid_blue, family=typo.family, first=True)
        letter_space(p, 120)
        y += 0.32
    if has_body:
        bh = bottom - y
        inner = sw - 0.42
        size = 15
        for c in cells:
            for sz in range(size, MIN_PT - 1, -1):
                if _cell_h(c, inner, sz) <= bh - 0.3:
                    size = sz
                    break
            else:
                size = MIN_PT
        warn_small("value_chain", title, size, "Shorten the bullets per stage.")
        need = max(_cell_h(c, inner, size) for c in cells) + 0.36
        bh = min(bh, max(need, 1.2))
        for i, (s, c) in enumerate(zip(stages, cells)):
            x = left + i * sw + 0.06
            own = bool(s.get("own"))
            box = add_rect(slide, x, y, sw - 0.12, bh, fill=tint(pal.light_blue, 0.9) if own
                           else pal.white)
            if own:
                box.line.color.rgb = pal.status_red
                box.line.width = Pt(1.5)
            else:
                set_dashed(box, pal.footer_gray, 1.0)
            if _cell(c) != ([], []):
                _write_cell(slide, theme, x + 0.15, y + 0.15, inner, bh - 0.3, c, size,
                            anchor=MSO_ANCHOR.TOP)
    return slide
