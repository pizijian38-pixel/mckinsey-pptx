"""Planning templates: roadmap (workstreams x periods) and timeline (dated events)."""
from __future__ import annotations

from typing import Dict, Optional, Sequence

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from ..base import add_chrome, add_line, add_oval, add_rect, add_textbox, blank_slide, write_paragraph
from ..design import (add_callout_bar, fit_one_line, fit_size, text_height_in, text_width_pt, tone_rgb,
                      warn_small, write_rich_paragraph)
from ..theme import Theme, DEFAULT_THEME


def _frame(prs, theme, title, subtitle, insight, insight_label, **chrome):
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, **chrome)
    pal, typo, layout = theme.palette, theme.typography, theme.layout
    left = layout.margin_left_in
    width = layout.slide_width_in - layout.margin_left_in - layout.margin_right_in
    top = layout.body_top_in + 0.05
    bottom = layout.footer_top_in - 0.25
    if subtitle:
        tb = add_textbox(slide, left, top, width, 0.35)
        write_paragraph(tb.text_frame, subtitle, size=typo.section_title_size,
                        bold=True, color=pal.text_dark, family=typo.family, first=True)
        top += 0.45
    if insight:
        h = 0.7
        add_callout_bar(slide, theme, left, bottom - h, width, h, insight,
                        label=insight_label)
        bottom -= h + 0.25
    return slide, left, top, width, bottom


def _diamond(slide, cx, cy, d, rgb):
    s = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(cx - d / 2), Inches(cy - d / 2),
                               Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def add_roadmap(prs, *,
                title: str = "[Roadmap / Insert action title]",
                periods: Sequence[str],
                lanes: Sequence[Dict],
                milestones: Sequence[Dict] = (),
                subtitle: Optional[str] = None,
                insight: Optional[str] = None,
                insight_label: Optional[str] = "Key insight",
                page_number=None, section_marker=None,
                source=None, footnote=None,
                theme: Theme = DEFAULT_THEME):
    """periods: column labels, e.g. ["Q1", "Q2", "Q3", "Q4"] or months.
    lanes: [{"name": str, "items": [{"label": str, "start": num, "end": num,
             "tone"?: str}]}] — start/end are 0-based period indices and the
             bar runs from the start of period `start` to the end of period
             `end` (start=1, end=2 -> Q2 to Q3). Add .5 to begin mid-period.
    milestones: [{"label": str, "at": num, "tone"?: str}] — `at` is a position
             on the same axis: 0 = start of the first period, 4 = end of the
             fourth (at=2 -> the boundary between period 2 and 3).
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    lane_w = 2.3
    gx, gw = left + lane_w, width - lane_w
    n_p = max(len(periods), 1)
    pw = gw / n_p
    head_h = 0.42
    ms_h = 0.85 if milestones else 0

    # sub-rows per lane so overlapping bars don't collide
    lane_rows = []
    for ln in lanes:
        rows_end = []
        placed = []
        for it in sorted(ln.get("items", []), key=lambda t: t["start"]):
            # a bar occupies [start, end + 1) in period units
            for ri, end in enumerate(rows_end):
                if it["start"] >= end:
                    rows_end[ri] = it["end"] + 1
                    placed.append((ri, it))
                    break
            else:
                rows_end.append(it["end"] + 1)
                placed.append((len(rows_end) - 1, it))
        lane_rows.append((max(len(rows_end), 1), placed))
    total_sub = sum(n for n, _ in lane_rows) or 1
    avail = bottom - top - head_h - ms_h
    sub_h = min(0.85, avail / total_sub)
    bar_h = sub_h * 0.72

    # header
    for i, p in enumerate(periods):
        x = gx + i * pw
        add_rect(slide, x + 0.02, top, pw - 0.04, head_h, fill=pal.deep_navy)
        tb = add_textbox(slide, x, top, pw, head_h, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, p, size=min(14, fit_one_line(p, pw - 0.1, 14, 9, True)),
                        bold=True, color=pal.white, family=typo.family,
                        align=PP_ALIGN.CENTER, first=True)
    y = top + head_h

    def px(at):
        return gx + (at + 0.0) * pw

    if milestones:
        line_y = y + 0.2
        add_line(slide, gx, line_y, gx + gw, line_y, color=pal.rule_gray, width_pt=0.75)
        row_right = [-1e9, -1e9]          # two label rows; drop to row 2 on collision
        for m in sorted(milestones, key=lambda m: m["at"]):
            cx = px(min(max(m["at"], 0), n_p))
            _diamond(slide, cx, line_y, 0.22, tone_rgb(theme, m.get("tone", "amber")))
            size = 11
            text_w = text_width_pt(m["label"], size, True) / 72 + 0.15
            lx = min(max(cx - text_w / 2, gx), gx + gw - text_w)
            row = 0 if lx > row_right[0] + 0.05 else 1
            row_right[row] = lx + text_w
            tb = add_textbox(slide, lx, line_y + 0.13 + row * 0.28, text_w, 0.26)
            write_paragraph(tb.text_frame, m["label"], size=size, bold=True,
                            color=pal.text_dark, family=typo.family,
                            align=PP_ALIGN.CENTER, first=True)
        y += ms_h

    grid_top = y
    grid_bottom = y + sum(n for n, _ in lane_rows) * sub_h
    for i in range(1, n_p):          # behind the bars
        add_line(slide, gx + i * pw, grid_top, gx + i * pw, grid_bottom,
                 color=pal.grid_gray, width_pt=0.5)
    sizes = []
    for (n_sub, placed), ln in zip(lane_rows, lanes):
        lane_h = n_sub * sub_h
        add_rect(slide, left, y + 0.03, lane_w - 0.1, lane_h - 0.06, fill=pal.light_gray)
        tb = add_textbox(slide, left + 0.15, y, lane_w - 0.35, lane_h, anchor=MSO_ANCHOR.MIDDLE)
        ls = fit_size([ln["name"]], lane_w - 0.35, lane_h - 0.06, max_size=14, min_size=10)
        write_rich_paragraph(tb.text_frame, ln["name"], size=ls, theme=theme,
                             bold=True, first=True)
        for ri, it in placed:
            x0 = px(max(0, it["start"])) + 0.04
            x1 = px(min(n_p, it["end"] + 1)) - 0.04
            by = y + ri * sub_h + (sub_h - bar_h) / 2
            rgb = tone_rgb(theme, it.get("tone", "blue"))
            add_rect(slide, x0, by, x1 - x0, bar_h, fill=rgb)
            s = fit_one_line(it["label"], x1 - x0 - 0.16, 14 if bar_h > 0.5 else 13, 9, bold=True)
            sizes.append(s)
            tb = add_textbox(slide, x0 + 0.08, by, x1 - x0 - 0.16, bar_h,
                             anchor=MSO_ANCHOR.MIDDLE)
            write_paragraph(tb.text_frame, it["label"], size=s, bold=True,
                            color=pal.text_dark if it.get("tone") in ("light_blue", "amber")
                            else pal.white, family=typo.family, first=True)
        y += lane_h
        add_line(slide, left, y, left + width, y, color=pal.grid_gray, width_pt=0.5)
    if sizes:
        warn_small("roadmap", title, min(sizes),
                   "Shorten bar labels or give the item a longer span.")
    return slide


def add_timeline(prs, *,
                 title: str = "[Timeline / Insert action title]",
                 events: Sequence[Dict],
                 subtitle: Optional[str] = None,
                 insight: Optional[str] = None,
                 insight_label: Optional[str] = "Key insight",
                 page_number=None, section_marker=None,
                 source=None, footnote=None,
                 theme: Theme = DEFAULT_THEME):
    """events (chronological, 3-7): [{"date": str, "title": str, "body"?: str,
    "tone"?: str}] on a horizontal line; text alternates above and below."""
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    n = max(len(events), 1)
    mid = (top + bottom) / 2
    add_line(slide, left, mid, left + width, mid, color=pal.deep_navy, width_pt=2.5)
    slot = width / n
    box_w = min(slot * 1.7, 3.6)
    box_h = (bottom - top) / 2 - 0.35
    paras_all = [[e.get("title", "")] + ([e["body"]] if e.get("body") else []) for e in events]
    size = min(fit_size(p, box_w, box_h - 0.4, max_size=18, min_size=10, para_gap_pt=4)
               for p in paras_all) if events else 13
    warn_small("timeline", title, size, "Shorten event text or use fewer events.")
    for i, e in enumerate(events):
        cx = left + slot * (i + 0.5)
        tone = e.get("tone", "blue")
        add_oval(slide, cx - 0.14, mid - 0.14, 0.28, 0.28, fill=tone_rgb(theme, tone))
        above = i % 2 == 0
        add_line(slide, cx, mid + (-0.14 if above else 0.14), cx,
                 mid + (-0.32 if above else 0.32), color=pal.rule_gray, width_pt=0.75)
        bx = min(max(cx - box_w / 2, left), left + width - box_w)
        by = mid - 0.35 - box_h if above else mid + 0.35
        tb = add_textbox(slide, bx, by, box_w, box_h,
                         anchor=MSO_ANCHOR.BOTTOM if above else MSO_ANCHOR.TOP)
        write_paragraph(tb.text_frame, e.get("date", ""), size=size, bold=True,
                        color=tone_rgb(theme, tone), family=typo.family,
                        align=PP_ALIGN.CENTER, first=True)
        write_rich_paragraph(tb.text_frame, e.get("title", ""), size=size, theme=theme,
                             bold=True, align=PP_ALIGN.CENTER, space_before=2)
        if e.get("body"):
            write_rich_paragraph(tb.text_frame, e["body"], size=size - 1, theme=theme,
                                 align=PP_ALIGN.CENTER, space_before=2)
    return slide


# ---------- phase grid ----------

def add_phase_grid(prs, *,
                   title: str = "[Phase grid / Insert action title]",
                   stages: Sequence[str],
                   rows: Sequence[Dict],
                   side: Optional[Dict] = None,
                   timeline: bool = True,
                   subtitle: Optional[str] = None,
                   insight: Optional[str] = None,
                   insight_label: Optional[str] = "Key insight",
                   page_number=None, section_marker=None,
                   source=None, footnote=None,
                   theme: Theme = DEFAULT_THEME):
    """Implementation plan as a grid: row labels (channels, partners, goals,
    resources ...) x stages (Stage 1 / Stage 2, Short / Mid / Long term, years).
    stages: header per column, left -> right.
    rows: [{"label": str, "cells": [cell per stage], "kpi"?: str}]
          cell = None | str | [bullets] | {"body"?, "bullets"?, "span"?: int, "tone"?}
          span: the cell covers that many stages (a bar across years);
          kpi: a full-width band under the row ("KPI: 2M revenue, 300% CAGR").
    side: {"title": str, "items": [str | {"title", "bullets"? | "body"?}]} —
          a panel on the right (risk mitigation, enablers), split by a dashed line.
    """
    from ..design import add_stage_banners, set_dashed, tint
    from ..labels import loc
    from .logic_slides import _cell, _cell_h, _write_cell
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    sw_ = width * 0.29 if side else 0.0
    sg = 0.45 if side else 0.0
    gw = width - sw_ - sg
    lw = 1.55
    n = max(len(stages), 1)
    cg = 0.12
    cw = (gw - lw - 0.12 - cg * (n - 1)) / n
    xs = [left + lw + 0.12 + i * (cw + cg) for i in range(n)]

    # headers (+ timeline)
    head_h = 0.38
    hy = top + (0.12 if timeline else 0)
    if timeline:
        ly = hy + head_h + 0.12
        add_line(slide, xs[0], ly, xs[-1] + cw, ly, color=pal.grid_gray, width_pt=2.5)
        for x in xs:
            add_oval(slide, x + cw / 2 - 0.07, ly - 0.07, 0.14, 0.14, fill=pal.white,
                     line=pal.mid_blue, line_width=1.5)
    hs = min(fit_size([t], cw - 0.3, head_h - 0.04, max_size=14, min_size=10, bold=True)
             for t in stages)
    for x, t in zip(xs, stages):
        tw = min(cw, text_width_pt(t, hs, True) / 72 / 0.9 + 0.5)
        add_rect(slide, x + (cw - tw) / 2, hy, tw, head_h, fill=pal.deep_navy)
        tb = add_textbox(slide, x + (cw - tw) / 2, hy, tw, head_h, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, t, size=hs, bold=True, color=pal.white,
                        family=typo.family, align=PP_ALIGN.CENTER, first=True)
    body_top = hy + head_h + (0.3 if timeline else 0.15)

    # lay out cells per row with spans
    def layout_row(r):
        out, col = [], 0
        for c in r.get("cells", []):
            span = int(c.get("span", 1)) if isinstance(c, dict) else 1
            if col >= n:
                break
            span = max(1, min(span, n - col))
            out.append((col, span, c))
            col += span
        return out

    laid = [layout_row(r) for r in rows]
    kpi_h = 0.4
    g = 0.14
    m = max(len(rows), 1)

    def width_of(span):
        return span * cw + (span - 1) * cg

    def need(i, s):
        hh = [_cell_h(c, width_of(sp) - 0.3, s) + 0.24 for _, sp, c in laid[i] if c]
        lab = text_height_in([rows[i].get("label", "")], lw - 0.2, s + 1, bold=True) + 0.2
        return max(hh + [lab, 0.5])

    n_kpi = sum(1 for r in rows if r.get("kpi"))
    avail = bottom - body_top - g * (m - 1) - n_kpi * (kpi_h + 0.06)
    size = 10
    for s in range(15, 9, -1):
        if sum(need(i, s) for i in range(len(rows))) <= avail:
            size = s
            break
    warn_small("phase_grid", title, size, "Shorten the cells or split the plan over two slides.")
    needs = [need(i, size) for i in range(len(rows))]
    extra = avail - sum(needs)
    add = min(extra / m, 0.5) if extra > 0 else 0.0
    y = body_top
    lsize = min(size + 1, min(fit_one_line(max(str(r.get("label", "")).split() or [""], key=len),
                                           lw - 0.2, 15, 10, True) for r in rows))
    for i, r in enumerate(rows):
        rh = needs[i] + add
        add_rect(slide, left, y, lw, rh, fill=pal.light_gray)
        tb = add_textbox(slide, left + 0.1, y, lw - 0.2, rh, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, r.get("label", ""), size=lsize, theme=theme,
                             bold=True, align=PP_ALIGN.CENTER, first=True)
        for col, sp, c in laid[i]:
            if not c:
                continue
            x, w = xs[col], width_of(sp)
            tone = c.get("tone") if isinstance(c, dict) else None
            solid = tone in ("navy", "blue", "mid_blue")
            if solid:
                add_rect(slide, x, y, w, rh, fill=tone_rgb(theme, tone))
            elif tone:
                add_rect(slide, x, y, w, rh, fill=tint(tone_rgb(theme, tone), 0.82))
            else:
                add_rect(slide, x, y, w, rh, fill=pal.white, line=pal.grid_gray, line_width=1.0)
            body, bullets = _cell(c)
            _write_cell(slide, theme, x + 0.15, y + 0.1, w - 0.3, rh - 0.2, c, size,
                        color=pal.white if solid else None, bold=solid,
                        align=PP_ALIGN.CENTER if (sp > 1 and not bullets) else PP_ALIGN.LEFT)
        y += rh
        if r.get("kpi"):
            y += 0.06
            add_rect(slide, left, y, lw, kpi_h, fill=pal.status_red)
            tb = add_textbox(slide, left, y, lw, kpi_h, anchor=MSO_ANCHOR.MIDDLE)
            write_paragraph(tb.text_frame, loc(theme, "KPI"), size=13, bold=True, italic=True,
                            color=pal.white, family=typo.family, align=PP_ALIGN.CENTER,
                            first=True)
            kx = xs[0]
            add_rect(slide, kx, y, xs[-1] + cw - kx, kpi_h, fill=tint(pal.status_red, 0.15))
            ks = fit_size([r["kpi"]], xs[-1] + cw - kx - 0.3, kpi_h - 0.04, max_size=13,
                          min_size=10)
            tb = add_textbox(slide, kx + 0.15, y, xs[-1] + cw - kx - 0.3, kpi_h,
                             anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, r["kpi"], size=ks, theme=theme, color=pal.white,
                                 bold=True, align=PP_ALIGN.CENTER, first=True)
            y += kpi_h
        y += g

    if side:
        sx = left + gw + sg
        ln = add_line(slide, sx - sg / 2, top, sx - sg / 2, bottom, color=pal.footer_gray,
                      width_pt=1.0)
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        st = loc(theme, side.get("title", "Risks and mitigation"))
        add_rect(slide, sx + 0.2, top, sw_ - 0.4, 0.38, fill=pal.deep_navy)
        tb = add_textbox(slide, sx + 0.2, top, sw_ - 0.4, 0.38, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, st, size=13, bold=True, color=pal.white,
                        family=typo.family, align=PP_ALIGN.CENTER, first=True)
        items = [it if isinstance(it, dict) else {"title": str(it)} for it in side.get("items", [])]
        paras = []
        for it in items:
            paras.append(it.get("title", ""))
            paras += ([it["body"]] if it.get("body") else []) + list(it.get("bullets", []))
        ss = fit_size(paras or [" "], sw_ - 0.3, bottom - top - 0.55, max_size=14, min_size=10,
                      para_gap_pt=5, indent_in=0.22)
        warn_small("phase_grid", title, ss, "Shorten the side panel.")
        tb = add_textbox(slide, sx, top + 0.55, sw_, bottom - top - 0.55)
        first = True
        for it in items:
            if it.get("title"):
                p = write_rich_paragraph(tb.text_frame, "✓ " + it["title"], size=ss, theme=theme,
                                         color=pal.deep_navy, bold=True, first=first,
                                         space_before=None if first else 8)
                first = False
            for t in ([it["body"]] if it.get("body") else []) + list(it.get("bullets", [])):
                write_rich_paragraph(tb.text_frame, t, size=ss, theme=theme, bullet=True,
                                     first=first, space_before=None if first else 4)
                first = False
    return slide
