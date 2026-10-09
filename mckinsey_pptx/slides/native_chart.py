"""Native PowerPoint chart + key-insight panel.

The chart is a real PowerPoint chart (data editable via "Edit Data"), unlike
the shape-drawn column/line templates. Use it whenever the user may want to
update the numbers later.
"""
from __future__ import annotations

from typing import Dict, Optional, Sequence

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.util import Inches, Pt

from ..base import add_chrome, add_textbox, blank_slide, write_paragraph
from ..design import add_insight_panel, tone_rgb
from ..theme import Theme, DEFAULT_THEME

_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "grouped_column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_column_100": XL_CHART_TYPE.COLUMN_STACKED_100,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}
_DEFAULT_SERIES_TONES = ("navy", "blue", "mid_blue", "light_blue", "amber", "gray")


def draw_chart(slide, theme: Theme, x, y, w, h, *, chart_type: str = "column",
               categories: Sequence, series: Sequence[Dict],
               number_format: str = "General", show_values: bool = True,
               highlight: Optional[Dict] = None, y_max: Optional[float] = None):
    """Native chart inside a box (see add_native_chart for the arguments)."""
    pal, typo = theme.palette, theme.typography
    data = CategoryChartData()
    data.categories = [str(c) for c in categories]
    for s in series:
        data.add_series(s["name"], [None if v is None else float(v) for v in s["values"]])
    xl_type = _TYPES[chart_type]
    gf = slide.shapes.add_chart(xl_type, Inches(x), Inches(y),
                                Inches(w), Inches(h), data)
    chart = gf.chart
    chart.has_title = False          # the slide title / subtitle carry it
    chart.font.size = Pt(12)
    chart.font.name = typo.family
    chart.font.color.rgb = pal.text_dark

    is_pie = chart_type in ("pie", "doughnut")
    hl = highlight or {}
    plot = chart.plots[0]
    if not is_pie:
        if chart_type.startswith(("column", "grouped", "stacked", "bar")):
            plot.gap_width = 60 if len(series) > 1 else 80
            if chart_type.startswith("stacked"):
                plot.overlap = 100
        va = chart.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = pal.grid_gray
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(11)
        va.tick_labels.font.color.rgb = pal.footer_gray
        if y_max is not None:
            va.maximum_scale = y_max
        va.minimum_scale = 0 if chart_type != "line" else None
        ca = chart.category_axis
        ca.tick_labels.font.size = Pt(12)
        ca.format.line.color.rgb = pal.rule_gray

    multi = len(series) > 1 or is_pie
    chart.has_legend = multi
    if multi:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)

    hl_series = hl.get("series")
    # Emphasis is the accent (bright blue). Red means "bad" - only when asked for.
    hl_tone = hl.get("tone") or "blue"
    for i, (s, ps) in enumerate(zip(series, plot.series)):
        tone = s.get("tone") or _DEFAULT_SERIES_TONES[i % len(_DEFAULT_SERIES_TONES)]
        if hl_series is not None:
            tone = hl_tone if s["name"] == hl_series else ("gray" if not s.get("tone") else tone)
        rgb = tone_rgb(theme, tone)
        if chart_type == "line":
            ps.format.line.color.rgb = rgb
            ps.format.line.width = Pt(3.5 if s["name"] == hl_series else 2.25)
            ps.smooth = False
            ps.marker.format.fill.solid()
            ps.marker.format.fill.fore_color.rgb = rgb
            ps.marker.format.line.color.rgb = rgb
        elif not is_pie:
            ps.format.fill.solid()
            ps.format.fill.fore_color.rgb = rgb
        if is_pie:
            for pi, pt in enumerate(ps.points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = tone_rgb(
                    theme, _DEFAULT_SERIES_TONES[pi % len(_DEFAULT_SERIES_TONES)])
        if hl.get("point") is not None and not is_pie and chart_type != "line":
            pt = ps.points[int(hl["point"])]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = tone_rgb(theme, hl_tone)

    # Crowded line charts: label only the highlighted / explicitly toned series.
    label_only = None
    if chart_type == "line" and len(series) > 2:
        label_only = {s["name"] for s in series
                      if s["name"] == hl_series or s.get("tone")}
    if show_values and label_only:
        for s, ps in zip(series, plot.series):
            if s["name"] not in label_only:
                continue
            dl = ps.data_labels
            dl.number_format = number_format
            dl.number_format_is_linked = False
            dl.font.size = Pt(11)
            dl.font.bold = True
            dl.font.color.rgb = tone_rgb(theme, hl_tone if s["name"] == hl_series
                                         else s.get("tone"))
            dl.position = XL_LABEL_POSITION.ABOVE
            dl.show_value = True
    elif show_values:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = number_format
        dl.number_format_is_linked = False
        dl.font.size = Pt(11)
        if chart_type.startswith("stacked"):
            dl.position = XL_LABEL_POSITION.CENTER
            dl.font.color.rgb = pal.white
        elif chart_type == "line":
            dl.position = XL_LABEL_POSITION.ABOVE
        elif is_pie:
            dl.show_value = True
            dl.show_percentage = False
            dl.font.bold = True
            dl.font.size = Pt(13)
            dl.font.color.rgb = pal.white
        else:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END

    return gf


def add_native_chart(prs, *,
                     title: str = "[Chart / Insert action title]",
                     chart_type: str = "column",
                     categories: Sequence,
                     series: Sequence[Dict],
                     subtitle: Optional[str] = None,
                     number_format: str = "General",
                     show_values: bool = True,
                     highlight: Optional[Dict] = None,
                     insight: Optional[str] = None,
                     insight_bullets: Sequence[str] = (),
                     insight_title: str = "Key insight",
                     y_max: Optional[float] = None,
                     page_number=None, section_marker=None,
                     source=None, footnote=None,
                     theme: Theme = DEFAULT_THEME):
    """chart_type: column | grouped_column | stacked_column | stacked_column_100
                   | bar | stacked_bar | line | pie | doughnut
    series: [{"name": str, "values": [num], "tone"?: str}]
    highlight: {"series": name} -> that series in the accent, others muted (line/bar);
               {"point": index} -> that category's bar in the accent (single series).
               Add "tone": "red" (or green / amber) only when the item is bad (good)
               news - red is a judgement, not emphasis.
    number_format: Excel format for value labels, e.g. '0.0', '0"%"', '#,##0'.
    """
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo, layout = theme.palette, theme.typography, theme.layout
    left = layout.margin_left_in
    width = layout.slide_width_in - layout.margin_left_in - layout.margin_right_in
    top = layout.body_top_in + 0.05
    bottom = layout.footer_top_in - 0.25

    if subtitle:
        tb = add_textbox(slide, left, top, width, 0.35)
        write_paragraph(tb.text_frame, subtitle, size=typo.section_title_size,
                        bold=True, color=pal.text_dark, family=typo.family,
                        first=True)
        top += 0.45

    has_panel = bool(insight or insight_bullets)
    panel_w = 3.9 if has_panel else 0
    gap = 0.35 if has_panel else 0
    chart_w = width - panel_w - gap

    draw_chart(slide, theme, left, top, chart_w, bottom - top,
               chart_type=chart_type, categories=categories, series=series,
               number_format=number_format, show_values=show_values,
               highlight=highlight, y_max=y_max)

    if has_panel:
        add_insight_panel(slide, theme, left + chart_w + gap, top, panel_w,
                          bottom - top, title=insight_title, text=insight,
                          bullets=insight_bullets)
    return slide
