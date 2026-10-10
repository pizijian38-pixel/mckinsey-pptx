"""2x2 matrix: four labelled quadrants with items, or plotted points."""
from __future__ import annotations

from typing import Dict, Optional, Sequence

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from ..base import add_chrome, add_line, add_oval, add_rect, add_textbox, blank_slide, write_paragraph
from ..design import (HAIRLINE_PT, add_callout_bar, check_focus, eyebrow, fit_size, has_focus,
                      is_focus, muted_fill, tint, tone_rgb, warn_small, write_rich_paragraph)
from ..theme import Theme, DEFAULT_THEME


def add_matrix_2x2(prs, *,
                   title: str = "[2x2 matrix / Insert action title]",
                   x_label: str,
                   y_label: str,
                   quadrants: Sequence[Dict] = (),
                   points: Sequence[Dict] = (),
                   x_ends: Sequence[str] = ("Low", "High"),
                   y_ends: Sequence[str] = ("Low", "High"),
                   highlight: Optional[int] = None,
                   focus=None,
                   subtitle: Optional[str] = None,
                   insight: Optional[str] = None,
                   insight_label: Optional[str] = "Key insight",
                   page_number=None, section_marker=None,
                   source=None, footnote=None,
                   theme: Theme = DEFAULT_THEME):
    """quadrants (order: top-left, top-right, bottom-left, bottom-right):
         [{"title": str, "items"?: [str], "tone"?: str}]
    points: [{"label": str, "x": 0-1, "y": 0-1, "tone"?: str}] plotted on the grid
    focus: what the title is about — a quadrant (index 0-3 or its title) and/or
           point labels. The focal quadrant is tinted and outlined in the accent;
           focal points are drawn in the accent, the other points muted.
    highlight: older name for a focal quadrant index (still accepted).
    Axes: x_label left→right (x_ends), y_label bottom→top (y_ends).
    """
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo, layout = theme.palette, theme.typography, theme.layout
    left = layout.margin_left_in
    width = layout.slide_width_in - layout.margin_left_in - layout.margin_right_in
    top = layout.body_top_in + 0.05
    bottom = layout.footer_top_in - 0.25
    if subtitle:
        tb = add_textbox(slide, left, top, width, 0.35)
        write_paragraph(tb.text_frame, subtitle, size=typo.section_title_size,
                        bold=True, color=pal.text_dark, family=typo.family, first=True)
        top += 0.45
    if insight:
        h = 0.7
        add_callout_bar(slide, theme, left, bottom - h, width, h, insight,
                        label=insight_label)
        bottom -= h + 0.25

    axis_w, axis_h = 0.55, 0.45
    gx, gy = left + axis_w, top
    gw, gh = width - axis_w, bottom - top - axis_h
    qw, qh = (gw - 0.12) / 2, (gh - 0.12) / 2
    origins = [(gx, gy), (gx + qw + 0.12, gy), (gx, gy + qh + 0.12),
               (gx + qw + 0.12, gy + qh + 0.12)]

    if focus is None and highlight is not None:
        focus = highlight
    q_titles = [q.get("title", "") for q in quadrants]
    p_labels = [p.get("label", "") for p in points]
    bad = [k for k in ((focus,) if isinstance(focus, (str, int)) else (focus or ()))
           if not (isinstance(k, int) and 0 <= k < 4)
           and not any(is_focus(k, i, t) for i, t in enumerate(q_titles + p_labels))]
    if bad:
        check_focus("matrix_2x2", title, bad, [])
    focal_q = [i for i in range(4)
               if is_focus(focus, i, q_titles[i] if i < len(q_titles) else None)]
    point_focus = [k for k in ((focus,) if isinstance(focus, (str, int)) else (focus or ()))
                   if isinstance(k, str)]

    q_paras = [[q.get("title", "")] + list(q.get("items", [])) for q in quadrants]
    size = min((fit_size(p, qw - 0.5, qh - 0.4, max_size=16, min_size=10, para_gap_pt=4,
                         indent_in=0.25) for p in q_paras), default=14)
    warn_small("matrix_2x2", title, size, "Shorten quadrant items.")
    for i, (x, y) in enumerate(origins):
        f = i in focal_q
        if f:
            add_rect(slide, x, y, qw, qh, fill=tint(pal.bright_blue, 0.88),
                     line=pal.bright_blue, line_width=1.5)
        else:
            add_rect(slide, x, y, qw, qh, fill=pal.soft_gray)
        if i < len(quadrants):
            q = quadrants[i]
            tone = q.get("tone")
            tb = add_textbox(slide, x + 0.25, y + 0.2, qw - 0.5, qh - 0.4)
            write_rich_paragraph(tb.text_frame, q.get("title", ""), size=size + 2,
                                 theme=theme, bold=True,
                                 color=tone_rgb(theme, tone) if tone else
                                 (pal.mid_blue if f else pal.text_dark),
                                 first=True, space_after=4)
            for it in q.get("items", []):
                write_rich_paragraph(tb.text_frame, it, size=size, theme=theme,
                                     bullet=True, space_before=3)
    # points: the focal ones in the accent, the rest muted (when any is focal)
    pts_focused = any(is_focus(point_focus, i, l) for i, l in enumerate(p_labels))
    for i, p in enumerate(points):
        f = is_focus(point_focus, i, p.get("label"))
        cx = gx + p["x"] * gw
        cy = gy + (1 - p["y"]) * gh
        d = 0.3 if f else 0.26
        fill = (pal.bright_blue if f else muted_fill(theme)) if pts_focused \
            else tone_rgb(theme, p.get("tone", "blue"))
        add_oval(slide, cx - d / 2, cy - d / 2, d, d, fill=fill, line=pal.white, line_width=1.0)
        tb = add_textbox(slide, cx + d / 2 + 0.05, cy - 0.15, 2.2, 0.3, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, p["label"], size=12, bold=f or not pts_focused,
                        color=pal.text_dark if (f or not pts_focused) else pal.footer_gray,
                        family=typo.family, first=True)
    # axes
    add_line(slide, gx, gy + gh + 0.08, gx + gw, gy + gh + 0.08, color=pal.text_dark,
             width_pt=HAIRLINE_PT)
    add_line(slide, gx - 0.08, gy, gx - 0.08, gy + gh, color=pal.text_dark, width_pt=HAIRLINE_PT)
    from ..labels import loc
    x_ends = [loc(theme, t) for t in x_ends]
    y_ends = [loc(theme, t) for t in y_ends]
    for txt, x, al in ((x_ends[0], gx, PP_ALIGN.LEFT), (x_ends[1], gx + gw - 2, PP_ALIGN.RIGHT)):
        eyebrow(slide, theme, x, gy + gh + 0.14, 2, txt, align=al)
    tb = add_textbox(slide, gx + gw / 2 - 2.5, gy + gh + 0.12, 5, 0.3)
    write_paragraph(tb.text_frame, x_label + "  →", size=12, bold=True, color=pal.text_dark,
                    family=typo.family, align=PP_ALIGN.CENTER, first=True)
    # vertical y-axis label: a wide box rotated 270° around its centre
    cx, cy = left + (axis_w - 0.1) / 2, gy + gh / 2
    bw, bh = min(5.0, gh), 0.35
    tb = add_textbox(slide, cx - bw / 2, cy - bh / 2, bw, bh, anchor=MSO_ANCHOR.MIDDLE)
    tb.rotation = 270
    write_paragraph(tb.text_frame, y_label + "  →", size=12, bold=True,
                    color=pal.text_dark, family=typo.family, align=PP_ALIGN.CENTER,
                    first=True)
    for txt, y in ((y_ends[1], gy), (y_ends[0], gy + gh - 0.3)):
        eyebrow(slide, theme, left - 0.1, y, axis_w, txt, align=PP_ALIGN.CENTER, size=9)
    return slide
