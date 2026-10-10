"""Performance templates used across business decks.

- waterfall: bridge from a start value to an end value through positive and
  negative steps (revenue / EBIT / cost / headcount bridges). Native chart:
  a stacked column with an invisible base series, so data stays editable.
- scorecard: metrics with target vs. actual and a status pill per row.
"""
from __future__ import annotations

from typing import Dict, Optional, Sequence

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from ..base import add_chrome, add_line, add_rect, add_textbox, blank_slide, write_paragraph
from ..design import (add_insight_panel, fit_size, tone_rgb, warn_small,
                      write_rich_paragraph)
from ..theme import Theme, DEFAULT_THEME


def _body_box(theme, subtitle, slide):
    pal, typo, layout = theme.palette, theme.typography, theme.layout
    left = layout.margin_left_in
    width = layout.slide_width_in - layout.margin_left_in - layout.margin_right_in
    top = layout.body_top_in + 0.05
    bottom = layout.footer_top_in - 0.25
    if subtitle:
        tb = add_textbox(slide, left, top, width, 0.35)
        write_paragraph(tb.text_frame, subtitle, size=typo.section_title_size,
                        bold=True, color=pal.text_dark, family=typo.family,
                        first=True)
        top += 0.45
    return left, top, width, bottom


def add_waterfall(prs, *,
                  title: str = "[Waterfall / Insert action title]",
                  steps: Sequence[Dict],
                  subtitle: Optional[str] = None,
                  number_format: str = "#,##0",
                  insight: Optional[str] = None,
                  insight_bullets: Sequence[str] = (),
                  insight_title: str = "Key insight",
                  page_number=None, section_marker=None,
                  source=None, footnote=None,
                  theme: Theme = DEFAULT_THEME):
    """steps: [{"label": str, "value": num, "total"?: bool}]
    - total=True marks a bar drawn from zero (start, subtotals, end).
    - other steps are deltas: positive (green) or negative (red).
    Use only values the source gives; the end total should equal start + deltas.
    """
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    left, top, width, bottom = _body_box(theme, subtitle, slide)
    has_panel = bool(insight or insight_bullets)
    panel_w = 3.9 if has_panel else 0
    gap = 0.35 if has_panel else 0
    chart_w = width - panel_w - gap

    base, up, down, tot = [], [], [], []
    running = 0.0
    for s in steps:
        v = float(s["value"])
        if s.get("total"):
            running = v
            base.append(0); up.append(0); down.append(0); tot.append(v)
        elif v >= 0:
            base.append(running); up.append(v); down.append(0); tot.append(0)
            running += v
        else:
            running += v
            base.append(running); up.append(0); down.append(-v); tot.append(0)

    data = CategoryChartData()
    data.categories = [s["label"] for s in steps]
    data.add_series("_base", base)   # invisible helper; "_" = not data
    data.add_series("Increase", up)
    data.add_series("Decrease", down)
    data.add_series("Total", tot)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(left),
                                Inches(top), Inches(chart_w), Inches(bottom - top),
                                data)
    chart = gf.chart
    chart.has_title = False
    chart.has_legend = False
    chart.font.size = Pt(12)
    chart.font.name = typo.family
    plot = chart.plots[0]
    plot.gap_width = 45
    plot.overlap = 100
    colors = [None, pal.status_green, pal.status_red, pal.deep_navy]
    pos_fmt = number_format.split(";")[0]
    fmts = [None, f"+{pos_fmt};;;", f"-{pos_fmt};;;", f"{pos_fmt};-{pos_fmt};;"]
    for ps, rgb, fmt in zip(plot.series, colors, fmts):
        if rgb is None:
            ps.format.fill.background()
            ps.format.line.fill.background()
            continue
        ps.format.fill.solid()
        ps.format.fill.fore_color.rgb = rgb
        dl = ps.data_labels
        dl.number_format = fmt
        dl.number_format_is_linked = False
        dl.show_value = True
        dl.font.size = Pt(12)
        dl.font.bold = True
        dl.font.color.rgb = pal.white
        dl.position = XL_LABEL_POSITION.CENTER
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = pal.grid_gray
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(11)
    va.tick_labels.font.color.rgb = pal.footer_gray
    va.minimum_scale = 0
    chart.category_axis.tick_labels.font.size = Pt(12)

    if has_panel:
        add_insight_panel(slide, theme, left + chart_w + gap, top, panel_w,
                          bottom - top, title=insight_title, text=insight,
                          bullets=insight_bullets)
    return slide


_STATUS = {"green": "green", "on track": "green", "achieved": "green",
           "amber": "amber", "at risk": "amber", "watch": "amber",
           "red": "red", "off track": "red", "missed": "red", "behind": "red"}


def add_scorecard(prs, *,
                  title: str = "[Scorecard / Insert action title]",
                  metrics: Sequence[Dict],
                  columns: Sequence[str] = ("Metric", "Target", "Actual", "Status", "Comment"),
                  subtitle: Optional[str] = None,
                  page_number=None, section_marker=None,
                  source=None, footnote=None,
                  theme: Theme = DEFAULT_THEME):
    """metrics: [{"metric": str, "target": str, "actual": str,
                  "status": "green"|"amber"|"red" (or "on track"/"at risk"/"off track"),
                  "status_label"?: str, "comment"?: str}]
    Values are shown as given — don't compute variances the source doesn't state.
    """
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    left, top, width, bottom = _body_box(theme, subtitle, slide)

    has_comment = any(m.get("comment") for m in metrics)
    rel = [3.0, 1.3, 1.3, 1.5] + ([4.2] if has_comment else [])
    from ..labels import loc
    cols = [loc(theme, c) for c in list(columns)[:len(rel)]]
    total = sum(rel)
    ws = [width * r / total for r in rel]
    xs = [left + sum(ws[:i]) for i in range(len(ws))]

    head_h = 0.5
    n = max(len(metrics), 1)
    row_h = min(0.95, (bottom - top - head_h) / n)
    size = min(fit_size([m.get("metric", "")], ws[0] - 0.3, row_h - 0.12, max_size=16,
                        min_size=10) for m in metrics) if metrics else 14
    if has_comment:
        size = min([size] + [fit_size([m.get("comment", "")], ws[-1] - 0.3, row_h - 0.12,
                                      max_size=16, min_size=10) for m in metrics])
    warn_small("scorecard", title, size, "Shorten metric names or comments.")

    add_rect(slide, left, top, width, head_h, fill=pal.deep_navy)
    for k, (x, w, c) in enumerate(zip(xs, ws, cols)):
        left_aligned = k == 0 or k == 4          # metric and comment columns
        tb = add_textbox(slide, x + 0.15, top, w - 0.3, head_h, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, c, size=size, bold=True, color=pal.white,
                        family=typo.family,
                        align=PP_ALIGN.LEFT if left_aligned else PP_ALIGN.CENTER,
                        first=True)
    y = top + head_h
    for i, m in enumerate(metrics):
        if i % 2:
            add_rect(slide, left, y, width, row_h, fill=pal.soft_gray)
        cells = [m.get("metric", ""), m.get("target", ""), m.get("actual", "")]
        for k, (x, w, val) in enumerate(zip(xs[:3], ws[:3], cells)):
            tb = add_textbox(slide, x + 0.15, y, w - 0.3, row_h, anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, str(val), size=size, theme=theme,
                                 bold=(k == 0 or k == 2),
                                 align=PP_ALIGN.LEFT if k == 0 else PP_ALIGN.CENTER,
                                 first=True)
        tone = _STATUS.get(str(m.get("status", "")).lower(), "gray")
        label = m.get("status_label") or loc(theme, {"green": "On track", "amber": "At risk",
                                                     "red": "Off track"}.get(tone, str(m.get("status", ""))))
        pw, ph = min(ws[3] - 0.3, 1.35), min(0.38, row_h - 0.16)
        px, py = xs[3] + (ws[3] - pw) / 2, y + (row_h - ph) / 2
        add_rect(slide, px, py, pw, ph, fill=tone_rgb(theme, tone))
        tb = add_textbox(slide, px, py, pw, ph, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, label, size=max(size - 2, 10), bold=True,
                        color=pal.text_dark if tone == "amber" else pal.white,
                        family=typo.family, align=PP_ALIGN.CENTER, first=True)
        if has_comment:
            tb = add_textbox(slide, xs[4] + 0.15, y, ws[4] - 0.3, row_h,
                             anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, m.get("comment", ""), size=size,
                                 theme=theme, first=True)
        add_line(slide, left, y + row_h, left + width, y + row_h,
                 color=pal.grid_gray, width_pt=0.5)
        y += row_h
    return slide
