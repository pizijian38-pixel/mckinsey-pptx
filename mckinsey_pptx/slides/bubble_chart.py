"""Bubble / scatter slides — 4 variants:
 - bubble chart (full width)
 - bubble chart with takeaways (chart left, bullets right)
 - growth-share matrix (BCG style 2x2)
 - prioritization / assessment matrix (3x3 with status colors)
"""
from __future__ import annotations
from typing import Sequence, Optional, Tuple, List, Dict, Literal
import math

from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.util import Inches, Pt

from ..base import (
    blank_slide, add_chrome, add_rect, add_oval, add_line, add_textbox,
    write_paragraph,
)
from ..theme import Theme, DEFAULT_THEME
from ..labels import loc
from ..design import (HAIRLINE_PT, check_focus, eyebrow, has_focus, is_focus, legend_strip,
                      mark_axis, muted_fill, tint)
from .column_chart import _draw_takeaway, _draw_description_header


def _round_up_nice(v):
    if v <= 0:
        return 1
    mag = 10 ** math.floor(math.log10(v))
    n = v / mag
    if n <= 1: pick = 1
    elif n <= 2: pick = 2
    elif n <= 2.5: pick = 2.5
    elif n <= 5: pick = 5
    else: pick = 10
    return pick * mag


def _draw_xy_axis(slide, theme, *, plot_box, x_max, y_max,
                  x_label, x_unit, y_label, y_unit, n_ticks_x=10, n_ticks_y=7):
    pal, typo = theme.palette, theme.typography
    pl, pt, pw, ph = plot_box
    pr = pl + pw
    pb = pt + ph

    x_ticks = [i * x_max / (n_ticks_x - 1) for i in range(n_ticks_x)]
    y_ticks = [i * y_max / (n_ticks_y - 1) for i in range(n_ticks_y)]

    # Grid
    for tv in y_ticks:
        ty = pb - (tv / y_max) * ph
        add_line(slide, pl, ty, pr, ty, color=pal.grid_gray, width_pt=0.5)
    for tv in x_ticks:
        tx = pl + (tv / x_max) * pw
        add_line(slide, tx, pt, tx, pb, color=pal.grid_gray, width_pt=0.5)

    # Axis labels (numbers)
    for tv in y_ticks:
        ty = pb - (tv / y_max) * ph
        tb = mark_axis(add_textbox(slide, pl - 0.55, ty - 0.12, 0.45, 0.24,
                                   anchor=MSO_ANCHOR.MIDDLE))
        write_paragraph(tb.text_frame, f"{int(round(tv))}",
                        size=typo.chart_axis_size, color=pal.text_dark,
                        family=typo.family, align=PP_ALIGN.RIGHT, first=True)
    for tv in x_ticks:
        tx = pl + (tv / x_max) * pw
        tb = mark_axis(add_textbox(slide, tx - 0.30, pb + 0.05, 0.6, 0.22))
        write_paragraph(tb.text_frame, f"{int(round(tv))}",
                        size=typo.chart_axis_size, color=pal.text_dark,
                        family=typo.family, align=PP_ALIGN.CENTER, first=True)

    # Axis titles
    tb = add_textbox(slide, pl - 0.6, pt - 0.45, 3.5, 0.3)
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = f"[{y_label}], "
    r.font.size = Pt(typo.section_title_size); r.font.bold = True
    r.font.color.rgb = pal.text_dark; r.font.name = typo.family
    r2 = p.add_run(); r2.text = f"[{y_unit}]"
    r2.font.size = Pt(typo.section_title_size)
    r2.font.color.rgb = pal.placeholder_gray; r2.font.name = typo.family

    tb = add_textbox(slide, pr - 2.5, pb + 0.30, 2.5, 0.3)
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = f"[{x_label}], "
    r.font.size = Pt(typo.section_title_size); r.font.bold = True
    r.font.color.rgb = pal.text_dark; r.font.name = typo.family
    r2 = p.add_run(); r2.text = f"[{x_unit}]"
    r2.font.size = Pt(typo.section_title_size)
    r2.font.color.rgb = pal.placeholder_gray; r2.font.name = typo.family


def _draw_bubble(slide, theme, *, plot_box, x_max, y_max, bubble,
                 size_min=0.25, size_max=0.95):
    pal = theme.palette
    pl, pt, pw, ph = plot_box
    pb = pt + ph

    sz_vals = [b.get("size", 1) for b in bubble]
    s_max = max(sz_vals) if sz_vals else 1
    s_min = min(sz_vals) if sz_vals else 1

    color_map = {
        "blue_light": pal.bright_blue,
        "blue_dark": pal.dark_navy,
        "blue_royal": pal.royal_blue,
        "navy": pal.dark_navy,
    }

    for b in bubble:
        x = pl + (b["x"] / x_max) * pw
        y = pb - (b["y"] / y_max) * ph
        s = b.get("size", 1)
        if s_max == s_min:
            d = (size_min + size_max) / 2
        else:
            d = size_min + (size_max - size_min) * (s - s_min) / (s_max - s_min)
        fill = color_map.get(b.get("group"), pal.bright_blue)
        if isinstance(b.get("color"), str):
            fill = color_map.get(b["color"], fill)
        add_oval(slide, x - d / 2, y - d / 2, d, d, fill=fill)
        # Label — placement via b["label_pos"] in {"right","left","top","bottom"}
        # (default "right"). Use when clusters would otherwise overlap.
        if b.get("label"):
            pos = b.get("label_pos", "right")
            lw, lh = 1.2, 0.24
            if pos == "left":
                lx, ly, align = x - d / 2 - lw - 0.05, y - 0.12, PP_ALIGN.RIGHT
            elif pos == "top":
                lx, ly, align = x - lw / 2, y - d / 2 - lh - 0.02, PP_ALIGN.CENTER
            elif pos == "bottom":
                lx, ly, align = x - lw / 2, y + d / 2 + 0.02, PP_ALIGN.CENTER
            else:  # "right"
                lx, ly, align = x + d / 2 + 0.05, y - 0.12, PP_ALIGN.LEFT
            tb = add_textbox(slide, lx, ly, lw, lh, anchor=MSO_ANCHOR.MIDDLE)
            write_paragraph(tb.text_frame, b["label"],
                            size=theme.typography.chart_label_size,
                            color=pal.text_dark,
                            family=theme.typography.family,
                            align=align, first=True)


def _draw_legend_groups(slide, theme, *, top_left, groups, with_size_swatch=True):
    """groups: [(color_name, label)]"""
    pal, typo = theme.palette, theme.typography
    color_map = {
        "blue_light": pal.bright_blue,
        "blue_dark": pal.dark_navy,
        "blue_royal": pal.royal_blue,
        "navy": pal.dark_navy,
    }
    x, y = top_left
    for cname, label in groups:
        rgb = color_map.get(cname, pal.bright_blue)
        d = 0.28
        add_oval(slide, x, y - 0.05, d, d, fill=rgb)
        tb = add_textbox(slide, x + d + 0.08, y - 0.06, 1.5, 0.30,
                         anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, f"[{label}]", size=typo.chart_label_size,
                        color=pal.text_dark, family=typo.family, first=True)
        x += d + 1.5
    if with_size_swatch:
        d = 0.28
        add_oval(slide, x, y - 0.05, d, d, fill=None,
                 line=pal.placeholder_gray, line_width=0.75)
        tb = add_textbox(slide, x + d + 0.08, y - 0.06, 2.4, 0.30,
                         anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, "Size = [Insert dimension]",
                        size=typo.chart_label_size, color=pal.text_dark,
                        family=typo.family, first=True)


# ---------- Public builders ----------

def add_bubble_chart(prs, *, title="[Bubble chart / Insert action title]",
                    bubbles: Sequence[Dict],
                    x_max=900, y_max=3000,
                    x_label="Dimension 2", x_unit="Unit",
                    y_label="Dimension 1", y_unit="Unit",
                    groups: Sequence[Tuple[str, str]] = (
                        ("blue_light", "Insert group 1"),
                        ("blue_dark", "Insert group 2"),
                        ("blue_royal", "Insert group 3"),
                    ),
                    diagonal=True,
                    state_top_left: Optional[str] = "State e.g. profitable",
                    state_bottom_right: Optional[str] = "State, e.g. unprofitable",
                    page_number=None, section_marker=None,
                    source="xx", footnote="1. xx",
                    theme: Theme = DEFAULT_THEME):
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo, layout = theme.palette, theme.typography, theme.layout

    # Legend top
    _draw_legend_groups(slide, theme, top_left=(3.4, 1.55), groups=groups)

    plot_box = (1.05, 2.05, 11.7, 4.6)
    _draw_xy_axis(slide, theme, plot_box=plot_box, x_max=x_max, y_max=y_max,
                  x_label=x_label, x_unit=x_unit,
                  y_label=y_label, y_unit=y_unit)
    pl, pt, pw, ph = plot_box

    # Diagonal reference
    if diagonal:
        line = add_line(slide, pl, pt + ph, pl + pw, pt,
                         color=pal.text_dark, width_pt=0.75,
                         dash=MSO_LINE_DASH_STYLE.DASH)

    # State labels
    if state_top_left:
        tb = add_textbox(slide, pl + 0.3, pt + 0.15, 3.5, 0.3)
        write_paragraph(tb.text_frame, f"[{state_top_left}]",
                        size=typo.chart_label_size, italic=True,
                        color=pal.placeholder_gray, family=typo.family,
                        first=True)
    if state_bottom_right:
        tb = add_textbox(slide, pl + pw - 3.5, pt + ph - 0.45, 3.4, 0.3)
        write_paragraph(tb.text_frame, f"[{state_bottom_right}]",
                        size=typo.chart_label_size, italic=True,
                        color=pal.placeholder_gray, family=typo.family,
                        align=PP_ALIGN.RIGHT, first=True)

    _draw_bubble(slide, theme, plot_box=plot_box, x_max=x_max, y_max=y_max,
                 bubble=bubbles)
    return slide


def add_bubble_chart_with_takeaways(prs, *,
                                    title="[Bubble chart with takeaways / Insert action title]",
                                    bubbles, x_max=900, y_max=3000,
                                    x_label="Dimension 2", x_unit="Unit",
                                    y_label="Dimension 1", y_unit="Unit",
                                    groups=(
                                        ("blue_dark", "Insert group 1"),
                                        ("blue_light", "Insert group 2"),
                                        ("mid_blue", "Insert group 3"),
                                    ),
                                    takeaways=(),
                                    description: str = "Description",
                                    takeaway_header: str = "Key takeaways/main conclusion",
                                    page_number=None, section_marker=None,
                                    source="xx", footnote=None,
                                    theme: Theme = DEFAULT_THEME):
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography

    # Description header on left side
    _draw_description_header(slide, theme, left=0.45, top=1.45, width=8.5,
                              label=description)

    plot_box = (1.05, 2.30, 7.9, 3.55)
    _draw_xy_axis(slide, theme, plot_box=plot_box, x_max=x_max, y_max=y_max,
                  x_label=x_label, x_unit=x_unit,
                  y_label=y_label, y_unit=y_unit)
    _draw_bubble(slide, theme, plot_box=plot_box, x_max=x_max, y_max=y_max,
                 bubble=bubbles)
    # Legend below chart (kept clear of axis title)
    _draw_legend_groups(slide, theme, top_left=(1.05, 6.55), groups=groups)
    # Takeaway right side
    _draw_takeaway(slide, theme, takeaways=takeaways,
                   header=takeaway_header,
                   box=(9.45, 1.45, 3.45, 5.0))
    return slide


def add_growth_share_matrix(prs, *,
                            title="[Growth-share matrix / Insert action title]",
                            bus: Sequence[Dict],
                            x_max=100, y_max=50,
                            x_label: str = "Relative market share (%)",
                            y_label: str = "Market growth (%)",
                            focus=None,
                            focus_note: Optional[str] = None,
                            size_label: Optional[str] = None,
                            page_number=None, section_marker=None,
                            source="xx", footnote="1. xx",
                            theme: Theme = DEFAULT_THEME):
    """bus: list of {"name", "x" (market share), "y" (growth), "size", "label_pos"?}
    focus: the unit(s) the title is about (name or index) — drawn in the accent,
           every other bubble muted. focus_note: one-line annotation on it.
    size_label: what bubble area encodes ("revenue") — shown in the legend.
    Bubble AREA is proportional to size.
    """
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    names = [b.get("name", "") for b in bus]
    check_focus("growth_share", title, focus, names)
    focused = has_focus(focus)
    legend = focused or bool(size_label)

    pl, pt, pw = 1.15, 1.85, 11.0
    ph = 4.2 if legend else 4.55
    pr, pb = pl + pw, pt + ph
    midx, midy = pl + pw / 2, pt + ph / 2

    # structure as hairlines: two axes + the quadrant split, no filled quadrants
    add_line(slide, pl, pb, pr, pb, color=pal.text_dark, width_pt=HAIRLINE_PT)
    add_line(slide, pl, pt, pl, pb, color=pal.text_dark, width_pt=HAIRLINE_PT)
    add_line(slide, midx, pt, midx, pb, color=pal.rule_gray, width_pt=HAIRLINE_PT)
    add_line(slide, pl, midy, pr, midy, color=pal.rule_gray, width_pt=HAIRLINE_PT)

    ystep, xstep = _round_up_nice(y_max / 5), _round_up_nice(x_max / 5)
    v = 0
    while v <= y_max + 1e-9:
        ty = pb - v / y_max * ph
        tb = mark_axis(add_textbox(slide, pl - 0.55, ty - 0.11, 0.45, 0.22,
                                   anchor=MSO_ANCHOR.MIDDLE))
        write_paragraph(tb.text_frame, f"{v:g}", size=typo.chart_axis_size,
                        color=pal.footer_gray, family=typo.family, align=PP_ALIGN.RIGHT,
                        first=True)
        v += ystep
    v = 0
    while v <= x_max + 1e-9:
        tx = pl + v / x_max * pw
        tb = mark_axis(add_textbox(slide, tx - 0.3, pb + 0.05, 0.6, 0.22))
        write_paragraph(tb.text_frame, f"{v:g}", size=typo.chart_axis_size,
                        color=pal.footer_gray, family=typo.family, align=PP_ALIGN.CENTER,
                        first=True)
        v += xstep
    eyebrow(slide, theme, pl - 0.55, pt - 0.45, 6, y_label, color=pal.text_dark)
    eyebrow(slide, theme, pr - 6, pb + 0.3, 6, x_label, color=pal.text_dark,
            align=PP_ALIGN.RIGHT)

    # bubbles: area ~ size; focal bubble in the accent, the rest muted
    sizes = [max(float(b.get("size", 1)), 0.0) for b in bus] or [1.0]
    s_max = max(sizes) or 1.0
    d_max = 1.1
    base = muted_fill(theme) if focused else pal.dark_navy
    order = sorted(range(len(bus)), key=lambda i: is_focus(focus, i, names[i]))
    for i in order:
        b = bus[i]
        f = is_focus(focus, i, names[i])
        d = max(0.22, d_max * math.sqrt(sizes[i] / s_max))
        x = pl + b["x"] / x_max * pw
        y = pb - b["y"] / y_max * ph
        add_oval(slide, x - d / 2, y - d / 2, d, d, fill=pal.bright_blue if f else base,
                 line=pal.white, line_width=1.0)
        if b.get("name"):
            pos = b.get("label_pos", "right")
            lw, lh = 1.9, 0.26
            if pos == "left":
                lx, ly, al = x - d / 2 - lw - 0.06, y - lh / 2, PP_ALIGN.RIGHT
            elif pos == "top":
                lx, ly, al = x - lw / 2, y - d / 2 - lh - 0.02, PP_ALIGN.CENTER
            elif pos == "bottom":
                lx, ly, al = x - lw / 2, y + d / 2 + 0.02, PP_ALIGN.CENTER
            else:
                lx, ly, al = x + d / 2 + 0.06, y - lh / 2, PP_ALIGN.LEFT
            tb = add_textbox(slide, lx, ly, lw, lh, anchor=MSO_ANCHOR.MIDDLE)
            strong = f or not focused
            write_paragraph(tb.text_frame, b["name"], size=typo.chart_label_size + 1,
                            bold=strong, color=pal.deep_navy if strong else pal.footer_gray,
                            family=typo.family, align=al, first=True)
        if f and focus_note:
            nx, ny = x + d / 2 + 0.2, max(pt + 0.05, y - d / 2 - 0.6)
            add_line(slide, x + d * 0.3, y - d * 0.4, nx, ny + 0.3, color=pal.bright_blue,
                     width_pt=HAIRLINE_PT)
            tb = add_textbox(slide, nx + 0.05, ny, 3.6, 0.5)
            write_paragraph(tb.text_frame, focus_note, size=typo.chart_label_size + 1,
                            italic=True, color=pal.deep_navy, family=typo.family, first=True)

    # quadrant names last, so they sit on top
    for name, x, y, al in (("Question marks", pl + 0.15, pt + 0.08, PP_ALIGN.LEFT),
                           ("Stars", pr - 3.15, pt + 0.08, PP_ALIGN.RIGHT),
                           ("Dogs", pl + 0.15, pb - 0.34, PP_ALIGN.LEFT),
                           ("Cash cows", pr - 3.15, pb - 0.34, PP_ALIGN.RIGHT)):
        eyebrow(slide, theme, x, y, 3.0, name, align=al)

    if legend:
        items = ([("dot", pal.bright_blue, "Focus"), ("dot", base, "Other units")]
                 if focused else [])
        legend_strip(slide, theme, items, pb + 0.58,
                     note=f"{loc(theme, 'Bubble area')} = {size_label}" if size_label else None)
    return slide


def add_prioritization_matrix(prs, *,
                              title="[Prioritization or assessment matrix / Insert action title]",
                              items: Sequence[Dict],
                              page_number=None, section_marker=None,
                              source="xx", footnote="1. xx",
                              description: str = "[Description]",
                              legend: Sequence[str] = ("[Insert status/group]",) * 3,
                              focus=None,
                              theme: Theme = DEFAULT_THEME):
    """items: [{"name", "x_band": 0|1|2 (Low/Med/High), "y_band": 0|1|2 (Short/Med/Long),
               "status": green|amber|red, "ox"?, "oy"?, "d"?}]
    legend: labels for the green / amber / red status dots, in that order.
    focus: the item(s) the title is about (name or index) — ringed in navy, the
           rest drawn lighter. The short-time x high-impact cell is the target zone.
    """
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, page_number=page_number,
               section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    names = [it.get("name", "") for it in items]
    check_focus("prioritization_matrix", title, focus, names)
    focused = has_focus(focus)

    if description:
        tb = add_textbox(slide, 0.45, 1.42, 12.4, 0.3)
        write_paragraph(tb.text_frame, description, size=typo.section_title_size, bold=True,
                        color=pal.text_dark, family=typo.family, first=True)

    pl, pt, pw, ph = 1.85, 2.15, 11.0, 3.9
    cell_w, cell_h = pw / 3, ph / 3
    # target zone: short time to impact x high impact — tinted, labelled, no solid block
    add_rect(slide, pl + 2 * cell_w, pt, cell_w, cell_h, fill=tint(pal.bright_blue, 0.86))
    eyebrow(slide, theme, pl + 2 * cell_w + 0.1, pt + 0.06, cell_w - 0.2, "Priority",
            color=pal.mid_blue, align=PP_ALIGN.RIGHT, size=9)
    for i in range(1, 3):
        add_line(slide, pl + i * cell_w, pt, pl + i * cell_w, pt + ph, color=pal.grid_gray,
                 width_pt=HAIRLINE_PT)
        add_line(slide, pl, pt + i * cell_h, pl + pw, pt + i * cell_h, color=pal.grid_gray,
                 width_pt=HAIRLINE_PT)
    add_line(slide, pl, pt + ph, pl + pw, pt + ph, color=pal.text_dark, width_pt=HAIRLINE_PT)
    add_line(slide, pl, pt, pl, pt + ph, color=pal.text_dark, width_pt=HAIRLINE_PT)

    for i, label in enumerate(["Short", "Medium", "Long"]):
        tb = add_textbox(slide, pl - 1.0, pt + i * cell_h, 0.9, cell_h, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, loc(theme, label), size=typo.chart_label_size,
                        color=pal.footer_gray, family=typo.family, align=PP_ALIGN.RIGHT,
                        first=True)
    for i, label in enumerate(["Low", "Medium", "High"]):
        tb = add_textbox(slide, pl + i * cell_w, pt + ph + 0.06, cell_w, 0.26)
        write_paragraph(tb.text_frame, loc(theme, label), size=typo.chart_label_size,
                        color=pal.footer_gray, family=typo.family, align=PP_ALIGN.CENTER,
                        first=True)
    eyebrow(slide, theme, pl, pt - 0.3, 4, "TIME TO IMPACT", color=pal.text_dark)
    eyebrow(slide, theme, pl + pw - 5, pt + ph + 0.34, 5, "LEVEL OF IMPACT",
            color=pal.text_dark, align=PP_ALIGN.RIGHT)

    color_map = {"green": pal.status_green, "amber": pal.status_amber, "red": pal.status_red}
    order = sorted(range(len(items)), key=lambda i: is_focus(focus, i, names[i]))
    for i in order:
        it = items[i]
        f = is_focus(focus, i, names[i])
        x = pl + (it.get("x_band", 1) + it.get("ox", 0.5)) * cell_w
        y = pt + (it.get("y_band", 1) + it.get("oy", 0.5)) * cell_h
        d = it.get("d", 0.85)
        st = it.get("status", "green")
        fill = color_map.get(st, pal.status_green)
        if focused and not f:
            fill = tint(fill, 0.45)
        add_oval(slide, x - d / 2, y - d / 2, d, d, fill=fill,
                 line=pal.deep_navy if f else pal.white, line_width=2.5 if f else 1.0)
        tb = add_textbox(slide, x - d / 2, y - d / 2, d, d, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, it.get("name", ""), size=typo.chart_label_size,
                        bold=f, color=pal.text_dark if st == "amber" or (focused and not f)
                        else pal.white, family=typo.family, align=PP_ALIGN.CENTER, first=True)

    used = [k for k in ("green", "amber", "red") if any(it.get("status", "green") == k
                                                          for it in items)]
    labels = dict(zip(("green", "amber", "red"), legend))
    legend_strip(slide, theme, [("dot", color_map[k], labels[k]) for k in used],
                 pt + ph + 0.62)
    return slide


