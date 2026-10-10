"""Part-of-whole and flow diagrams where AREA or WIDTH is the data.

- marimekko : two-way part-of-whole. Column width = the column's share of the
              grand total; segment height = the series' share within the column;
              so every segment's area is its share of the whole.
- treemap   : one-way part-of-whole, squarified cells, area = value.
- sankey    : a quantity splitting and merging across 2-4 stages; node height
              and ribbon width = value, on one scale for the whole figure.

Honesty rules (enforced, not advisory):
- values must be >= 0; a negative value raises ValueError (area cannot be negative);
- a node in a sankey whose inflow and outflow differ prints a WARNING;
- every data mark is named "<kind>:<label>" so tests can measure it.
"""
from __future__ import annotations

import math
from typing import Dict, List, Optional, Sequence

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from ..base import add_line, add_rect, add_textbox, write_paragraph
from ..components import _polygon
from ..design import (HAIRLINE_PT, check_focus, eyebrow, fit_size, fmt_num, has_focus,
                      is_focus, legend_strip, mark_axis, mark_derived, pct, set_alpha,
                      text_height_in, tint, warn_small)
from ..labels import loc
from ..theme import Theme, DEFAULT_THEME
from .evaluation_slides import _frame

MIN_PT = 10


def _check_values(where, title, vals):
    bad = [v for v in vals if v is None or float(v) < 0]
    if bad:
        raise ValueError(f"{where} \"{str(title)[:40]}\": values must be >= 0 (got {bad[:3]}); "
                         "area cannot encode a negative amount — use a waterfall or bar chart.")


def _ramp(theme, k, i, muted=False):
    """Fill for series / rank i of k: strongest first. Muted = the light ramp used
    behind a focal accent."""
    t = (0.55 + 0.33 * i / max(k - 1, 1)) if muted else (0.72 * i / max(k - 1, 1))
    return tint(theme.palette.dark_navy, t), t < 0.45


def _cell_text(slide, theme, x, y, w, h, lines, *, light, max_size=12, derived=True):
    """Name / value lines inside a cell, top-left; nothing if it doesn't fit at 10pt."""
    pal = theme.palette
    pad = 0.08
    lines = [l for l in lines if l]
    while lines:
        size = fit_size(lines, w - 2 * pad, h - 2 * pad, max_size=max_size, min_size=MIN_PT,
                        para_gap_pt=1)
        if size >= MIN_PT and text_height_in(lines, w - 2 * pad, size, para_gap_pt=1) <= h - pad:
            break
        lines = lines[:-1]
    if not lines:
        return False
    tb = add_textbox(slide, x + pad, y + pad * 0.6, w - 2 * pad, h - pad)
    if derived:                # shares and totals are computed, not typed
        mark_derived(tb)
    col = pal.white if light else pal.text_dark
    for j, t in enumerate(lines):
        write_paragraph(tb.text_frame, t, size=size, bold=j == 0, color=col,
                        family=theme.typography.family, first=j == 0)
    return True


# ---------------------------------------------------------------- marimekko

def add_marimekko(prs, *,
                  title: str = "[Marimekko / Insert action title]",
                  series: Sequence[str],
                  columns: Sequence[Dict],
                  fmt: Optional[str] = None,
                  cell_label: str = "share",
                  show_totals: bool = True,
                  focus=None,
                  subtitle: Optional[str] = None,
                  insight: Optional[str] = None,
                  insight_label: Optional[str] = "Key insight",
                  page_number=None, section_marker=None,
                  source=None, footnote=None,
                  theme: Theme = DEFAULT_THEME):
    """series: names of the stacked parts, top to bottom (2-5).
    columns: [{"name", "values": [one per series, >= 0]}] (3-8), left to right.
    cell_label: "share" (within the column), "value", "both" or "none".
    fmt: format string for values, e.g. "${:,.1f}B".
    focus: a column name, a series name, or one cell as "Column / Series".
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    k, n = len(series), len(columns)
    vals = [[float(v) for v in (c.get("values") or [])] + [0.0] * (k - len(c.get("values") or []))
            for c in columns]
    _check_values("marimekko", title, [v for row in vals for v in row])
    if not 2 <= k <= 5 or not 2 <= n <= 8:
        warn_small("marimekko", title, 0, f"{n} columns x {k} series; keep to 3-8 x 2-5 "
                   "(fold the tail into 'Other').")
    names = [c.get("name", "") for c in columns]
    cells = [f"{c} / {s}" for c in names for s in series]
    check_focus("marimekko", title, focus, names + list(series) + cells)
    focused = has_focus(focus)

    totals = [sum(r) for r in vals]
    grand = sum(totals) or 1.0
    axis_w = 0.5
    px, pw = left + axis_w, width - axis_w
    py = top + (0.3 if show_totals else 0.05)
    ph = bottom - py - 0.62 - 0.45
    gut = 0.04
    live = [i for i in range(n) if totals[i] > 0]
    usable = pw - gut * (len(live) - 1)

    for v, lab in ((0, "0%"), (0.5, "50%"), (1, "100%")):
        tb = mark_axis(add_textbox(slide, left, py + (1 - v) * ph - 0.11, axis_w - 0.1, 0.22,
                                   anchor=MSO_ANCHOR.MIDDLE))
        write_paragraph(tb.text_frame, lab, size=typo.chart_axis_size, color=pal.footer_gray,
                        family=typo.family, align=PP_ALIGN.RIGHT, first=True)
    x = px
    for i in live:
        cw = usable * totals[i] / grand
        y = py
        for j, s in enumerate(series):
            v = vals[i][j]
            if v <= 0:
                continue          # absent series: omitted, never drawn at zero height
            h = ph * v / totals[i]
            f = focused and (is_focus(focus, -1, names[i]) or is_focus(focus, -1, s)
                             or is_focus(focus, -1, f"{names[i]} / {s}"))
            fill, light = (pal.bright_blue, True) if f else _ramp(theme, k, j, muted=focused)
            box = add_rect(slide, x, y, cw, h, fill=fill, line=pal.white, line_width=1.0)
            box.name = f"mekko:{names[i]}|{s}"
            if cell_label != "none":
                share = pct(v / totals[i])
                val = fmt_num(v, fmt)
                body = {"share": share, "value": val, "both": f"{val} · {share}"}[cell_label]
                _cell_text(slide, theme, x, y, cw, h, [body], light=light, max_size=12,
                           derived=cell_label != "value")
            y += h
        if show_totals:
            tb = mark_derived(add_textbox(slide, x, top, cw, 0.26, anchor=MSO_ANCHOR.BOTTOM))
            write_paragraph(tb.text_frame, fmt_num(totals[i], fmt), size=typo.chart_label_size,
                            bold=True, color=pal.text_dark, family=typo.family,
                            align=PP_ALIGN.CENTER, first=True)
        tb = mark_derived(add_textbox(slide, x - 0.05, py + ph + 0.06, cw + 0.1, 0.5))
        fcol = focused and is_focus(focus, -1, names[i])
        write_paragraph(tb.text_frame, names[i], size=typo.chart_label_size, bold=True,
                        color=pal.mid_blue if fcol else pal.text_dark, family=typo.family,
                        align=PP_ALIGN.CENTER, first=True)
        write_paragraph(tb.text_frame, pct(totals[i] / grand), size=typo.chart_label_size,
                        color=pal.footer_gray, family=typo.family, align=PP_ALIGN.CENTER)
        x += cw + gut

    items = []
    for j, s in enumerate(series):
        f = focused and is_focus(focus, -1, s)
        items.append(("box", pal.bright_blue if f else _ramp(theme, k, j, muted=focused)[0], s))
    note = loc(theme, "Column width = share of total")
    legend_strip(slide, theme, items, bottom - 0.42, note=note)
    return slide


# ---------------------------------------------------------------- treemap

def _squarify(areas, x, y, w, h):
    """Squarified layout (Bruls et al. 2000). areas sorted descending, summing to w*h."""
    rects, i = [], 0

    def worst(row, side):
        s = sum(row)
        return max(max(side * side * r / (s * s), (s * s) / (side * side * r)) for r in row)
    while i < len(areas):
        side = min(w, h)
        row = [areas[i]]
        i += 1
        while i < len(areas) and worst(row + [areas[i]], side) <= worst(row, side):
            row.append(areas[i])
            i += 1
        s = sum(row)
        if w >= h:
            rw, yy = s / h, y
            for a in row:
                rects.append((x, yy, rw, a / rw))
                yy += a / rw
            x, w = x + rw, w - rw
        else:
            rh, xx = s / w, x
            for a in row:
                rects.append((xx, y, a / rh, rh))
                xx += a / rh
            y, h = y + rh, h - rh
    return rects


def add_treemap(prs, *,
                title: str = "[Treemap / Insert action title]",
                items: Sequence[Dict],
                fmt: Optional[str] = None,
                max_cells: int = 8,
                other_label: str = "Other",
                focus=None,
                subtitle: Optional[str] = None,
                insight: Optional[str] = None,
                insight_label: Optional[str] = "Key insight",
                page_number=None, section_marker=None,
                source=None, footnote=None,
                theme: Theme = DEFAULT_THEME):
    """items: [{"name", "value" >= 0}] (4-8). Squarified cells, area = value.
    Beyond max_cells the smallest items are folded into one "Other" cell (named
    in the note under the map) — never dropped. focus: the item the title is
    about (accent; the rest muted). Without focus cells shade by rank.
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal = theme.palette
    data = [(str(it.get("name", "")), float(it.get("value", 0))) for it in items]
    _check_values("treemap", title, [v for _, v in data])
    data = sorted([d for d in data if d[1] > 0], key=lambda d: -d[1])
    folded = []
    if len(data) > max_cells:
        folded = data[max_cells - 1:]
        data = data[:max_cells - 1] + [(loc(theme, other_label), sum(v for _, v in folded))]
        warn_small("treemap", title, 0, f"{len(folded)} smallest items folded into "
                   f"'{other_label}' (named in the note).")
    check_focus("treemap", title, focus, [d[0] for d in data])
    focused = has_focus(focus)
    total = sum(v for _, v in data) or 1.0
    ph = bottom - top - 0.5
    rects = _squarify([v / total * width * ph for _, v in data], left, top, width, ph)
    unlabeled = []
    for r, ((name, v), (x, y, w, h)) in enumerate(zip(data, rects)):
        f = is_focus(focus, r, name)
        fill, light = (pal.bright_blue, True) if f else _ramp(theme, len(data), r, muted=focused)
        box = add_rect(slide, x, y, w, h, fill=fill, line=pal.white, line_width=1.5)
        box.name = f"treemap:{name}"
        share = pct(v / total)
        if not _cell_text(slide, theme, x, y, w, h,
                          [name, f"{fmt_num(v, fmt)} · {share}"], light=light, max_size=14):
            unlabeled.append(f"{name} {share}")
    notes = []
    if unlabeled:
        notes.append(f"{loc(theme, 'Not labelled')}: " + ", ".join(unlabeled))
    if folded:
        notes.append(f"{loc(theme, other_label)}: " + ", ".join(n for n, _ in folded))
    note = "; ".join(notes) or loc(theme, "Area = value")
    tb = add_textbox(slide, left, top + ph + 0.1, width, 0.3, anchor=MSO_ANCHOR.MIDDLE)
    write_paragraph(tb.text_frame, note, size=MIN_PT, italic=True, color=pal.footer_gray,
                    family=theme.typography.family, first=True)
    return slide


# ---------------------------------------------------------------- sankey

def _flow_label(f):
    return f"{f['from']} → {f['to']}"


def add_sankey(prs, *,
               title: str = "[Sankey / Insert action title]",
               flows: Sequence[Dict],
               stages: Optional[Sequence[str]] = None,
               fmt: Optional[str] = None,
               focus=None,
               subtitle: Optional[str] = None,
               insight: Optional[str] = None,
               insight_label: Optional[str] = "Key insight",
               page_number=None, section_marker=None,
               source=None, footnote=None,
               theme: Theme = DEFAULT_THEME):
    """flows: [{"from", "to", "value" >= 0}]. Stages (columns) follow from the
    links: sources on the left, each node one column right of its furthest source
    (2-4 columns, <= 10 nodes, <= 14 flows). Node height and ribbon width share
    one scale. stages: optional column headers.
    focus: a node (its bar and every ribbon touching it) or one flow "A → B".
    A middle node whose inflow and outflow differ prints a WARNING.
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    fl = [dict(f, value=float(f.get("value", 0))) for f in flows]
    _check_values("sankey", title, [f["value"] for f in fl])
    fl = [f for f in fl if f["value"] > 0]
    order: List[str] = []
    for f in fl:
        for nme in (f["from"], f["to"]):
            if nme not in order:
                order.append(nme)
    depth = {nme: 0 for nme in order}
    for _ in range(len(order)):
        for f in fl:
            depth[f["to"]] = max(depth[f["to"]], depth[f["from"]] + 1)
    ncol = max(depth.values(), default=0) + 1
    inflow = {nme: sum(f["value"] for f in fl if f["to"] == nme) for nme in order}
    outflow = {nme: sum(f["value"] for f in fl if f["from"] == nme) for nme in order}
    for nme in order:
        if inflow[nme] and outflow[nme] and abs(inflow[nme] - outflow[nme]) > 0.005 * max(
                inflow[nme], outflow[nme]):
            warn_small("sankey", title, 0, f"'{nme}' takes in {fmt_num(inflow[nme], fmt)} but "
                       f"sends on {fmt_num(outflow[nme], fmt)} — add the missing flow "
                       "(e.g. an 'Other' / 'Lost' node) so volume is conserved.")
    if not 2 <= ncol <= 4 or len(order) > 10 or len(fl) > 14:
        warn_small("sankey", title, 0, f"{ncol} stages, {len(order)} nodes, {len(fl)} flows; "
                   "keep to 2-4 / 10 / 14 or split into two sankeys.")
    check_focus("sankey", title, focus, order + [_flow_label(f) for f in fl])
    focused = has_focus(focus)
    thru = {nme: max(inflow[nme], outflow[nme]) for nme in order}
    cols = [[nme for nme in order if depth[nme] == c] for c in range(ncol)]

    head_h = 0.3 if stages else 0.0
    y0, y1 = top + head_h + (0.35 if ncol > 2 else 0.05), bottom - 0.1
    lab_w = 1.9
    bar_w = 0.16
    xs = [left + lab_w + (width - 2 * lab_w - bar_w) * c / max(ncol - 1, 1) for c in range(ncol)]
    gap = 0.36
    k = min((y1 - y0 - gap * (len(c) - 1)) / (sum(thru[x] for x in c) or 1) for c in cols)

    pos: Dict[str, List[float]] = {}

    def place(col):
        tot = sum(thru[x] for x in col) * k + gap * (len(col) - 1)
        y = y0 + (y1 - y0 - tot) / 2
        for nme in col:
            pos[nme] = [y, y + thru[nme] * k]
            y += thru[nme] * k + gap
    place(cols[0])
    for c in range(1, ncol):   # barycentre order against the sources, then place
        def bary(nme):
            src = [(f["value"], sum(pos[f["from"]]) / 2) for f in fl if f["to"] == nme
                   and f["from"] in pos]
            return sum(v * y for v, y in src) / sum(v for v, _ in src) if src else 0
        cols[c].sort(key=bary)
        place(cols[c])

    # ribbons: outgoing stacked by target position, incoming by source position
    out_y = {nme: pos[nme][0] for nme in order}
    in_y = {nme: pos[nme][0] for nme in order}
    def is_hot(f):
        return focused and (is_focus(focus, -1, f["from"]) or is_focus(focus, -1, f["to"])
                            or is_focus(focus, -1, _flow_label(f))
                            or is_focus(focus, -1, f"{f['from']} -> {f['to']}"))
    srt = sorted(fl, key=lambda f: (depth[f["from"]], pos[f["from"]][0], pos[f["to"]][0]))
    geo = []
    for f in srt:
        w_ = f["value"] * k
        a0 = out_y[f["from"]]
        out_y[f["from"]] += w_
        geo.append([f, a0, w_])
    for g in sorted(geo, key=lambda g: (pos[g[0]["to"]][0], pos[g[0]["from"]][0])):
        g.append(in_y[g[0]["to"]])
        in_y[g[0]["to"]] += g[2]
    for hot_pass in (False, True):   # accent ribbons last, so they read on top
        for f, a0, w_, b0 in geo:
            if is_hot(f) != hot_pass:
                continue
            xa = xs[depth[f["from"]]] + bar_w
            xb = xs[depth[f["to"]]]
            mx = (xa + xb) / 2
            top_e, bot_e = [], []
            for t in [i / 24 for i in range(25)]:
                u = 1 - t
                bx = u ** 3 * xa + 3 * u * u * t * mx + 3 * u * t * t * mx + t ** 3 * xb
                by = u ** 3 * a0 + 3 * u * u * t * a0 + 3 * u * t * t * b0 + t ** 3 * b0
                top_e.append((bx, by))
                bot_e.append((bx, by + w_))
            shp = _polygon(slide, top_e + bot_e[::-1],
                           pal.bright_blue if hot_pass else tint(pal.dark_navy, 0.55))
            set_alpha(shp, 0.6 if hot_pass else 0.35)
            shp.name = f"sankey:{_flow_label(f)}"
    for nme in order:
        c = depth[nme]
        ya, yb = pos[nme]
        f = focused and is_focus(focus, -1, nme)
        bar = add_rect(slide, xs[c], ya, bar_w, yb - ya, fill=pal.bright_blue if f else pal.deep_navy)
        bar.name = f"sankey-node:{nme}"
        txt = [nme, fmt_num(thru[nme], fmt)]
        if c == 0:
            tb = add_textbox(slide, xs[c] - lab_w + 0.05, ya, lab_w - 0.15, max(yb - ya, 0.45),
                             anchor=MSO_ANCHOR.MIDDLE)
            al = PP_ALIGN.RIGHT
        elif c == ncol - 1:
            tb = add_textbox(slide, xs[c] + bar_w + 0.1, ya, lab_w - 0.15, max(yb - ya, 0.45),
                             anchor=MSO_ANCHOR.MIDDLE)
            al = PP_ALIGN.LEFT
        else:   # middle stage: one line in the gap above the bar
            txt = [f"{nme} · {fmt_num(thru[nme], fmt)}"]
            lw = 1.9
            tb = add_textbox(slide, xs[c] + bar_w / 2 - lw / 2, ya - 0.3, lw, 0.26,
                             anchor=MSO_ANCHOR.MIDDLE)
            al = PP_ALIGN.CENTER
        if not any(abs(f["value"] - thru[nme]) < 1e-9 for f in fl):
            mark_derived(tb)    # a sum over several flows, not a typed value
        for j, t in enumerate(txt):
            write_paragraph(tb.text_frame, t, size=typo.chart_label_size + (1 if j == 0 else 0),
                            bold=j == 0, color=pal.text_dark if j == 0 else pal.footer_gray,
                            family=typo.family, align=al, first=j == 0)
    if stages:
        for c, s in enumerate(list(stages)[:ncol]):
            al = PP_ALIGN.RIGHT if c == 0 else (PP_ALIGN.LEFT if c == ncol - 1 else PP_ALIGN.CENTER)
            x = xs[c] - 1.7 + bar_w if c == 0 else (xs[c] if c == ncol - 1 else xs[c] - 0.9)
            eyebrow(slide, theme, x, top, 1.8 if c in (0, ncol - 1) else 1.96, s, align=al)
    return slide
