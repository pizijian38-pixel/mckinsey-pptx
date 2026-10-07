"""Logic slides: one page that carries an argument, not a list.

Patterns taken from strategy-case decks (situation analysis -> challenge ->
options -> recommendation), where most pages make a reader follow a chain
of reasoning in one look.

- logic_grid          : rows x logical stages (situation -> capability ->
                        advantage), flow headers, an arrow into a dark
                        conclusion per row. direction="down" turns it into
                        columns x stages (PEST / five forces: "current
                        situation" -> "influence on us").
- strategic_challenge : drivers -> possible results -> converge on one key
                        threat -> the key question ("How can X ..., given ...?")
- storyline_summary   : executive summary that tells the deck's story in one
                        page: situation -> key question -> options (with the
                        recommended one marked) -> recommendation.
"""
from __future__ import annotations

from typing import Dict, List, Optional, Sequence, Union

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from ..base import add_line, add_rect, add_textbox, write_paragraph
from ..components import letter_space
from ..design import (add_arrow, add_fade_wedge, add_icon, add_stage_banners, add_wedge,
                      fit_one_line, fit_size, plain, set_dashed, text_height_in, text_width_pt,
                      tint, tone_rgb, warn_small, write_rich_paragraph)
from ..theme import Theme, DEFAULT_THEME
from .evaluation_slides import _frame

Cell = Union[None, str, Sequence[str], Dict]
MAX_PT, MIN_PT = 16, 10
BULLET_INDENT = 0.22


# ---------- shared text helpers ----------

def _cell(cell: Cell):
    """-> (body paragraphs, bullets) for a str / list / {body, bullets} cell."""
    if cell is None or cell == "":
        return [], []
    if isinstance(cell, str):
        return [cell], []
    if isinstance(cell, dict):
        body = cell.get("body")
        return ([body] if body else []), list(cell.get("bullets", []))
    return [], [str(c) for c in cell]


def _cell_h(cell: Cell, w: float, size: int) -> float:
    body, bullets = _cell(cell)
    if not body and not bullets:
        return 0.0
    h = 0.0
    if body:
        h += text_height_in(body, w, size, para_gap_pt=5)
    if bullets:
        h += text_height_in(bullets, w, size, para_gap_pt=5, indent_in=BULLET_INDENT)
        if body:
            h += 5 / 72
    return h


def _write_cell(slide, theme, x, y, w, h, cell: Cell, size: int, *, color=None,
                bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    body, bullets = _cell(cell)
    tb = add_textbox(slide, x, y, w, h, anchor=anchor)
    first = True
    for t in body:
        write_rich_paragraph(tb.text_frame, t, size=size, theme=theme, color=color,
                             bold=bold, align=align, first=first,
                             space_before=None if first else 5)
        first = False
    for t in bullets:
        write_rich_paragraph(tb.text_frame, t, size=size, theme=theme, color=color,
                             bold=bold, bullet=True, first=first,
                             space_before=None if first else 5)
        first = False
    return tb


def _longest_word_size(texts, w, size, min_size=10, bold=True):
    """Shrink so that the longest single word of every text fits on one line."""
    for t in texts:
        words = plain(t).split() or [""]
        size = min(size, fit_one_line(max(words, key=len), w, size, min_size, bold=bold))
    return size


# ---------- logic grid ----------

def add_logic_grid(prs, *,
                   title: str = "[Logic grid / Insert action title]",
                   stages: Sequence[str],
                   rows: Sequence[Dict],
                   conclusion_label: Optional[str] = "Implication",
                   conclusion: Optional[Cell] = None,
                   direction: str = "across",
                   arrows: str = "conclusion",
                   header_style: str = "chevron",
                   label_header: Optional[str] = None,
                   subtitle: Optional[str] = None,
                   insight: Optional[str] = None,
                   insight_label: Optional[str] = "Key insight",
                   page_number=None, section_marker=None,
                   source=None, footnote=None,
                   theme: Theme = DEFAULT_THEME):
    """stages: headers of the analysis columns, in reasoning order
               (e.g. ["External situation", "Our capability"]).
    rows: [{"label"?: str, "icon"?: str, "cells": [cell per stage],
            "conclusion"?: str, "tone"?: str}]
          cell = str | [bullets] | {"body"?, "bullets"?}; markup allowed.
    conclusion_label: header of the dark conclusion column ("Competitive
          advantage", "Impact on us", "Implication").
    conclusion: one conclusion for all rows (a single dark box spanning the
          rows) instead of a conclusion per row.
    direction: "across" (stages left -> right, one row per topic) or "down"
          (stages top -> bottom, one column per topic: PEST, five forces).
    arrows: "conclusion" (arrow into the conclusion only), "all" (also
          between stages), "none".
    header_style: "chevron" | "rule" | "bar".
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    if direction == "down":
        _grid_down(slide, theme, left, top, width, bottom, list(stages), list(rows),
                   conclusion_label, arrows, title)
    else:
        _grid_across(slide, theme, left, top, width, bottom, list(stages), list(rows),
                     conclusion_label, arrows, header_style, label_header, title,
                     shared=conclusion)
    return slide


def _grid_across(slide, theme, left, top, width, bottom, stages, rows, conclusion_label,
                 arrows, header_style, label_header, where, shared=None):
    pal = theme.palette
    n, m = max(len(stages), 1), max(len(rows), 1)
    has_label = any(r.get("label") for r in rows)
    has_icon = any(r.get("icon") for r in rows)
    has_concl = shared is not None or any(r.get("conclusion") for r in rows)
    lw = 1.55 if has_label else 0.0
    lg = 0.15 if has_label else 0.0
    cw = min(2.9, width * 0.23) if has_concl else 0.0
    ag = 0.5 if has_concl else 0.0
    cg = 0.42 if arrows == "all" else 0.15
    cell_w = (width - lw - lg - cw - ag - cg * (n - 1)) / n
    pad_x, pad_y = 0.16, 0.12
    inner = cell_w - 2 * pad_x
    c_inner = cw - 0.36
    xs = [left + lw + lg + i * (cell_w + cg) for i in range(n)]
    cx = left + width - cw

    head_h = 0.44
    spans = [(xs[i], cell_w, stages[i], "stage") for i in range(n)]
    if has_concl and conclusion_label:
        spans.append((cx, cw, conclusion_label, "result"))
    add_stage_banners(slide, theme, spans, top, head_h, style=header_style, where=where)
    if label_header and has_label:
        tb = add_textbox(slide, left, top, lw, head_h, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, label_header, size=12, bold=True, color=pal.mid_blue,
                        family=theme.typography.family, align=PP_ALIGN.CENTER, first=True)

    body_top = top + head_h + 0.16
    avail = bottom - body_top
    rg = 0.18 if m > 1 else 0.0
    icon_d = 0.48

    def label_size(s):
        return _longest_word_size([r.get("label", "") for r in rows], lw - 0.2,
                                  min(s + 2, 17), 11)

    def need(r, s):
        cells = list(r.get("cells", []))
        hs = [_cell_h(c, inner, s) + 2 * pad_y for c in cells if _cell(c) != ([], [])]
        if r.get("conclusion") and shared is None:
            hs.append(_cell_h(r["conclusion"], c_inner, s + 1) * 1.08 + 0.4)
        if has_label:
            hs.append((icon_d + 0.1 if has_icon else 0) + 0.3 +
                      text_height_in([r.get("label", "")], lw - 0.2, label_size(s), bold=True))
        return max(hs + [0.6])

    size = MIN_PT
    for s in range(MAX_PT, MIN_PT - 1, -1):
        if sum(need(r, s) for r in rows) + rg * (m - 1) <= avail:
            size = s
            break
    warn_small("logic_grid", where, size,
               "Too much text: shorten cells to 2-3 short bullets, drop the subtitle "
               "or insight bar, or split the rows over two slides.")
    needs = [need(r, size) for r in rows]
    extra = avail - sum(needs) - rg * (m - 1)
    add = min(extra / m, 0.8) if extra > 0 else 0.0
    heights = [h + add for h in needs]
    lsize = label_size(size)

    y = body_top
    for r, rh in zip(rows, heights):
        if has_label:
            add_rect(slide, left, y, lw, rh, fill=pal.light_gray)
            lab_h = text_height_in([r.get("label", "")], lw - 0.2, lsize, bold=True)
            grp = (icon_d + 0.1 if r.get("icon") else 0) + lab_h
            gy = y + (rh - grp) / 2
            if r.get("icon"):
                add_icon(slide, r["icon"], left + (lw - icon_d) / 2, gy, icon_d, theme,
                         r.get("label_tone", "navy"))
                gy += icon_d + 0.1
            tb = add_textbox(slide, left + 0.1, gy, lw - 0.2, lab_h + 0.05)
            write_rich_paragraph(tb.text_frame, r.get("label", ""), size=lsize, theme=theme,
                                 color=pal.deep_navy, bold=True, align=PP_ALIGN.CENTER,
                                 first=True)
        cells = list(r.get("cells", []))
        for i in range(n):
            cell = cells[i] if i < len(cells) else None
            add_rect(slide, xs[i], y, cell_w, rh, fill=pal.soft_gray)
            if _cell(cell) != ([], []):
                _write_cell(slide, theme, xs[i] + pad_x, y + pad_y, inner, rh - 2 * pad_y,
                            cell, size)
            if arrows == "all" and i < n - 1:
                add_wedge(slide, xs[i] + cell_w + (cg - 0.16) / 2, y + rh / 2 - 0.17, 0.16,
                          0.34, tint(pal.mid_blue, 0.45))
        if r.get("conclusion") and shared is None:
            _conclusion_box(slide, theme, cx, y, cw, rh, ag, r["conclusion"], size + 1,
                            r.get("tone"), arrows)
        elif shared is not None and arrows != "none":
            wh = min(rh * 0.6, 0.9)
            add_wedge(slide, cx - ag + 0.14, y + (rh - wh) / 2, 0.2, wh,
                      tint(pal.mid_blue, 0.55))
        y += rh + rg
    if shared is not None:
        block = sum(heights) + rg * (m - 1)
        ssize = min(size + 3, fit_size(_cell(shared)[0] + _cell(shared)[1], c_inner,
                                       block - 0.4, max_size=18, min_size=11, bold=True,
                                       para_gap_pt=5, indent_in=BULLET_INDENT))
        ch = min(block, _cell_h(shared, c_inner, ssize) * 1.08 + 0.6)
        cy = body_top + (block - ch) / 2
        add_rect(slide, cx, cy, cw, ch, fill=pal.deep_navy)
        _write_cell(slide, theme, cx + 0.18, cy, c_inner, ch, shared, ssize,
                    color=pal.white, bold=True)


def _conclusion_box(slide, theme, cx, y, cw, rh, ag, cell, size, tone, arrows):
    pal = theme.palette
    c_inner = cw - 0.36
    body, bullets = _cell(cell)
    ch = min(rh, _cell_h(cell, c_inner, size) * 1.08 + 0.5)
    cy = y + (rh - ch) / 2
    if arrows != "none":
        wh = min(rh * 0.8, max(ch, 0.6))
        add_wedge(slide, cx - ag + 0.12, y + (rh - wh) / 2, 0.24, wh, tint(pal.mid_blue, 0.35))
    add_rect(slide, cx, cy, cw, ch, fill=tone_rgb(theme, tone or "navy"))
    _write_cell(slide, theme, cx + 0.18, cy, c_inner, ch, cell, size, color=pal.white,
                bold=True, align=PP_ALIGN.CENTER if not bullets else PP_ALIGN.LEFT)


def _grid_down(slide, theme, left, top, width, bottom, stages, items, conclusion_label,
               arrows, where):
    """Columns = topics (rows' items), bands = stages top -> bottom, conclusion band last."""
    pal = theme.palette
    k, m = max(len(stages), 1), max(len(items), 1)
    has_concl = any(it.get("conclusion") for it in items)
    sw = 1.35
    sg = 0.15
    cg = 0.18
    col_w = (width - sw - sg - cg * (m - 1)) / m
    xs = [left + sw + sg + j * (col_w + cg) for j in range(m)]
    pad_x, pad_y = 0.16, 0.1
    inner = col_w - 2 * pad_x

    head_h = 0.44
    add_stage_banners(slide, theme, [(xs[j], col_w, items[j].get("label", ""), "result")
                                     for j in range(m)], top, head_h, style="bar", where=where)
    for j, it in enumerate(items):
        if it.get("icon"):
            pass  # icons would crowd the header bar; labels carry the topic
    body_top = top + head_h + 0.14
    avail = bottom - body_top
    bands = k + (1 if has_concl else 0)
    bg = 0.32  # room for the down arrow between bands

    def cell_of(it, i):
        cells = list(it.get("cells", []))
        return cells[i] if i < len(cells) else None

    def band_need(i, s):
        if i < k:
            return max([_cell_h(cell_of(it, i), inner, s) + 2 * pad_y for it in items] + [0.5])
        return max([text_height_in([it.get("conclusion", "")], inner, s, bold=True) + 0.3
                    for it in items] + [0.5])

    lab_texts = list(stages) + ([conclusion_label or ""] if has_concl else [])
    size = MIN_PT
    for s in range(MAX_PT, MIN_PT - 1, -1):
        if sum(band_need(i, s) for i in range(bands)) + bg * (bands - 1) <= avail:
            size = s
            break
    warn_small("logic_grid", where, size,
               "Too much text: shorten cells or use fewer topics / stages.")
    needs = [band_need(i, size) for i in range(bands)]
    extra = avail - sum(needs) - bg * (bands - 1)
    add = min(extra / bands, 1.2) if extra > 0 else 0.0
    heights = [h + add for h in needs]
    lsize = _longest_word_size(lab_texts, sw - 0.2, 13, 10)
    lsize = min([lsize] + [fit_size([t], sw - 0.2, h - 0.1, max_size=lsize, min_size=10,
                                    bold=True) for t, h in zip(lab_texts, heights)])

    y = body_top
    for i, bh in enumerate(heights):
        concl = i >= k
        # stage label strip
        add_line(slide, left + sw, y, left + sw, y + bh, color=pal.rule_gray, width_pt=1.0)
        tb = add_textbox(slide, left, y, sw - 0.15, bh, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, lab_texts[i], size=lsize, theme=theme,
                             color=pal.deep_navy if concl else pal.mid_blue, bold=True,
                             first=True)
        if concl:
            add_rect(slide, xs[0], y, xs[-1] + col_w - xs[0], bh,
                     fill=tint(pal.light_blue, 0.82))
        for j, it in enumerate(items):
            if concl:
                if it.get("conclusion"):
                    tb = add_textbox(slide, xs[j] + pad_x, y + pad_y, inner, bh - 2 * pad_y,
                                     anchor=MSO_ANCHOR.MIDDLE)
                    write_rich_paragraph(tb.text_frame, it["conclusion"], size=size,
                                         theme=theme, color=pal.deep_navy, bold=True,
                                         first=True)
                continue
            add_rect(slide, xs[j], y, col_w, bh, fill=pal.soft_gray)
            cell = cell_of(it, i)
            if _cell(cell) != ([], []):
                _write_cell(slide, theme, xs[j] + pad_x, y + pad_y, inner, bh - 2 * pad_y,
                            cell, size)
        if i < bands - 1 and arrows != "none":
            add_wedge(slide, left + sw / 2 - 0.22, y + bh + (bg - 0.2) / 2, 0.44, 0.2,
                      tint(pal.mid_blue, 0.35), direction="down")
        y += bh + bg


# ---------- strategic challenge ----------

def _as_driver(d):
    return d if isinstance(d, dict) else {"body": str(d)}


def add_strategic_challenge(prs, *,
                            title: str = "[Strategic challenge / Insert action title]",
                            drivers: Sequence[Union[str, Dict]],
                            question: str,
                            threat: Optional[str] = None,
                            labels: Sequence[str] = ("Drivers", "Possible results", "Key threat"),
                            question_label: Optional[str] = "Strategic challenge",
                            directions: Optional[Sequence[Union[str, Dict]]] = None,
                            directions_label: str = "Ways to address it",
                            subtitle: Optional[str] = None,
                            page_number=None, section_marker=None,
                            source=None, footnote=None,
                            theme: Theme = DEFAULT_THEME):
    """drivers: [{"title"?: str, "body": str, "result"?: str}] (2-4) — the forces
              at work; `result` is what each could do to us.
    threat: the one consequence they converge on (optional).
    question: the key question the rest of the deck answers, usually
              "How can <company> <goal>, given <constraint>?"
    directions: optional 2-4 strands of the answer (they lead into options).
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, None, None, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    drivers = [_as_driver(d) for d in drivers]
    n = max(len(drivers), 1)
    has_res = any(d.get("result") for d in drivers)
    has_threat = bool(threat)

    # --- bottom blocks first: question (+ directions)
    q_w = width - 1.2
    q_size = fit_size([question], q_w, 0.85, max_size=20, min_size=14, bold=True)
    q_h = min(1.5, max(0.95, text_height_in([question], q_w, q_size, bold=True) + 0.5))
    dirs = [d if isinstance(d, dict) else {"body": str(d)} for d in (directions or [])]
    d_h = 0.0
    if dirs:
        dw = (width - 0.2 * (len(dirs) - 1)) / len(dirs)
        d_texts = [(f"**{d['title']}** " if d.get("title") else "") + d.get("body", "")
                   for d in dirs]
        d_size = min(fit_size([t], dw - 0.3, 0.9, max_size=14, min_size=10) for t in d_texts)
        d_h = max(text_height_in([t], dw - 0.3, d_size) for t in d_texts) + 0.2 + 0.32
        d_h = min(d_h, 1.3)
    wedge_h = 0.24
    head_h = 0.4
    top_block = bottom - top - head_h - 0.1 - (wedge_h + 0.28) - q_h - (d_h + 0.15 if dirs else 0)

    # --- columns
    g1 = 0.55 if has_res else 0.0  # number badge sits in this gap
    g2 = 0.75 if has_threat else 0.0
    avail_w = width - g1 - g2
    if has_res and has_threat:
        dw_, rw_, tw_ = avail_w * 0.37, avail_w * 0.37, avail_w * 0.26
    elif has_res:
        dw_, rw_, tw_ = avail_w * 0.5, avail_w * 0.5, 0.0
    elif has_threat:
        dw_, rw_, tw_ = avail_w * 0.66, 0.0, avail_w * 0.34
    else:
        dw_, rw_, tw_ = avail_w, 0.0, 0.0
    dx = left
    rx = dx + dw_ + g1
    tx = left + width - tw_

    spans = [(dx, dw_, labels[0], "stage")]
    if has_res:
        spans.append((rx, rw_, labels[1], "stage"))
    if has_threat:
        spans.append((tx, tw_, labels[2] if len(labels) > 2 else "Key threat", "result"))
    add_stage_banners(slide, theme, spans, top, head_h, style="rule", where=title)

    def d_text(d):
        return (f"**{d['title']}:** " if d.get("title") else "") + d.get("body", "")

    texts_d = [d_text(d) for d in drivers]
    texts_r = [d.get("result", "") for d in drivers]
    rg = 0.18 if n > 2 else 0.22
    row_cap = (top_block - rg * (n - 1)) / n
    size = MIN_PT
    for s in range(15, MIN_PT - 1, -1):
        hs = [max(text_height_in([a], dw_ - 0.3, s),
                  text_height_in([b], rw_ - 0.3, s) if (has_res and b) else 0) + 0.18
              for a, b in zip(texts_d, texts_r)]
        if max(hs) <= row_cap:
            size = s
            break
    warn_small("strategic_challenge", title, size, "Shorten drivers / results.")
    row_h = max(max(text_height_in([a], dw_ - 0.3, size),
                    text_height_in([b], rw_ - 0.3, size) if (has_res and b) else 0) + 0.28
                for a, b in zip(texts_d, texts_r))
    row_h = min(max(row_h, 0.62), row_cap)
    block = n * row_h + rg * (n - 1)
    y0 = top + head_h + 0.1 + max(0.0, (top_block - block) / 2)
    mid = y0 + block / 2

    if has_threat:
        t_size = min(size + 2, fit_size([threat], tw_ - 0.4, block, max_size=17,
                                        min_size=11, bold=True))
        t_h = min(block, max(1.0, text_height_in([threat], tw_ - 0.4, t_size, bold=True)
                             + 0.5))
        add_rect(slide, tx, mid - t_h / 2, tw_, t_h, fill=pal.deep_navy)
        tb = add_textbox(slide, tx + 0.2, mid - t_h / 2, tw_ - 0.4, t_h,
                         anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, threat, size=t_size, theme=theme,
                             color=pal.white, bold=True, first=True)

    for i, d in enumerate(drivers):
        y = y0 + i * (row_h + rg)
        add_rect(slide, dx, y, dw_, row_h, fill=pal.white, line=pal.mid_blue, line_width=1.0)
        tb = add_textbox(slide, dx + 0.15, y, dw_ - 0.3, row_h, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, texts_d[i], size=size, theme=theme, first=True)
        end_x, end_y = dx + dw_, y + row_h / 2
        if has_res:
            bd = 0.38
            add_rect(slide, rx - bd - 0.06, y + (row_h - bd) / 2, bd, bd, fill=pal.deep_navy)
            tb = add_textbox(slide, rx - bd - 0.06, y + (row_h - bd) / 2, bd, bd,
                             anchor=MSO_ANCHOR.MIDDLE)
            write_paragraph(tb.text_frame, str(i + 1), size=13, bold=True, color=pal.white,
                            family=typo.family, align=PP_ALIGN.CENTER, first=True)
            if d.get("result"):
                add_rect(slide, rx, y, rw_, row_h, fill=pal.white, line=pal.mid_blue,
                         line_width=1.0)
                tb = add_textbox(slide, rx + 0.15, y, rw_ - 0.3, row_h,
                                 anchor=MSO_ANCHOR.MIDDLE)
                write_rich_paragraph(tb.text_frame, d["result"], size=size, theme=theme,
                                     first=True)
            end_x = rx + rw_
        if has_threat:
            add_arrow(slide, end_x + 0.04, end_y, tx - 0.04, mid, color=pal.mid_blue,
                      width_pt=1.25)

    # --- wedge + question
    wy = y0 + block + 0.16
    add_fade_wedge(slide, theme, left + width * 0.32, wy, width * 0.36, wedge_h)
    qy = wy + wedge_h + 0.1
    add_rect(slide, left, qy + 0.18, width, q_h - 0.18, fill=pal.white, line=pal.deep_navy,
             line_width=1.5)
    if question_label:
        lab_w = text_width_pt(question_label, 16, True) / 72 / 0.92 + 0.5
        tb = add_textbox(slide, left + (width - lab_w) / 2, qy, lab_w, 0.36, fill=pal.white,
                         anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, question_label, size=16, bold=True,
                        color=pal.deep_navy, family=typo.family, align=PP_ALIGN.CENTER,
                        first=True)
    tb = add_textbox(slide, left + 0.6, qy + 0.36, q_w, q_h - 0.36 - 0.1,
                     anchor=MSO_ANCHOR.MIDDLE)
    write_rich_paragraph(tb.text_frame, question, size=q_size, theme=theme,
                         color=pal.mid_blue, bold=True, align=PP_ALIGN.CENTER, first=True)

    if dirs:
        dy = qy + q_h + 0.15
        tb = add_textbox(slide, left, dy, width, 0.3)
        p = write_paragraph(tb.text_frame, directions_label.upper(), size=10, bold=True,
                            color=pal.mid_blue, family=typo.family, first=True)
        letter_space(p, 120)
        dw = (width - 0.2 * (len(dirs) - 1)) / len(dirs)
        for j, t in enumerate(d_texts):
            bx = left + j * (dw + 0.2)
            add_rect(slide, bx, dy + 0.34, dw, d_h - 0.36, fill=pal.soft_gray)
            add_rect(slide, bx, dy + 0.34, 0.06, d_h - 0.36, fill=pal.mid_blue)
            tb = add_textbox(slide, bx + 0.18, dy + 0.34, dw - 0.3, d_h - 0.36,
                             anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, t, size=d_size, theme=theme, first=True)
    return slide


# ---------- storyline summary ----------

def _as_block(b):
    return b if isinstance(b, dict) else {"body": str(b)}


def add_storyline_summary(prs, *,
                          title: str = "[Executive summary / Insert action title]",
                          situation: Sequence[Union[str, Dict]],
                          question: str,
                          options: Optional[Sequence[Union[str, Dict]]] = None,
                          recommended: Union[None, int, Sequence[int]] = None,
                          recommendation: Union[None, str, Sequence[Union[str, Dict]]] = None,
                          labels: Optional[Sequence[str]] = ("Situation", "Key question",
                                                             "Options", "Recommendation"),
                          option_prefix: str = "Option",
                          badge: str = "Recommended",
                          subtitle: Optional[str] = None,
                          page_number=None, section_marker=None,
                          source=None, footnote=None,
                          theme: Theme = DEFAULT_THEME):
    """situation: 1-4 blocks [{"title", "body"?, "bullets"?}] (what we found).
    question: the key question (one sentence).
    options: [{"name"?: str, "body": str, "badge"?: str}] — names default to
             "Option A", "Option B", ...
    recommended: index or list of indices of the chosen option(s); badge text
             per option via "badge" (e.g. "Top priority", "2nd priority").
    recommendation: str, [str] or [{"title", "body"}] (1-3 items).
    labels: the four band labels on the left; None hides the label strip.
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, None, None, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    sit = [_as_block(b) for b in situation]
    opts = [_as_block(o) for o in (options or [])]
    for i, o in enumerate(opts):
        o.setdefault("name", f"{option_prefix} {chr(65 + i)}")
    rec_idx = set([recommended] if isinstance(recommended, int) else (recommended or []))
    if isinstance(recommendation, str):
        recs = [{"body": recommendation}]
    else:
        recs = [_as_block(r) for r in (recommendation or [])]

    lw = 1.45 if labels else 0.0
    x0 = left + lw + (0.25 if labels else 0)
    cw = width - (x0 - left)
    gap_band = 0.26  # room for the down chevrons
    bands = ["S", "Q"] + (["O"] if opts else []) + (["R"] if recs else [])

    q_size = fit_size([question], cw - 0.8, 0.9, max_size=17, min_size=12, bold=True)
    q_h = min(1.0, max(0.62, text_height_in([question], cw - 0.8, q_size, bold=True) + 0.3))

    g = 0.2
    s_w = (cw - g * (len(sit) - 1)) / max(len(sit), 1)
    o_w = (cw - g * (len(opts) - 1)) / max(len(opts), 1) if opts else 0
    r_w = (cw - g * (len(recs) - 1)) / max(len(recs), 1) if recs else 0
    pad = 0.12
    s_head = 0.36
    o_head = 0.3

    def sit_need(s):
        return max(_cell_h({"body": b.get("body"), "bullets": b.get("bullets", [])},
                           s_w - 2 * pad, s) for b in sit) + 2 * pad + s_head + 0.06

    def opt_need(s):
        return max(text_height_in([o.get("body", "")], o_w - 2 * pad, s)
                   for o in opts) + 2 * pad + o_head + 0.06

    def rec_text(r):
        return (f"**{r['title']}:** " if r.get("title") else "") + r.get("body", "")

    def rec_need(s):
        return max(text_height_in([rec_text(r)], r_w - 2 * pad - 0.3, s + 1) for r in recs) \
            + 2 * pad + 0.08

    avail = bottom - top - q_h - gap_band * (len(bands) - 1)
    size = MIN_PT
    for s in range(15, MIN_PT - 1, -1):
        tot = sit_need(s) + (opt_need(s) if opts else 0) + (rec_need(s) if recs else 0)
        if tot <= avail:
            size = s
            break
    warn_small("storyline_summary", title, size,
               "Shorten the situation blocks or option texts; keep 1 line per option.")
    hs = {"S": sit_need(size), "Q": q_h}
    if opts:
        hs["O"] = opt_need(size)
    if recs:
        hs["R"] = rec_need(size)
    extra = avail - sum(v for k, v in hs.items() if k != "Q")
    if extra > 0:
        flex = [k for k in hs if k != "Q"]
        for k in flex:
            hs[k] += min(extra / len(flex), 0.5)
    block = sum(hs.values()) + gap_band * (len(bands) - 1)
    y = top + max(0.0, (bottom - top - block) / 2)

    label_map = dict(zip(["S", "Q", "O", "R"], list(labels or []) + [""] * 4))
    lsize = _longest_word_size([label_map[b] for b in bands], lw - 0.05, 14, 10) if labels else 0
    head_size = _longest_word_size([b.get("title", "") for b in sit], s_w - 0.3, 14, 11)

    for bi, band in enumerate(bands):
        bh = hs[band]
        if labels:
            tb = add_textbox(slide, left, y, lw, bh, anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, label_map[band], size=lsize, theme=theme,
                                 color=pal.mid_blue, bold=True, first=True)
            add_line(slide, x0 - 0.13, y, x0 - 0.13, y + bh, color=pal.rule_gray,
                     width_pt=1.25)
        if band == "S":
            for j, b in enumerate(sit):
                bx = x0 + j * (s_w + g)
                add_rect(slide, bx, y, s_w, s_head, fill=pal.white, line=pal.deep_navy,
                         line_width=1.25)
                tb = add_textbox(slide, bx + 0.1, y, s_w - 0.2, s_head, anchor=MSO_ANCHOR.MIDDLE)
                write_rich_paragraph(tb.text_frame, b.get("title", ""), size=head_size,
                                     theme=theme, color=pal.deep_navy, bold=True,
                                     align=PP_ALIGN.CENTER, first=True)
                by = y + s_head + 0.06
                add_rect(slide, bx, by, s_w, bh - s_head - 0.06, fill=pal.soft_gray)
                _write_cell(slide, theme, bx + pad, by + pad, s_w - 2 * pad,
                            bh - s_head - 0.06 - 2 * pad,
                            {"body": b.get("body"), "bullets": b.get("bullets", [])}, size)
        elif band == "Q":
            add_rect(slide, x0, y, cw, bh, fill=pal.deep_navy)
            tb = add_textbox(slide, x0 + 0.4, y, cw - 0.8, bh, anchor=MSO_ANCHOR.MIDDLE)
            write_rich_paragraph(tb.text_frame, question, size=q_size, theme=theme,
                                 color=pal.white, bold=True, align=PP_ALIGN.CENTER, first=True)
        elif band == "O":
            for j, o in enumerate(opts):
                bx = x0 + j * (o_w + g)
                rec = j in rec_idx
                # header row: option name left, badge right (chosen options only)
                bw = 0.0
                if rec:
                    btxt = str(o.get("badge") or badge).upper()
                    bw = min(o_w * 0.6, text_width_pt(btxt, 9, True) / 72 / 0.9 + 0.3)
                    add_rect(slide, bx + o_w - bw, y + (o_head - 0.24) / 2, bw, 0.24,
                             fill=pal.bright_blue)
                    tb = add_textbox(slide, bx + o_w - bw, y + (o_head - 0.24) / 2, bw, 0.24,
                                     anchor=MSO_ANCHOR.MIDDLE)
                    p = write_paragraph(tb.text_frame, btxt, size=9, bold=True,
                                        color=pal.white, family=typo.family,
                                        align=PP_ALIGN.CENTER, first=True)
                    letter_space(p, 80)
                nw = o_w - bw - 0.1
                tb = add_textbox(slide, bx + 0.02, y, nw, o_head, anchor=MSO_ANCHOR.MIDDLE)
                write_rich_paragraph(tb.text_frame, o["name"],
                                     size=min(size + 1, 15, fit_one_line(o["name"], nw, 15, 10,
                                                                         True)),
                                     theme=theme, color=pal.deep_navy if rec else pal.mid_blue,
                                     bold=True, first=True)
                by = y + o_head + 0.06
                bhh = bh - o_head - 0.06
                box = add_rect(slide, bx, by, o_w, bhh,
                               fill=tint(pal.light_blue, 0.88) if rec else pal.white,
                               line=pal.bright_blue if rec else None,
                               line_width=2.25 if rec else None)
                if not rec:
                    set_dashed(box, pal.rule_gray, 1.0)
                tb = add_textbox(slide, bx + pad, by + pad, o_w - 2 * pad, bhh - 2 * pad,
                                 anchor=MSO_ANCHOR.MIDDLE)
                write_rich_paragraph(tb.text_frame, o.get("body", ""), size=size, theme=theme,
                                     align=PP_ALIGN.CENTER, first=True)
        elif band == "R":
            for j, r in enumerate(recs):
                bx = x0 + j * (r_w + g)
                add_rect(slide, bx, y, r_w, bh, fill=tint(pal.light_blue, 0.82))
                add_rect(slide, bx, y, 0.07, bh, fill=pal.deep_navy)
                tb = add_textbox(slide, bx + pad + 0.1, y + 0.06, r_w - 2 * pad - 0.1,
                                 bh - 0.12, anchor=MSO_ANCHOR.MIDDLE)
                write_rich_paragraph(tb.text_frame, rec_text(r), size=size + 1, theme=theme,
                                     color=pal.deep_navy, first=True)
        if bi < len(bands) - 1:
            add_wedge(slide, x0 + cw / 2 - 0.25, y + bh + (gap_band - 0.13) / 2, 0.5, 0.13,
                      tint(pal.mid_blue, 0.35), direction="down")
        y += bh + gap_band
    return slide
