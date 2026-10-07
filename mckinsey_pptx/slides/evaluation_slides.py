"""Evaluation templates: compare options, score them, register risks.

- option_profiles : 2-4 options side by side with the same facets (no
                    preference unless `recommended` is set)
- decision_matrix : criteria x weights x options, raw + weighted scores,
                    best per row, weighted totals
- risk_register   : risk cards with severity, description and mitigations
"""
from __future__ import annotations

from typing import Dict, Optional, Sequence, Union

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from ..base import add_chrome, add_rect, add_textbox, blank_slide, write_paragraph
from ..components import (HEADING_H, comp_metrics, comp_pros_cons, comp_text,
                          draw_heading, letter_space, pros_cons_height)
from ..design import (add_callout_bar, add_icon, add_insight_panel, fit_one_line,
                      fit_size, text_height_in, text_width_pt, tone_rgb, warn_small,
                      write_rich_paragraph)
from ..theme import Theme, DEFAULT_THEME
from ..labels import loc
from .table_slides import draw_table

GAP = 0.25


def _frame(prs, theme, title, subtitle, insight, insight_label, **chrome):
    slide = blank_slide(prs)
    add_chrome(slide, title=title, theme=theme, **chrome)
    pal, typo, layout = theme.palette, theme.typography, theme.layout
    left = layout.margin_left_in
    width = layout.slide_width_in - layout.margin_left_in - layout.margin_right_in
    top = layout.body_top_in + 0.05
    bottom = layout.footer_top_in - 0.25
    if subtitle:
        tb = add_textbox(slide, left, top, width, 0.35)
        write_paragraph(tb.text_frame, subtitle, size=typo.section_title_size,
                        bold=True, color=pal.text_dark, family=typo.family, first=True)
        top += 0.45
    if insight:
        h = 0.7
        add_callout_bar(slide, theme, left, bottom - h, width, h, insight,
                        label=insight_label)
        bottom -= h + GAP
    return slide, left, top, width, bottom


# ---------- option profiles ----------

def add_option_profiles(prs, *,
                        title: str = "[Options overview / Insert action title]",
                        options: Sequence[Dict],
                        recommended: Optional[int] = None,
                        subtitle: Optional[str] = None,
                        insight: Optional[str] = None,
                        insight_label: Optional[str] = "Preliminary view",
                        pros_label: str = "Advantages", cons_label: str = "Disadvantages",
                        page_number=None, section_marker=None,
                        source=None, footnote=None,
                        theme: Theme = DEFAULT_THEME):
    """options: [{"name", "tagline"?, "summary"?, "icon"?,
                  "metrics"?: [{"value", "label", "tone"?}] (0-3),
                  "pros"?: [str], "cons"?: [str]}]
    Every option gets the same facets in the same place. Leave `recommended`
    unset while the audience hasn't seen the evaluation yet.
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    n = max(len(options), 1)
    cw = (width - GAP * (n - 1)) / n
    pad = 0.22
    inner = cw - 2 * pad
    h_all = bottom - top

    head_h = 0.95
    has_summary = any(o.get("summary") for o in options)
    has_metrics = any(o.get("metrics") for o in options)
    sum_size = min((fit_size([o.get("summary", "")], inner, 0.75, max_size=14, min_size=10)
                    for o in options if o.get("summary")), default=13)
    sum_h = max((text_height_in([o.get("summary", "")], inner, sum_size) for o in options),
                default=0) + 0.12 if has_summary else 0
    met_h = 0.7 if has_metrics else 0
    pc_h = h_all - 2 * pad - head_h - sum_h - met_h - (0.08 if has_metrics else 0)

    # one pros/cons size and layout for all columns, so they read as a set;
    # stacked if it fits at a readable size, else side by side inside the card
    def best(layout):
        return next((sz for sz in range(15, 9, -1)
                     if all(pros_cons_height(o.get("pros", []), o.get("cons", []), inner,
                                             sz, layout) <= pc_h for o in options)), None)
    pc_layout, pc_size = "stacked", best("stacked")
    if pc_size is None or pc_size < 12:
        col_size = best("columns")
        if col_size and (pc_size is None or col_size > pc_size):
            pc_layout, pc_size = "columns", col_size
    if pc_size is None:
        pc_size = 10
        warn_small("option_profiles", title, 0,
                   "Pros / cons do not fit: use at most 2 short items each, drop the "
                   "summary or the insight bar.")
    # shrink the cards to their content and centre them (no half-empty cards)
    natural = 2 * pad + head_h + sum_h + (met_h + 0.08 if has_metrics else 0) + max(
        (pros_cons_height(o.get("pros", []), o.get("cons", []), inner, pc_size, pc_layout)
         for o in options), default=0) + (0.35 if recommended is not None else 0.1)
    if natural < h_all * 0.85:
        top += (h_all - natural) / 2
        h_all = natural
    warn_small("option_profiles", title, pc_size,
               "Use at most 3 short pros / cons per option.")

    for i, o in enumerate(options):
        x = left + i * (cw + GAP)
        rec = recommended == i
        add_rect(slide, x, top, cw, h_all, fill=pal.light_gray if rec else pal.soft_gray)
        y = top + pad
        tb = add_textbox(slide, x + pad, y - 0.04, 1.2, 0.42, anchor=MSO_ANCHOR.TOP)
        write_paragraph(tb.text_frame, f"{i + 1:02d}", size=22, bold=True,
                        color=pal.light_blue, family=typo.family, first=True)
        if o.get("icon"):
            add_icon(slide, o["icon"], x + cw - pad - 0.44, y, 0.44, theme,
                     "blue" if rec else "navy")
        nsize = min(18, fit_one_line(o.get("name", ""), inner, 18, 12, True))
        tb = add_textbox(slide, x + pad, y + 0.42, inner, 0.32)
        write_rich_paragraph(tb.text_frame, o.get("name", ""), size=nsize, theme=theme,
                             bold=True, first=True)
        if o.get("tagline"):
            tb = add_textbox(slide, x + pad, y + 0.72, inner, 0.25)
            write_paragraph(tb.text_frame, o["tagline"],
                            size=min(11, fit_one_line(o["tagline"], inner, 11, 9)),
                            italic=True, color=pal.footer_gray, family=typo.family,
                            first=True)
        y += head_h
        if has_summary:
            comp_text(slide, theme, x + pad, y, inner, sum_h, text=o.get("summary", ""),
                      size=sum_size)
            y += sum_h
        if has_metrics:
            comp_metrics(slide, theme, x + pad, y, inner, met_h, items=o.get("metrics", []))
            y += met_h + 0.08
        comp_pros_cons(slide, theme, x + pad, y, inner, top + h_all - pad - y,
                       pros=o.get("pros", []), cons=o.get("cons", []),
                       labels=(pros_label, cons_label), size=pc_size, layout=pc_layout)
        if rec:
            tb = add_textbox(slide, x + pad, top + h_all - 0.38, inner, 0.3)
            p = write_paragraph(tb.text_frame, loc(theme, "Recommended").upper(), size=10, bold=True,
                                color=pal.bright_blue, family=typo.family,
                                align=PP_ALIGN.RIGHT, first=True)
            letter_space(p, 150)
    return slide


# ---------- decision matrix ----------

def _weight(w) -> float:
    if isinstance(w, str):
        w = float(w.strip().rstrip("%"))
    return float(w)


def _fmt(v: float, dp: int) -> str:
    return f"{v:.{dp}f}"


def add_decision_matrix(prs, *,
                        title: str = "[Decision matrix / Insert action title]",
                        options: Sequence[str],
                        criteria: Sequence[Dict],
                        recommended: Optional[int] = None,
                        show_weighted: bool = True,
                        highlight_best: bool = True,
                        decimals: int = 2,
                        scale_note: Optional[str] = None,
                        subtitle: Optional[str] = None,
                        insight: Optional[str] = None,
                        insight_bullets: Sequence[str] = (),
                        insight_title: str = "What drives the result",
                        page_number=None, section_marker=None,
                        source=None, footnote=None,
                        theme: Theme = DEFAULT_THEME):
    """criteria: [{"name": str, "weight": 25 | 0.25 | "25%", "scores": [num per option]}]
    Weighted score = score x weight; totals are the sum of weighted scores
    (computed here from the given scores and weights — say so in the report).
    """
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, None, None, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal = theme.palette
    weights = [_weight(c["weight"]) for c in criteria]
    pct = any(w > 1 for w in weights)
    fracs = [w / 100 if pct else w for w in weights]
    n_opt = len(options)
    totals = [0.0] * n_opt
    rows = []
    for c, f, w in zip(criteria, fracs, weights):
        scores = list(c["scores"])
        best = max(scores) if highlight_best else None
        cells = [f"**{c['name']}**", f"{w:g}%" if pct else f"{w:g}"]
        for k, sc in enumerate(scores):
            totals[k] += sc * f
            txt = f"{sc:g}" + (f"  ({_fmt(sc * f, decimals)})" if show_weighted else "")
            cells.append(f"{{gold|{txt}}}" if best is not None and sc == best else txt)
        rows.append(cells)
    best_total = max(totals) if totals else None
    total_cells = [loc(theme, "Weighted total"), "100%" if pct else "1.0"]
    for t in totals:
        txt = _fmt(t, decimals)
        total_cells.append(f"{{gold|{txt}}}" if highlight_best and t == best_total else f"**{txt}**")
    rows.append(total_cells)

    has_panel = bool(insight or insight_bullets)
    panel_w = 3.6 if has_panel else 0
    gap = 0.35 if has_panel else 0
    table_w = width - panel_w - gap
    note_h = 0.3 if scale_note else 0
    th = draw_table(slide, theme, left, top, table_w, bottom - top - note_h,
                    columns=[loc(theme, "Criterion"), loc(theme, "Weight")] + list(options), rows=rows,
                    col_widths=[2.6, 0.9] + [1.5] * n_opt, highlight_col=(
                        recommended + 2 if recommended is not None else None),
                    total_row=True, font_size=13 if len(rows) <= 8 else 12)
    if scale_note:
        tb = add_textbox(slide, left, top + th + 0.05, table_w, 0.28)
        write_paragraph(tb.text_frame, scale_note, size=10, italic=True,
                        color=pal.footer_gray, family=theme.typography.family, first=True)
    if has_panel:
        add_insight_panel(slide, theme, left + table_w + gap, top, panel_w, th,
                          title=insight_title, text=insight, bullets=insight_bullets)
    return slide


# ---------- risk register ----------

_SEV = {"high": ("red", "HIGH"), "medium": ("amber", "MEDIUM"), "med": ("amber", "MEDIUM"),
        "low": ("green", "LOW"), "critical": ("red", "CRITICAL")}


def add_risk_register(prs, *,
                      title: str = "[Risk register / Insert action title]",
                      risks: Sequence[Dict],
                      columns: Optional[int] = None,
                      mitigation_label: str = "Mitigation",
                      subtitle: Optional[str] = None,
                      insight: Optional[str] = None,
                      insight_label: Optional[str] = "Bottom line",
                      page_number=None, section_marker=None,
                      source=None, footnote=None,
                      theme: Theme = DEFAULT_THEME):
    """risks: [{"title", "severity": high|medium|low|critical, "description"?,
                "mitigations"?: [str], "owner"?: str}]  (2-6 risks)"""
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    n = max(len(risks), 1)
    cols = columns or (n if n <= 3 else 2 if n == 4 else 3)
    rows_n = -(-n // cols)
    cw = (width - GAP * (cols - 1)) / cols
    ch = (bottom - top - GAP * (rows_n - 1)) / rows_n
    pad = 0.22
    inner = cw - 2 * pad
    head_h = 0.5

    def paras(r):
        return ([r["description"]] if r.get("description") else []) + list(r.get("mitigations", []))
    size = min((fit_size(paras(r) or [" "], inner - 0.2, ch - 2 * pad - head_h - HEADING_H - 0.2,
                         max_size=14, min_size=10, para_gap_pt=4, indent_in=0.25)
                for r in risks), default=12)
    warn_small("risk_register", title, size, "Shorten descriptions or mitigations.")

    def natural_h(r):
        h = 2 * pad + head_h + 0.08
        if r.get("description"):
            h += text_height_in([r["description"]], inner, size) + 0.1
        if r.get("mitigations"):
            h += HEADING_H + 0.04 + text_height_in(r["mitigations"], inner - 0.25, size,
                                                   para_gap_pt=4)
        return h + 0.1
    need = max((natural_h(r) for r in risks), default=ch)
    if need < ch * 0.85:
        block = rows_n * ch + GAP * (rows_n - 1)
        ch = need
        top += (block - (rows_n * ch + GAP * (rows_n - 1))) / 2

    for i, r in enumerate(risks):
        k, j = i % cols, i // cols
        x, y = left + k * (cw + GAP), top + j * (ch + GAP)
        tone, label = _SEV.get(str(r.get("severity", "")).lower(), ("gray", str(r.get("severity", "")).upper()))
        label = loc(theme, label.capitalize()).upper() if label.capitalize() != loc(theme, label.capitalize()) else label
        add_rect(slide, x, y, cw, ch, fill=pal.soft_gray)
        cy = y + pad
        add_rect(slide, x + pad, cy + 0.06, 0.42, 0.34, fill=pal.deep_navy)
        tb = add_textbox(slide, x + pad, cy + 0.06, 0.42, 0.34, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, f"R{i + 1}", size=11, bold=True, color=pal.white,
                        family=typo.family, align=PP_ALIGN.CENTER, first=True)
        pill_w = 1.05
        add_rect(slide, x + cw - pad - pill_w, cy + 0.08, pill_w, 0.3, fill=tone_rgb(theme, tone))
        tb = add_textbox(slide, x + cw - pad - pill_w, cy + 0.08, pill_w, 0.3,
                         anchor=MSO_ANCHOR.MIDDLE)
        p = write_paragraph(tb.text_frame, label, size=10, bold=True,
                            color=pal.text_dark if tone == "amber" else pal.white,
                            family=typo.family, align=PP_ALIGN.CENTER, first=True)
        letter_space(p, 120)
        tw = inner - 0.55 - pill_w - 0.15
        tsize = min(size + 2, fit_size([r.get("title", "")], tw, head_h, max_size=16, min_size=11))
        tb = add_textbox(slide, x + pad + 0.55, cy, tw, head_h, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, r.get("title", ""), size=tsize, theme=theme,
                             bold=True, first=True)
        cy += head_h + 0.08
        if r.get("description"):
            dh = text_height_in([r["description"]], inner, size) + 0.1
            tb = add_textbox(slide, x + pad, cy, inner, dh)
            write_rich_paragraph(tb.text_frame, r["description"], size=size, theme=theme,
                                 color=pal.footer_gray, first=True)
            cy += dh
        if r.get("mitigations"):
            cy += draw_heading(slide, theme, x + pad, cy, inner,
                               loc(theme, mitigation_label)
                               + (f" · {loc(theme, 'owner')}: {r['owner']}" if r.get("owner") else ""),
                               size=10)
            tb = add_textbox(slide, x + pad, cy + 0.04, inner, y + ch - pad - cy)
            for m_i, m in enumerate(r["mitigations"]):
                write_rich_paragraph(tb.text_frame, m, size=size, theme=theme, bullet=True,
                                     first=m_i == 0, space_before=0 if m_i == 0 else 4)
    return slide


# ---------- evaluation matrix (scores with reasons) ----------

def _dots(slide, theme, x, y, value, max_n, d, color, empty):
    """`value` filled dots out of max_n; a fractional last dot drawn as a pie."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    from ..base import add_oval
    gap = d * 0.28
    for k in range(max_n):
        cx = x + k * (d + gap)
        frac = max(0.0, min(1.0, value - k))
        if frac >= 0.999:
            add_oval(slide, cx, y, d, d, fill=color)
        elif frac <= 0.001:
            add_oval(slide, cx, y, d, d, fill=None, line=empty, line_width=0.75)
        else:
            add_oval(slide, cx, y, d, d, fill=None, line=color, line_width=0.75)
            pie = slide.shapes.add_shape(MSO_SHAPE.PIE, Inches(cx), Inches(y), Inches(d),
                                         Inches(d))
            pie.adjustments[0] = -90.0
            pie.adjustments[1] = -90.0 + 360.0 * frac
            pie.fill.solid()
            pie.fill.fore_color.rgb = color
            pie.line.fill.background()
            pie.shadow.inherit = False
    return max_n * (d + gap) - gap


def add_evaluation_matrix(prs, *,
                          title: str = "[Evaluation matrix / Insert action title]",
                          options: Sequence,
                          criteria: Sequence[Dict],
                          rating: Optional[str] = "number",
                          scale_max: int = 5,
                          total: Union[str, Sequence, None] = "auto",
                          total_label: Optional[str] = None,
                          basis: Optional[str] = None,
                          recommended: Union[None, int, Sequence[int]] = None,
                          criterion_label: str = "Criterion",
                          group_label: str = "Dimension",
                          decimals: int = 2,
                          scale_note: Optional[str] = None,
                          subtitle: Optional[str] = None,
                          insight: Optional[str] = None,
                          insight_label: Optional[str] = "Bottom line",
                          page_number=None, section_marker=None,
                          source=None, footnote=None,
                          theme: Theme = DEFAULT_THEME):
    """Options as columns, criteria as rows; each cell holds a rating and/or a
    one-line reason.
    options: [str | {"name", "badge"?}] (2-4); badge e.g. "Top priority".
    criteria: [{"name", "group"? (dimension, merged when consecutive),
                "weight"? (25 | "25%" | 0.25), "scores"?: [num per option],
                "notes"?: [str per option]}]
    rating: "number" | "dots" (filled circles out of scale_max) | None (text only).
    total: "auto" (weighted sum if weights, else average), "average", "sum",
           None, or a list of given totals (one per option).
    basis: banner above the table — why these criteria ("Profit = users x
           pay rate x ARPU - cost - risk").
    """
    from ..base import add_line
    from ..design import set_dashed, tint
    from ..labels import loc
    slide, left, top, width, bottom = _frame(
        prs, theme, title, subtitle, insight, insight_label, page_number=page_number,
        section_marker=section_marker, source=source, footnote=footnote)
    pal, typo = theme.palette, theme.typography
    opts = [o if isinstance(o, dict) else {"name": str(o)} for o in options]
    n = max(len(opts), 1)
    rec = set([recommended] if isinstance(recommended, int) else (recommended or []))
    has_group = any(c.get("group") for c in criteria)
    has_w = any(c.get("weight") is not None for c in criteria)
    has_notes = any(c.get("notes") for c in criteria)
    has_scores = rating is not None and any(c.get("scores") for c in criteria)

    if basis:
        bh = 0.5
        add_rect(slide, left, top, width, bh, fill=pal.white, line=pal.mid_blue, line_width=1.5)
        bs = fit_size([basis], width - 0.4, bh - 0.06, max_size=15, min_size=11, bold=True)
        tb = add_textbox(slide, left + 0.2, top, width - 0.4, bh, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, basis, size=bs, theme=theme, bold=True,
                             align=PP_ALIGN.CENTER, first=True)
        top += bh + 0.15
    note_h = 0.32 if scale_note else 0.0
    bottom -= note_h

    gw_ = 1.45 if has_group else 0.0
    cw_ = 2.0
    ww_ = 0.8 if has_w else 0.0
    ow = (width - gw_ - cw_ - ww_) / n
    xs = [left + gw_ + cw_ + ww_ + i * ow for i in range(n)]

    # weights
    def wfrac(w):
        if w is None:
            return None
        if isinstance(w, str):
            w = float(w.strip().rstrip("%"))
        w = float(w)
        return w / 100 if w > 1 else w
    fr = [wfrac(c.get("weight")) for c in criteria]
    if has_w and abs(sum(f or 0 for f in fr) - 1) > 0.011:
        warn_small("evaluation_matrix", title, 0,
                   f"weights add up to {sum(f or 0 for f in fr) * 100:.0f}%, not 100% — the "
                   "weighted total is not on the score scale; fix the weights or say why.")

    def num(v):
        if not isinstance(v, (int, float)):
            return str(v)
        txt = f"{v:.{max(decimals, 0)}f}"
        if "." in txt:
            txt = txt.rstrip("0").rstrip(".")
        if abs(float(txt) - v) > 1e-9 and decimals == 0:  # never hide a fraction
            txt = f"{v:.1f}"
        return txt

    # totals
    tot_vals = None
    if isinstance(total, (list, tuple)):
        tot_vals = list(total)
    elif total and has_scores:
        tot_vals = []
        for k in range(n):
            sc = [(c.get("scores") or [None] * n)[k] for c in criteria]
            pairs = [(s, f) for s, f in zip(sc, fr) if s is not None]
            if not pairs:
                tot_vals.append(None)
                continue
            mode = total
            if mode == "auto":
                mode = "weighted" if has_w else "average"
            if mode == "weighted":
                tot_vals.append(sum(s * (f or 0) for s, f in pairs))
            elif mode == "sum":
                tot_vals.append(sum(s for s, _ in pairs))
            else:
                tot_vals.append(sum(s for s, _ in pairs) / len(pairs))
    if total_label is None:
        total_label = ("Weighted total" if (total == "auto" and has_w) or total == "weighted"
                       else "Average" if total in ("auto", "average") else "Total")
    total_label = loc(theme, total_label)

    # header
    head_h = 0.62
    badge_h = 0.26 if any(o.get("badge") for o in opts) or rec else 0.0
    names = [o["name"] for o in opts]
    hsz = min(fit_size([t], ow - 0.25, head_h - 0.08, max_size=14, min_size=10, bold=True)
              for t in names)
    y = top + badge_h
    if has_group:
        tb = add_textbox(slide, left, y, gw_, head_h, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, loc(theme, group_label), size=12, bold=True,
                        color=pal.mid_blue, family=typo.family, align=PP_ALIGN.CENTER, first=True)
    tb = add_textbox(slide, left + gw_, y, cw_, head_h, anchor=MSO_ANCHOR.MIDDLE)
    write_paragraph(tb.text_frame, loc(theme, criterion_label), size=12, bold=True,
                    color=pal.mid_blue, family=typo.family, align=PP_ALIGN.CENTER, first=True)
    if has_w:
        tb = add_textbox(slide, left + gw_ + cw_, y, ww_, head_h, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, loc(theme, "Weight"), size=12, bold=True,
                        color=pal.mid_blue, family=typo.family, align=PP_ALIGN.CENTER, first=True)
    for i, (x, o) in enumerate(zip(xs, opts)):
        r = i in rec
        add_rect(slide, x + 0.04, y, ow - 0.08, head_h,
                 fill=pal.deep_navy if r else tint(pal.mid_blue, 0.82))
        tb = add_textbox(slide, x + 0.12, y, ow - 0.24, head_h, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, o["name"], size=hsz, theme=theme,
                             color=pal.white if r else pal.deep_navy, bold=True,
                             align=PP_ALIGN.CENTER, first=True)
        btxt = o.get("badge") or (loc(theme, "Recommended") if r else None)
        if btxt:
            bw = min(ow - 0.2, text_width_pt(str(btxt).upper(), 9, True) / 72 / 0.9 + 0.3)
            add_rect(slide, x + (ow - bw) / 2, top, bw, badge_h - 0.02, fill=pal.bright_blue)
            tb = add_textbox(slide, x + (ow - bw) / 2, top, bw, badge_h - 0.02,
                             anchor=MSO_ANCHOR.MIDDLE)
            p = write_paragraph(tb.text_frame, str(btxt).upper(), size=9, bold=True,
                                color=pal.white, family=typo.family, align=PP_ALIGN.CENTER,
                                first=True)
            letter_space(p, 80)
    y += head_h + 0.06

    # rows
    rows_n = len(criteria) + (1 if tot_vals else 0)
    score_w = 0.0
    if has_scores:
        score_w = (scale_max * 0.2 + 0.2) if rating == "dots" else 0.55
        score_w = min(score_w, ow * 0.5)
    note_w = ow - score_w - 0.3 if has_notes else 0

    def row_need(c, s):
        hs_ = [text_height_in([c.get("name", "")], cw_ - 0.2, s, bold=True)]
        for t in (c.get("notes") or []):
            if t:
                w_ = note_w if (has_scores and rating != "dots") else ow - 0.3
                hs_.append(text_height_in([t], w_, s - 1))
        if rating == "dots" and has_notes:
            hs_[-1:] = [max(hs_[-1:] or [0]) + 0.3]
        return max(hs_) + 0.22

    avail = bottom - y - (0.5 if tot_vals else 0)
    size = 10
    for s in range(15, 9, -1):
        if sum(row_need(c, s) for c in criteria) <= avail:
            size = s
            break
    warn_small("evaluation_matrix", title, size, "Shorten the reasons or split the criteria.")
    rhs = [row_need(c, size) for c in criteria]
    extra = avail - sum(rhs)
    cap = 0.35 if has_scores else 0.8
    rhs = [h + (min(extra / len(rhs), cap) if extra > 0 and rhs else 0) for h in rhs]

    # recommended column tint
    tot_h = 0.5 if tot_vals else 0.0
    for i in rec:
        if i < n:
            add_rect(slide, xs[i] + 0.04, y, ow - 0.08, sum(rhs) + tot_h,
                     fill=tint(pal.light_blue, 0.88))

    gy = y
    groups = []
    for c, rh in zip(criteria, rhs):
        if has_group:
            gname = c.get("group") or ""
            if groups and groups[-1][0] == gname:
                groups[-1][2] += rh
            else:
                groups.append([gname, gy, rh])
        gy += rh
    for gname, g0, gh in groups:
        add_rect(slide, left, g0 + 0.03, gw_ - 0.08, gh - 0.06, fill=pal.mid_blue)
        gs = min(13, fit_size([gname], gw_ - 0.3, gh - 0.1, max_size=13, min_size=9, bold=True))
        tb = add_textbox(slide, left + 0.06, g0, gw_ - 0.2, gh, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, gname, size=gs, theme=theme, color=pal.white,
                             bold=True, align=PP_ALIGN.CENTER, first=True)

    for ci, (c, rh) in enumerate(zip(criteria, rhs)):
        add_rect(slide, left + gw_, y + 0.03, cw_ - 0.08, rh - 0.06, fill=pal.white,
                 line=pal.mid_blue, line_width=1.25)
        tb = add_textbox(slide, left + gw_ + 0.1, y, cw_ - 0.28, rh, anchor=MSO_ANCHOR.MIDDLE)
        write_rich_paragraph(tb.text_frame, c.get("name", ""), size=size, theme=theme, bold=True,
                             align=PP_ALIGN.CENTER, first=True)
        if has_w:
            tb = add_textbox(slide, left + gw_ + cw_, y, ww_, rh, anchor=MSO_ANCHOR.MIDDLE)
            w = c.get("weight")
            wt = "" if w is None else (w if isinstance(w, str) else
                                       (f"{w:g}%" if float(w) > 1 else f"{float(w):g}"))
            write_paragraph(tb.text_frame, wt, size=size, color=pal.text_dark,
                            family=typo.family, align=PP_ALIGN.CENTER, first=True)
        scores = c.get("scores") or [None] * n
        notes = c.get("notes") or [None] * n
        best = max([s for s in scores if s is not None], default=None)
        for k in range(n):
            x = xs[k]
            sc, nt = (scores[k] if k < len(scores) else None), (notes[k] if k < len(notes) else None)
            if has_scores and sc is not None:
                if rating == "dots":
                    d = 0.2
                    dy = y + (0.14 if nt else (rh - d) / 2)
                    _dots(slide, theme, x + 0.2, dy, float(sc), scale_max, d, pal.bright_blue,
                          pal.grid_gray)
                else:
                    tb = add_textbox(slide, x + 0.15, y, score_w, rh, anchor=MSO_ANCHOR.MIDDLE)
                    is_best = sc == best and sum(1 for s_ in scores if s_ == best) < n
                    write_paragraph(tb.text_frame, num(sc), size=size + 3, bold=True,
                                    color=tone_rgb(theme, "gold") if is_best else pal.deep_navy,
                                    family=typo.family, align=PP_ALIGN.CENTER, first=True)
            if nt:
                if has_scores and rating == "dots":
                    tb = add_textbox(slide, x + 0.2, y + 0.42, ow - 0.35, rh - 0.48)
                elif has_scores:
                    tb = add_textbox(slide, x + 0.2 + score_w, y, note_w, rh,
                                     anchor=MSO_ANCHOR.MIDDLE)
                else:
                    tb = add_textbox(slide, x + 0.2, y, ow - 0.35, rh, anchor=MSO_ANCHOR.MIDDLE)
                write_rich_paragraph(tb.text_frame, nt, size=size - 1, theme=theme, first=True)
        ln = add_line(slide, left + gw_ + cw_, y + rh, left + width, y + rh, color=pal.grid_gray,
                      width_pt=0.75)
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        y += rh

    if tot_vals:
        th = 0.5
        add_rect(slide, left + gw_, y + 0.03, cw_ - 0.08, th - 0.06, fill=pal.deep_navy)
        tb = add_textbox(slide, left + gw_ + 0.1, y, cw_ - 0.28, th, anchor=MSO_ANCHOR.MIDDLE)
        write_paragraph(tb.text_frame, total_label, size=size + 1, bold=True, color=pal.white,
                        family=typo.family, align=PP_ALIGN.CENTER, first=True)
        if has_w:
            tb = add_textbox(slide, left + gw_ + cw_, y, ww_, th, anchor=MSO_ANCHOR.MIDDLE)
            tw_ = sum(f or 0 for f in fr)
            write_paragraph(tb.text_frame, f"{tw_ * 100:.0f}%", size=size, bold=True,
                            color=pal.text_dark, family=typo.family, align=PP_ALIGN.CENTER,
                            first=True)
        best_t = max([t for t in tot_vals if isinstance(t, (int, float))], default=None)
        for k, t in enumerate(tot_vals[:n]):
            if t is None:
                continue
            x = xs[k]
            if rating == "dots" and isinstance(t, (int, float)):
                d = 0.22
                _dots(slide, theme, x + 0.2, y + (th - d) / 2, float(t), scale_max, d,
                      pal.deep_navy, pal.grid_gray)
                tb = add_textbox(slide, x + 0.3 + scale_max * d * 1.28, y, 0.8, th,
                                 anchor=MSO_ANCHOR.MIDDLE)
                write_paragraph(tb.text_frame, num(t), size=size, bold=True,
                                color=tone_rgb(theme, "gold") if t == best_t else pal.deep_navy,
                                family=typo.family, first=True)
            else:
                under_score = has_notes and rating == "number"
                tb = add_textbox(slide, x + 0.15, y, score_w if under_score else ow - 0.3, th,
                                 anchor=MSO_ANCHOR.MIDDLE)
                write_paragraph(tb.text_frame, num(t), size=size + 3, bold=True,
                                color=tone_rgb(theme, "gold") if t == best_t else pal.deep_navy,
                                family=typo.family, align=PP_ALIGN.CENTER, first=True)
        add_line(slide, left + gw_, y + th, left + width, y + th, color=pal.deep_navy,
                 width_pt=1.25)
        y += th
    if scale_note:
        tb = add_textbox(slide, left, y + 0.05, width, 0.28)
        write_paragraph(tb.text_frame, scale_note, size=12, italic=True, color=pal.footer_gray,
                        family=typo.family, first=True)
    return slide
