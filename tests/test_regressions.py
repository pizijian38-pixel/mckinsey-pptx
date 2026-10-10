"""Regression tests for bugs found while building real decks.

    python tests/test_regressions.py      (or: pytest tests/test_regressions.py)
"""
import contextlib
import io
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from pptx import Presentation  # noqa: E402

from mckinsey_pptx import PresentationBuilder  # noqa: E402
import deck_check  # noqa: E402


def _body_sizes(slide):
    sizes = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and len(r.text) > 15:
                        sizes.append(r.font.size.pt)
    return sizes


def test_cards_in_wide_short_region_stay_readable():
    """Three icon cards stacked in a wide composite column used to shrink the
    body to 10pt although each body is one short line (coal ESG deck, slide 5)."""
    b = PresentationBuilder()
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        b.add("composite", title="Wide, short cards",
              columns=[{"type": "chart", "chart_type": "bar", "heading": "Share, %",
                        "categories": ["Global", "EU", "US"],
                        "series": [{"name": "Share", "values": [70, 69, 59]}]},
                       {"type": "cards", "columns": 1, "cards": [
                           {"title": "Short-term reversals for energy security", "icon": "zap",
                            "body": "Gas price spikes, geopolitics and weather force grids back to "
                                    "coal — e.g. India ordered idle plants to full capacity in heatwaves"},
                           {"title": "Delayed decommissioning in mature markets", "icon": "clock",
                            "body": "EU retirements postponed amid the gas supply crisis; US delays "
                                    "driven by data-centre and electrification load growth"},
                           {"title": "Corporate greenhushing", "icon": "message",
                            "body": "Utilities say less about ESG and frame transition spending as "
                                    "operational resilience, cost efficiency and system reliability"}]}],
              widths=[0.9, 1.25], insight="Energy security frequently overrides near-term "
                                          "emission goals; commitments and execution drift apart")
    assert "WARNING" not in err.getvalue(), err.getvalue()
    sizes = _body_sizes(b.prs.slides[-1])
    assert sizes and min(sizes) >= 12, sizes


def test_card_grid_regular_layout_unchanged():
    """2x2 card grids keep the stacked layout (icon + header above the body)."""
    b = PresentationBuilder()
    b.add("card_grid", title="Four cards",
          cards=[{"title": f"Card {i}", "icon": "check", "bullets": ["One point", "Another point"]}
                 for i in range(4)])
    sizes = _body_sizes(b.prs.slides[-1])
    assert not sizes or min(sizes) >= 12


def _docx(path: Path, text: str):
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           f'<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", xml)


def test_docx_source_entities_are_decoded():
    """Quotes and ampersands in a .docx source are XML entities; the checker
    flagged “Greener & Smarter” as a quoted term not in the source."""
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "outline.docx"
        _docx(src, "Banpu executes &quot;Greener &amp; Smarter&quot; strategies")
        text = deck_check._docx_text(src)
        assert '"Greener & Smarter"' in text, text
        b = PresentationBuilder()
        b.add("card_grid", title="Banpu strategy",
              cards=[{"title": "Strategy", "body": "Banpu executes “Greener & Smarter” strategies"}],
              source="", footnote="")
        deck = Path(d) / "deck.pptx"
        b.save(str(deck))
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            deck_check.main([str(deck), "--source", str(src)])
        report = out.getvalue()
        section = report.split("b) quoted terms")[1].split("c) vocabulary")[0]
        assert "Greener" not in section, section
        assert " amp" not in report, report


def test_table_decimals_survive_numbering_strip():
    """A docx table cell holding only "24.4" sits on its own line; the regex that
    strips list numbering ("3. Market") ate "24." and left "4", so every decimal
    in a source table was reported as invented (Crest deck, slide 3)."""
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "outline.docx"
        cells = "".join(f"<w:tc><w:p><w:r><w:t>{v}</w:t></w:r></w:p></w:tc>"
                        for v in ["YNBY", "24.4", "24.6", "8.0", "1.", "2)"])
        xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
               '<w:body><w:p><w:r><w:t>3. Market share</w:t></w:r></w:p>'
               f'<w:tbl><w:tr>{cells}</w:tr></w:tbl></w:body></w:document>')
        with zipfile.ZipFile(src, "w") as z:
            z.writestr("word/document.xml", xml)
        nums = deck_check.numbers(deck_check.source_text(src))
        assert {"24.4", "24.6", "8"} <= nums or {"24.4", "24.6", "8.0"} <= nums, nums
        assert "4" not in nums and "6" not in nums, nums
        assert "3" not in nums, "heading numbering must still be stripped"


def test_catalog_examples_run():
    """Every python example in CATALOG.md must run as written: the agent copies
    them. Old examples used `...` as a placeholder (TypeError / AttributeError)."""
    import os
    import re
    cat = (ROOT / "mckinsey_pptx" / "agent" / "CATALOG.md").read_text(encoding="utf8")
    blocks = re.findall(r"```python\n(.*?)```", cat, re.S)
    assert len(blocks) > 50, len(blocks)
    bad = [blk[:80] for blk in blocks if re.search(r"(?<![.\w])\.\.\.(?![.\w])", blk)]
    assert not bad, f"'...' placeholder in catalog examples: {bad}"
    failures = []
    with tempfile.TemporaryDirectory() as d:
        cwd = os.getcwd()
        os.chdir(d)
        os.makedirs("output", exist_ok=True)
        try:
            shared = PresentationBuilder(nav=["Overview", "Situation"])
            for blk in blocks:
                if "b.add(" not in blk and "make_theme(" not in blk:
                    continue
                env = {"b": shared}
                err = io.StringIO()
                try:
                    with contextlib.redirect_stderr(err), \
                            contextlib.redirect_stdout(io.StringIO()):
                        exec(blk, env)
                except Exception as e:  # noqa: BLE001
                    failures.append(f"{type(e).__name__}: {e} :: {blk[:70]!r}")
                # examples are copied as-is: they must also build without warnings
                warns = [w for w in err.getvalue().splitlines() if "WARNING" in w]
                if warns:
                    failures.append(f"warning: {warns[0][:160]} :: {blk[:70]!r}")
            shared.save("output/all_examples.pptx")
            # every example must also pass the checker's readable-size rule [7]:
            # a template whose own example fails it fails every deck that uses it
            prs = Presentation("output/all_examples.pptx")
            for i, sl in enumerate(prs.slides, start=1):
                bs = deck_check.body_size(sl, prs.slide_height)
                tpl = deck_check.template_of(sl)
                floor = deck_check.size_floor(sl)
                if bs is not None and bs < floor and i > 1:
                    failures.append(f"[7] {tpl}: {bs:g}pt < {floor}pt")
        finally:
            os.chdir(cwd)
    assert not failures, "\n".join(failures)


def test_no_theme_shadows():
    """Connectors reference theme effect style 1, which python-pptx's default
    theme defines as an outer shadow: every rule and arrow had a drop shadow
    in PowerPoint. The saved deck must contain no shadow effect anywhere."""
    b = PresentationBuilder()
    b.add("cycle", title="Shadows", kind="vicious", steps=["A", "B", "C"], center="X")
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "s.pptx"
        b.save(str(path))
        with zipfile.ZipFile(path) as z:
            hits = [n for n in z.namelist() if n.endswith(".xml")
                    and b"outerShdw" in z.read(n)]
    assert not hits, f"shadow effects in: {hits}"


def _warnings(**kw):
    b = PresentationBuilder()
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        b.add("growth_share", **kw)
    return [w for w in err.getvalue().splitlines() if "WARNING" in w]


def test_focus_follows_title():
    """A title that names one diagram item should put the accent on it: the
    build warns when focus is missing, and when focus names no item."""
    bus = [{"name": "Batteries", "x": 12, "y": 37, "size": 4},
           {"name": "Cables", "x": 70, "y": 6, "size": 5}]
    t = "Batteries are the only question mark worth funding"
    w = _warnings(title=t, bus=bus)
    assert w and "focus='Batteries'" in w[0], w
    assert not _warnings(title=t, bus=bus, focus="Batteries")
    assert not _warnings(title=t, bus=bus, focus=[])          # deliberate opt-out
    assert not _warnings(title="Portfolio is balanced", bus=bus)
    w = _warnings(title=t, bus=bus, focus="Battery")
    assert w and "matches no item" in w[0], w
    # several items named = no single focal claim: no nudge to highlight the first
    assert not _warnings(title="Batteries and Cables both lose share", bus=bus)


def test_focus_draws_one_accent():
    """With focus set, exactly the focal bubble is filled in the accent colour."""
    from mckinsey_pptx.theme import DEFAULT_THEME
    acc = str(DEFAULT_THEME.palette.bright_blue)
    bus = [{"name": n, "x": 10 + 15 * i, "y": 5 + 8 * i, "size": 1 + i}
           for i, n in enumerate(["A1", "B2", "C3", "D4"])]
    b = PresentationBuilder()
    b.add("growth_share", title="x", bus=bus, focus="C3")
    fills = [str(sh.fill.fore_color.rgb) for sh in b.prs.slides[0].shapes
             if sh.shape_type == 1 and sh.fill.type == 1 and sh.width == sh.height]
    # one accent bubble (+ one legend dot)
    assert fills.count(acc) == 2, fills


# ---------- data fidelity of the area / width / position charts ----------

def _marks(slide, prefix):
    return {sh.name[len(prefix):]: sh for sh in slide.shapes if sh.name.startswith(prefix)}


def _rel(a, b):
    return abs(a - b) / max(abs(b), 1e-9)


def test_marimekko_area_is_share_of_total():
    """Each segment's area must be its share of the grand total (column width =
    column share, height = share within column)."""
    series = ["A", "B", "C"]
    cols = [{"name": "X", "values": [10, 30, 0]}, {"name": "Y", "values": [5, 5, 5]},
            {"name": "Z", "values": [1, 2, 42]}]
    b = PresentationBuilder()
    b.add("marimekko", title="t", series=series, columns=cols)
    m = _marks(b.prs.slides[0], "mekko:")
    grand = sum(sum(c["values"]) for c in cols)
    assert "X|C" not in m, "zero segment must be omitted, not drawn"
    areas = {k: sh.width * sh.height for k, sh in m.items()}
    tot = sum(areas.values())
    for c in cols:
        for s_, v in zip(series, c["values"]):
            if v:
                assert _rel(areas[f"{c['name']}|{s_}"] / tot, v / grand) < 0.01, (c, s_)


def test_treemap_area_is_value():
    items = [{"name": n, "value": v} for n, v in
             [("a", 50), ("b", 20), ("c", 12), ("d", 9), ("e", 6), ("f", 3)]]
    b = PresentationBuilder()
    b.add("treemap", title="t", items=items)
    m = _marks(b.prs.slides[0], "treemap:")
    tot = sum(sh.width * sh.height for sh in m.values())
    for it in items:
        sh = m[it["name"]]
        assert _rel(sh.width * sh.height / tot, it["value"] / 100) < 0.01, it


def _ribbon_ends(shp):
    """Ribbon thickness at its source and target ends, in EMU."""
    from pptx.oxml.ns import qn
    path = shp._element.find(".//" + qn("a:path"))
    ys = [int(pt.get("y")) for pt in path.iter(qn("a:pt"))]
    sy = shp.height / int(path.get("h"))
    half = len(ys) // 2
    return (ys[-1] - ys[0]) * sy, (ys[half] - ys[half - 1]) * sy


def test_sankey_width_is_value_and_conserved():
    """Node height and ribbon width share one scale; a ribbon is equally thick at
    both ends; a node that loses volume is reported."""
    flows = [{"from": "In", "to": "A", "value": 60}, {"from": "In", "to": "B", "value": 40},
             {"from": "A", "to": "Win", "value": 45}, {"from": "A", "to": "Lose", "value": 15},
             {"from": "B", "to": "Win", "value": 10}, {"from": "B", "to": "Lose", "value": 30}]
    b = PresentationBuilder()
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        b.add("sankey", title="t", flows=flows)
    assert "WARNING" not in err.getvalue(), err.getvalue()
    s = b.prs.slides[0]
    nodes = _marks(s, "sankey-node:")
    k = nodes["In"].height / 100
    assert _rel(nodes["Win"].height, 55 * k) < 0.01 and _rel(nodes["A"].height, 60 * k) < 0.01
    ribbons = _marks(s, "sankey:")
    for f in flows:
        src, dst = _ribbon_ends(ribbons[f"{f['from']} → {f['to']}"])
        assert _rel(src, f["value"] * k) < 0.01 and _rel(dst, f["value"] * k) < 0.01, f
    leak = flows[:-1]                       # B sends on 10 of its 40
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        PresentationBuilder().add("sankey", title="t", flows=leak)
    assert "'B' takes in 40 but sends on 10" in err.getvalue(), err.getvalue()


def test_area_charts_reject_negative_values():
    for kind, kw in (("treemap", {"items": [{"name": "a", "value": 5},
                                             {"name": "b", "value": -1}]}),
                     ("marimekko", {"series": ["s", "t"],
                                    "columns": [{"name": "x", "values": [1, -2]}]}),
                     ("sankey", {"flows": [{"from": "a", "to": "b", "value": -3}]})):
        try:
            PresentationBuilder(on_error="raise").add(kind, title="t", **kw)
        except ValueError:
            continue
        raise AssertionError(f"{kind} accepted a negative value")


def test_slopegraph_axes_share_one_scale():
    """Every endpoint, on either axis, sits on one linear value -> y mapping."""
    series = [{"name": n, "start": a, "end": e} for n, a, e in
              [("p", 512, 288), ("q", 376, 264), ("r", 238, 431), ("s", 164, 121)]]
    b = PresentationBuilder()
    b.add("slopegraph", title="t", series=series)
    lines = _marks(b.prs.slides[0], "slope:")
    pts = []
    for sr in series:
        ln = lines[sr["name"]]
        pts += [(sr["start"], ln.begin_y), (sr["end"], ln.end_y)]
    (v0, y0), (v1, y1) = pts[0], pts[2]
    slope = (y1 - y0) / (v1 - v0)
    assert all(abs(y - (y0 + (v - v0) * slope)) < 2000 for v, y in pts), pts  # < 0.002 in


def test_heatmap_shade_orders_with_value():
    """Darker cell = larger value, on one ramp for the whole grid."""
    vals = [[1, 9, 4], [7, 2, 5]]
    b = PresentationBuilder()
    b.add("heatmap", title="t", rows=["r1", "r2"], columns=["a", "b", "c"], values=vals)
    m = _marks(b.prs.slides[0], "heatmap:")
    lum = {}
    for i, r in enumerate(["r1", "r2"]):
        for j, c in enumerate(["a", "b", "c"]):
            rgb = m[f"{r} / {c}"].fill.fore_color.rgb
            lum[vals[i][j]] = sum(rgb)
    ordered = [lum[v] for v in sorted(lum)]
    assert ordered == sorted(ordered, reverse=True), lum


def test_bump_ranks_sit_on_fixed_rows():
    series = [{"name": "a", "ranks": [1, 2, 3]}, {"name": "b", "ranks": [2, 1, 1]},
              {"name": "c", "ranks": [3, 3, 2]}]
    b = PresentationBuilder()
    b.add("bump", title="t", snapshots=["Q1", "Q2", "Q3"], series=series)
    m = _marks(b.prs.slides[0], "bump:")
    y = {}
    for sr in series:
        for k in range(2):
            ln = m[f"{sr['name']}|{k}"]
            for rank, yy in ((sr["ranks"][k], ln.begin_y), (sr["ranks"][k + 1], ln.end_y)):
                y.setdefault(rank, set()).add(round(yy / 1000))
    assert all(len(v) == 1 for v in y.values()), y           # one row per rank
    rows = [y[r].pop() for r in sorted(y)]
    steps = {rows[i + 1] - rows[i] for i in range(len(rows) - 1)}
    assert len(steps) == 1 and steps.pop() > 0, rows          # equal pitch, rank 1 on top
    try:
        PresentationBuilder(on_error="raise").add(
            "bump", title="t", snapshots=["a", "b", "c"],
            series=[{"name": "x", "ranks": [1, 1, 1]}, {"name": "y", "ranks": [1, 2, 2]}])
    except ValueError:
        return
    raise AssertionError("duplicate rank accepted")


def test_radar_is_native_and_on_one_scale():
    b = PresentationBuilder()
    b.add("radar", title="t", criteria=["a", "b", "c"], scale_max=5,
          series=[{"name": "x", "values": [1, 5, 3]}, {"name": "y", "values": [4, 2, 2]}])
    gf = [sh for sh in b.prs.slides[0].shapes if sh.has_chart][0]
    assert [list(s.values) for s in gf.chart.plots[0].series] == [[1, 5, 3], [4, 2, 2]]
    assert gf.chart.value_axis.maximum_scale == 5
    try:
        PresentationBuilder(on_error="raise").add(
            "radar", title="t", criteria=["a", "b", "c"], scale_max=5,
            series=[{"name": "x", "values": [1, 50, 3]}])
    except ValueError:
        return
    raise AssertionError("value off the shared scale accepted")


# ---------- agent ergonomics (found on a Gemini / Antigravity run) ----------

def test_chart_highlight_is_accent_not_red():
    """Highlight is emphasis, so it uses the accent; red only when asked."""
    from mckinsey_pptx.theme import DEFAULT_THEME
    pal = DEFAULT_THEME.palette

    def colours(chart_type, highlight):
        b = PresentationBuilder()
        b.add("chart", title="t", chart_type=chart_type, categories=["a", "b", "c"],
              series=[{"name": "S1", "values": [1, 2, 3]}, {"name": "S2", "values": [3, 2, 1]}]
              if chart_type == "line" else [{"name": "S1", "values": [1, 2, 3]}],
              highlight=highlight, source="")
        gf = [sh for sh in b.prs.slides[0].shapes if sh.has_chart][0]
        ser = gf.chart.plots[0].series
        if chart_type == "line":
            return [str(s.format.line.color.rgb) for s in ser]
        return [str(pt.format.fill.fore_color.rgb) for pt in [ser[0].points[1]]]

    assert colours("line", {"series": "S2"})[1] == str(pal.bright_blue)
    assert colours("line", {"series": "S2", "tone": "red"})[1] == str(pal.status_red)
    assert colours("column", {"point": 1}) == [str(pal.bright_blue)]
    assert colours("column", {"point": 1, "tone": "red"}) == [str(pal.status_red)]


def test_catalog_lookup():
    """scripts/catalog.py prints only the asked-for entries (by name, alias or
    number), an index of every template, and near matches for typos."""
    import catalog
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        assert catalog.main(["sankey", "mekko", "77"]) == 0
    text = out.getvalue()
    assert "## 70. Sankey" in text and "## 68. Marimekko" in text and "## 77. Venn" in text
    assert "## 71." not in text and "## 1. " not in text
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        catalog.main([])
    rows = [l for l in out.getvalue().splitlines() if l[:3].strip().rstrip(".").isdigit()]
    _, templates, _ = catalog.load()
    assert len(rows) == len(templates) >= 80, len(rows)
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        assert catalog.main(["sankee"]) == 1
    assert "Did you mean: sankey" in out.getvalue()
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        catalog.main(["--icons"])
    assert "`trend_up`" in out.getvalue() and "{red|" in out.getvalue()


def _check(path, *sources):
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        code = deck_check.main([str(path)] + sum((["--source", str(s)] for s in sources), []))
    return code, out.getvalue()


def test_text_layout_share_is_advisory():
    """[11] reminds when text layouts fill most content slides, but never fails
    the check (text is right for lists) - and stays quiet for a mixed deck."""
    items = [{"title": f"Point {k}", "body": "Short supporting line"} for k in "ABC"]
    with tempfile.TemporaryDirectory() as d:
        b = PresentationBuilder()
        b.add("cover_slide", title="Deck", subtitle="s", date="")
        for i in range(3):
            b.add("card_grid", title=f"Slide {i} makes a point", cards=items, source="")
        b.add("chart", title="One chart", categories=["a", "b"],
              series=[{"name": "s", "values": [1, 2]}], source="")
        path = Path(d) / "t.pptx"
        b.save(str(path))
        code, text = _check(path)
        assert "[11]" in text and "reminder: text layouts on 3 of 4" in text, text[-400:]
        assert code == 0, "the advisory section must not flag the deck"
        b = PresentationBuilder()
        b.add("card_grid", title="One", cards=items, source="")
        b.add("chart", title="Two", categories=["a", "b"],
              series=[{"name": "s", "values": [1, 2]}], source="")
        b.save(str(path))
        assert "ok (1 of 2" in _check(path)[1]


def test_step_numbers_are_not_data():
    """Sequence numbers (01 ... 05, Step 03, L4) are layout: the checker must
    not report them as numbers missing from the source."""
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "src.md"
        src.write_text("Plan, build, test, launch, scale. Five steps.", encoding="utf-8")
        b = PresentationBuilder()
        b.add("process_flow", title="Five steps take the product to scale", focus=[],
              steps=[{"name": n, "items": [n]} for n in ["Plan", "Build", "Test", "Launch", "Scale"]],
              source="")
        b.add("cycle", title="The loop", steps=["Plan", "Build", "Test", "Launch", "Scale"],
              center="Scale", focus=[], source="")
        path = Path(d) / "t.pptx"
        b.save(str(path))
        text = _check(path, src)[1]
        block = text.split("[1]")[1].split("[2]")[0]
        assert block.strip().endswith("none"), block


def test_run_deck_builds_checks_and_renders():
    """One command: build script -> deck found -> checker -> previews."""
    import shutil
    import run_deck
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "output"
        out.mkdir()
        script = out / "build_demo.py"
        script.write_text(
            "import sys\nfrom pathlib import Path\n"
            f"sys.path.insert(0, r'{ROOT}')\n"
            "from mckinsey_pptx import PresentationBuilder\n"
            "b = PresentationBuilder()\n"
            "b.add('chart', title='Revenue doubled in two years', categories=['2023', '2025'],\n"
            "      series=[{'name': 'Revenue', 'values': [10, 20]}], source='')\n"
            "b.save(str(Path(__file__).resolve().parent / 'demo.pptx'))\n", encoding="utf-8")
        render = shutil.which("soffice") or shutil.which("libreoffice")
        log = io.StringIO()
        with contextlib.redirect_stdout(log):
            code = run_deck.main([str(script)] + ([] if render else ["--no-render"]))
        text = log.getvalue()
        assert code == 0, text[-800:]
        assert "demo.pptx" in text and "build : ok" in text and "check : clean" in text
        # the full report is kept next to the deck; a short one needs no pointer
        assert (out / "report_demo.txt").read_text(encoding="utf-8").startswith("== 1. build")
        assert "\nreport: " not in text
        run_deck.REPORT_BUDGET, saved = 100, run_deck.REPORT_BUDGET
        try:
            with contextlib.redirect_stdout(io.StringIO()) as log2:
                run_deck.main([str(script), "--no-render"])
        finally:
            run_deck.REPORT_BUDGET = saved
        assert log2.getvalue().rstrip().splitlines()[-1].startswith("report: "), log2.getvalue()[-300:]
        if render:
            prev = out / "preview_demo"
            assert (prev / "contact_sheet.png").exists() and list(prev.glob("slide-*.png")), text[-600:]
        # a failing build is reported, not hidden
        script.write_text("raise SystemExit(3)\n", encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()) as log:
            assert run_deck.main([str(script), "--no-render"]) == 1
        assert "FAILED" in log.getvalue()


_MARKUP_SKIP = {"icon", "tone", "type", "template", "chart_type", "focus", "highlight", "kind", "direction",
                "group", "kicker", "nav", "layout", "style", "header_style", "fmt", "number_format",
                "sort", "cell_label", "rating", "mode", "from", "to", "legend", "source",
                "footnote", "section_marker", "title", "insight_label", "persona", "lang",
                "stages", "categories", "labels", "options", "series_names"}


def _with_markup(v, key=None):
    if key in _MARKUP_SKIP:
        return v
    if isinstance(v, str):
        return v + " **zz**" if len(v) > 3 and not v.isdigit() else v
    if isinstance(v, (list, tuple)):
        return type(v)(_with_markup(x) for x in v)
    if isinstance(v, dict):
        return {k: _with_markup(x, k) for k, x in v.items()}
    return v


def test_markup_scope_matches_templates():
    """SKILL.md, CATALOG.md and the old docs gave three different answers to
    "which templates render **bold** / {red|..}" (41-45, 41-55, rich templates).
    design.MARKUP_FULL / MARKUP_PARTIAL are the one answer; this measures it by
    adding ` **zz**` to every text value of each catalog example and looking
    for the asterisks on the slide."""
    import re
    from mckinsey_pptx import design

    class Probe:
        def __init__(self):
            self.real = PresentationBuilder(on_error="raise")

        def add(self, name, **kw):
            return self.real.add(name, **_with_markup(kw))

    cat = (ROOT / "mckinsey_pptx" / "agent" / "CATALOG.md").read_text(encoding="utf8")
    full, partial, literal = set(), set(), set()
    for part in re.split(r"\n## ", cat):
        m = re.search(r"\(`(\w+)`", part.splitlines()[0])
        blocks = [b for b in re.findall(r"```python\n(.*?)```", part, re.S) if "b.add(" in b]
        if not m or not blocks:
            continue
        probe = Probe()
        with contextlib.redirect_stderr(io.StringIO()), contextlib.redirect_stdout(io.StringIO()):
            exec(blocks[0], {"b": probe})
        text = deck_check.slide_text(probe.real.prs.slides[-1])
        shown, raw = text.count("zz"), text.count("**zz**")
        if not shown:
            continue                    # the probe's text never reached the slide
        (literal if raw == shown else partial if raw else full).add(m.group(1))
    assert full == set(design.MARKUP_FULL), (sorted(full ^ set(design.MARKUP_FULL)))
    assert partial == set(design.MARKUP_PARTIAL), (sorted(partial ^ set(design.MARKUP_PARTIAL)))
    assert literal, "no template prints markup literally? the probe is broken"


def test_literal_markup_is_flagged():
    """issue_tree prints `**bold**` as characters; the checker reports it in [4]."""
    with tempfile.TemporaryDirectory() as d:
        b = PresentationBuilder()
        b.add("issue_tree", title="Why margin fell",
              root="Margin fell **3 pts**",
              main_drivers=[{"label": "Freight", "secondaries": [
                  {"label": "Rates", "underlying": ["Rates {red|doubled}"]}]}],
              source="", footnote="")
        b.add("card_grid", title="Rendered markup is fine",
              cards=[{"title": "Cost", "bullets": ["Up **12%** in a year"]}], source="", footnote="")
        deck = Path(d) / "d.pptx"
        b.save(str(deck))
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            deck_check.main([str(deck)])
        lines = _section(out.getvalue(), "[4] Leftover", "[6] Layout")
        assert any(l.startswith("slide 1") and "**3 pts**" in l for l in lines), lines
        assert not any(l.startswith("slide 2") for l in lines), lines


def test_selection_guidance_is_consistent():
    """The same choice was written in five places and disagreed (operations spine
    "issue_tree / card_grid" vs fishbone, "rich templates 41-55 / 41-80 / 41-45",
    line_chart vs chart line). Keep every template name real and the retired
    wordings out."""
    import re
    from mckinsey_pptx import builder
    skill = (ROOT / "SKILL.md").read_text(encoding="utf8")
    cat = (ROOT / "mckinsey_pptx" / "agent" / "CATALOG.md").read_text(encoding="utf8")
    names = set(builder._REGISTRY)
    not_templates = {"kv_table", "default_section_marker", "highlight_col", "highlight_rows",
                     "insight_label", "make_brand_theme", "make_theme", "make_zh_theme",
                     "mckinsey_pptx", "section_marker", "takeaway_header", "trend_up"}
    stray = sorted(t for t in set(re.findall(r"`([a-z][a-z0-9_]+)(?:\(|`|=)", skill))
                   if "_" in t and t not in names and t not in not_templates)
    assert not stray, f"SKILL.md names templates that do not exist: {stray}"
    for retired in ("(41–45)", "(41–55)", "templates 41–55", "root causes `issue_tree` / `card_grid`",
                    "options `data_table` → recommendation", "this template only",
                    "Multi-series time series with continuous lines"):
        assert retired not in skill and retired not in cat, f"retired wording back: {retired!r}"


def test_run_deck_summary_names_failing_sections_and_slides():
    """The summary used to say "see sections [1] [4] [5a/b] [6]-[10]" whatever
    was wrong; it now names the sections that fired and their slides, and a
    build that lists failed slides says so."""
    import run_deck
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "output"
        out.mkdir()
        (out / "src.md").write_text("The market grew.", encoding="utf-8")
        head = ("import sys\nfrom pathlib import Path\n"
                f"sys.path.insert(0, r'{ROOT}')\n"
                "from mckinsey_pptx import PresentationBuilder\n"
                "b = PresentationBuilder()\n")
        save = "b.save(str(Path(__file__).resolve().parent / 'd.pptx'))\n"
        script = out / "build_d.py"
        script.write_text(head + "b.add('card_grid', title='Market', source='', footnote='',\n"
                          "      cards=[{'title': 'Size', 'bullets': ['Grew 37% in 2024']}])\n"
                          + save, encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()) as log:
            code = run_deck.main([str(script), "--source", str(out / "src.md"), "--no-render"])
        text = log.getvalue()
        assert code == 1 and "check : items to fix — [1] slide 1" in text, text[-400:]
        script.write_text(head + "b.add('dumbbell', title='t', rowz=[])\n" + save, encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()) as log:
            assert run_deck.main([str(script), "--no-render"]) == 1
        text = log.getvalue()
        assert "BUILD FAILED" in text and "listed above" in text, text[-500:]


_SHARE = {"YNBY": [24.4, 24.6, 25.1, 25.0], "DARLIE": [19.4, 18.6, 18.0, 17.2],
          "Crest": [8.8, 8.9, 8.5, 8.0], "LSL": [7.3, 8.0, 8.6, 9.0]}


def test_players_by_period_table_gets_a_reminder():
    """Players x years copied as a table (Crest deck, slide 3): [11] reminds that a
    line chart shows the trend - advisory only, the check still passes."""
    with tempfile.TemporaryDirectory() as d:
        b = PresentationBuilder()
        b.add("data_table", title="Crest has slipped to 8.0% share",
              columns=["Brand", "2022", "2023", "2024", "2025 Est"],
              rows=[[k] + [f"{v:.1f}" for v in vals] for k, vals in _SHARE.items()], source="")
        b.add("chart", chart_type="line", title="Same data as a chart",
              categories=["2022", "2023", "2024", "2025E"],
              series=[{"name": k, "values": v} for k, v in _SHARE.items()],
              highlight={"series": "Crest"}, number_format="0.0", source="")
        path = Path(d) / "t.pptx"
        b.save(str(path))
        code, text = _check(path)
        block = text.split("[11]")[1]
        assert "slide(s) 1 show numbers by period" in block, block
        assert code == 0


def test_multi_line_chart_names_lines_at_the_end():
    """3+ lines: no legend to decode; each line is named at its last point, the
    focal line with its value."""
    b = PresentationBuilder()
    b.add("chart", chart_type="line", title="Crest slipped as LSL rose",
          categories=["2022", "2023", "2024", "2025E"],
          series=[{"name": k, "values": v, **({"tone": "navy"} if k == "LSL" else {})}
                  for k, v in _SHARE.items()],
          highlight={"series": "Crest"}, number_format="0.0", source="")
    gf = [sh for sh in b.prs.slides[0].shapes if sh.has_chart][0]
    assert not gf.chart.has_legend
    ends = {}
    for ser in gf.chart.plots[0].series:
        dl = ser.points[3].data_label
        ends[ser.name] = dl.text_frame.text if dl.has_text_frame else ""
    assert ends["Crest"] == "Crest 8.0" and ends["LSL"] == "LSL 9.0", ends
    assert ends["YNBY"] == "YNBY", ends



def test_line_end_names_do_not_collide():
    """Lines that finish close together (Saky 5.5 / Colgate 4.9 on a 0-25 axis)
    get their end names moved below the point instead of printing on top of each
    other - without landing on the next line's name (Crest 8.0 vs Saky 5.5)."""
    from pptx.enum.chart import XL_LABEL_POSITION as P
    data = dict(_SHARE, Saky=[5.1, 5.2, 5.2, 5.5], Colgate=[5.2, 5.0, 4.9, 4.9])
    b = PresentationBuilder()
    b.add("chart", chart_type="line", title="Crest slipped as LSL rose",
          categories=["2022", "2023", "2024", "2025E"],
          series=[{"name": k, "values": v} for k, v in data.items()],
          highlight={"series": "Crest"}, number_format="0.0", source="")
    gf = [sh for sh in b.prs.slides[0].shapes if sh.has_chart][0]
    pos = {ser.name: ser.points[3].data_label.position for ser in gf.chart.plots[0].series}
    assert pos["YNBY"] == P.RIGHT and pos["DARLIE"] == P.RIGHT, pos
    assert pos["LSL"] == P.RIGHT and pos["Crest"] == P.BELOW, pos      # 9.0 / 8.0
    assert pos["Saky"] == P.RIGHT and pos["Colgate"] == P.BELOW, pos    # 5.5 / 4.9


def test_read_source_keeps_tables():
    """read_source.py replaces the agent's python -c snippets: paragraphs in
    order and every table cell, so no source number is lost."""
    import read_source
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "outline.docx"
        row = lambda vals: "<w:tr>" + "".join(
            f"<w:tc><w:p><w:r><w:t>{v}</w:t></w:r></w:p></w:tc>" for v in vals) + "</w:tr>"
        xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
               '<w:body><w:p><w:r><w:t>Slide 3: Competitive landscape</w:t></w:r></w:p>'
               f'<w:tbl>{row(["Brand", "2022", "2025 Est"])}{row(["YNBY", "24.4", "25.0"])}</w:tbl>'
               '<w:p><w:r><w:t>Slide 4: Strengths</w:t></w:r></w:p></w:body></w:document>')
        with zipfile.ZipFile(src, "w") as z:
            z.writestr("word/document.xml", xml)
        out = read_source.read_any(src)
        assert out.index("Slide 3") < out.index("| YNBY | 24.4 | 25.0 |") < out.index("Slide 4"), out
        assert "| Brand | 2022 | 2025 Est |" in out
        csvp = Path(d) / "t.csv"
        csvp.write_text("a,b\n1,2.5\n", encoding="utf-8")
        assert "| 1 | 2.5 |" in read_source.read_any(csvp)


def test_run_deck_env_report():
    import run_deck
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        assert run_deck.main(["--env"]) == 0
    text = out.getvalue()
    assert "python-pptx :" in text and "LibreOffice :" in text and "PDF -> PNG  :" in text


def _section(report: str, head: str, nxt: str) -> list:
    body = report.split(head)[1].split(nxt)[0]
    return [l.strip() for l in body.splitlines()[1:] if l.strip() and l.strip() != "none"]


def _check_report(deck: Path, src: Path) -> str:
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        deck_check.main([str(deck), "--source", str(src)])
    return out.getvalue()


def test_catalog_examples_raise_no_invented_numbers():
    """Every CATALOG example, checked against its own code as the source, must
    leave section [1] empty. Axis ticks, shares, totals, weighted scores and
    sequence numbers are drawn by the template, not typed by the agent; when
    they were flagged, a correct decision_matrix or marimekko deck could never
    pass, and agents rewrote them as plain tables to get a clean run."""
    import re
    import os
    cat = (ROOT / "mckinsey_pptx" / "agent" / "CATALOG.md").read_text(encoding="utf8")
    bad = []
    cwd = os.getcwd()
    with tempfile.TemporaryDirectory() as d:
        os.chdir(d)                     # examples save to output/deck.pptx
        os.makedirs("output", exist_ok=True)
        try:
            _raise_no_invented_numbers(cat, d, bad)
        finally:
            os.chdir(cwd)
    assert not bad, "\n".join(bad)


def _raise_no_invented_numbers(cat, d, bad):
    import re
    if True:
        for part in re.split(r"\n## ", cat):
            head = part.splitlines()[0][:50]
            for blk in re.findall(r"```python\n(.*?)```", part, re.S):
                if "b.add(" not in blk:
                    continue
                b = PresentationBuilder()
                with contextlib.redirect_stderr(io.StringIO()), \
                        contextlib.redirect_stdout(io.StringIO()):
                    exec(blk, {"b": b})
                deck, src = Path(d) / "d.pptx", Path(d) / "s.md"
                b.save(str(deck))
                # numeric lists in code are written "[1,2,3]"; a source writes "1, 2, 3"
                src.write_text(re.sub(r"\[[\d.,\s-]+\]",
                                      lambda m: m.group(0).replace(",", ", "), blk),
                               encoding="utf8")
                hits = _section(_check_report(deck, src), "[1] Numbers", "[2] Source")
                if hits:
                    bad.append(f"{head}: {hits}")


def test_invented_numbers_are_still_flagged():
    """Marking computed numbers must not hide typed ones: a number the agent
    typed into a card or a table input is still reported."""
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "s.md"
        src.write_text("Option A and Option B are scored on market size.", encoding="utf8")
        b = PresentationBuilder()
        b.add("card_grid", title="Market", cards=[{"title": "Size", "bullets": ["Grew 37% in 2024"]}],
              source="", footnote="")
        b.add("decision_matrix", title="Scores", options=["Option A", "Option B"],
              criteria=[{"name": "Market size", "weight": 50, "scores": [5, 3]},
                        {"name": "Fit", "weight": 50, "scores": [4, 2]}],
              source="", footnote="")
        deck = Path(d) / "d.pptx"
        b.save(str(deck))
        lines = _section(_check_report(deck, src), "[1] Numbers", "[2] Source")
        by_slide = {l.split(":")[0]: {x.strip() for x in l.split(":", 1)[1].split(",")}
                    for l in lines}
        assert "37" in by_slide.get("slide 1", set()), lines            # typed bullet number
        s2 = by_slide.get("slide 2", set())
        assert {"5", "3", "4", "50"} <= s2, lines     # typed scores and weight ("2" = page no.)
        assert not ({"2.5", "1.5", "3.5", "4.5", "100"} & s2), lines      # computed cells


def test_build_reports_every_failed_slide_at_once():
    """One run used to expose one mistake (TypeError, then KeyError, ...): three
    bad slides cost three round trips. save() now raises one BuildErrors that
    lists every slide that failed, with the unknown argument's near match, the
    accepted names and the template to look up."""
    from mckinsey_pptx.builder import BuildErrors
    b = PresentationBuilder()
    b.add("cover_slide", title="T", subtitle="S", date="2026")
    b.add("dumbbell", title="Gap", rowz=[{"name": "a", "a": 1, "b": 2}] * 3)
    b.add("dumbell", title="Typo")
    b.add("chart", title="C", chart_type="lines", categories=["a", "b"],
          series=[{"name": "s", "values": [1, 2]}])
    b.add("treemap", title="Negative", items=[{"name": "a", "value": 5}, {"name": "b", "value": -1}])
    with tempfile.TemporaryDirectory() as d:
        try:
            b.save(str(Path(d) / "x.pptx"))
        except BuildErrors as e:
            msg = str(e)
        else:
            raise AssertionError("save() wrote a deck with failed slides")
        assert not list(Path(d).iterdir()), "no deck may be written when slides failed"
    for needle in ("4 slide(s)", "slide 2 - dumbbell", "did you mean 'rows'?", "rows*",
                   "slide 3 - ", "did you mean 'dumbbell'?", "slide 4 - chart",
                   "did you mean 'line'?", "slide 5 - treemap", "must be >= 0",
                   "catalog.py dumbbell"):
        assert needle in msg, (needle, msg)


def test_build_on_error_raise_keeps_first_failure():
    for bad in (lambda b: b.add("dumbbell", title="t", rowz=[]),
                lambda b: b.add("nope", title="t")):
        try:
            bad(PresentationBuilder(on_error="raise"))
        except (TypeError, ValueError):
            continue
        raise AssertionError("on_error='raise' must raise from add()")


def test_silent_argument_slips_become_warnings():
    """A misspelled card key or a short table row used to build a slide with
    blank cells and no message."""
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        b = PresentationBuilder()
        b.add("card_grid", title="Cards",
              cards=[{"titel": "x", "bullets": ["a"]}, {"title": "ok", "bullets": ["b"]}])
        b.add("data_table", title="Table", columns=["a", "b", "c"], rows=[["1", "2"], ["1", "2", "3"]])
    text = err.getvalue()
    assert "key 'titel' is not used — did you mean 'title'?" in text, text
    assert "data_table row 1 has 2 cells for 3 columns" in text, text
    assert text.count("WARNING") == 2, text


def test_csv_source_commas_are_delimiters():
    """"APAC,980,1010" was read as the single number 9801010 (a thousands
    separator), so every CSV row with two adjacent 3-digit values was wrong."""
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "t.csv"
        p.write_text('region,2023,2024\nAPAC,980,1010\n"US","1,200","1,350"\n', encoding="utf-8")
        nums = deck_check.numbers(deck_check.source_text(p))
        assert {"980", "1010", "1200", "1350"} <= nums, nums
        assert not any(len(n) > 4 for n in nums), nums


def _catalog_first_calls():
    """template -> keyword arguments of its first CATALOG example, chrome removed."""
    import re
    from mckinsey_pptx import builder
    cat = (ROOT / "mckinsey_pptx" / "agent" / "CATALOG.md").read_text(encoding="utf8")

    class Rec:
        def __init__(self):
            self.calls = []

        def add(self, name, **kw):
            self.calls.append((name, kw))

    found = {}
    for blk in re.findall(r"```python\n(.*?)```", cat, re.S):
        if "b.add(" not in blk:
            continue
        r = Rec()
        try:
            exec(blk, {"b": r})
        except Exception:      # noqa: BLE001 - examples that need other names
            continue
        for name, kw in r.calls:
            kw = {k: v for k, v in kw.items() if k not in (
                "title", "subtitle", "insight", "insight_label", "kicker", "group", "nav",
                "source", "footnote", "section_marker")}
            found.setdefault(builder.template_name(builder._REGISTRY[name]), kw)
    return found


def _draw_diagram(template, w, h, x=1.0, y=1.5):
    """Draw a template's CATALOG example into a w x h region; return its shapes
    (in inches) and the build warnings."""
    from pptx.util import Emu
    from mckinsey_pptx import builder
    from mckinsey_pptx.base import blank_slide
    from mckinsey_pptx.components import comp_diagram
    from mckinsey_pptx.theme import DEFAULT_THEME
    prs = PresentationBuilder().prs
    slide = blank_slide(prs)
    n0 = len(slide.shapes)
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        comp_diagram(slide, DEFAULT_THEME, x, y, w, h, template=template,
                     **_catalog_first_calls()[template])
    shapes = [tuple(Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height))
              + (sh,) for sh in list(slide.shapes)[n0:]]
    return shapes, [ln for ln in err.getvalue().splitlines() if "WARNING" in ln]


def test_diagram_region_stays_inside_its_minimum():
    """A `diagram` region promises that the template draws inside the box it is
    given; DIAGRAM_MIN is the smallest box for which that was measured."""
    from mckinsey_pptx.components import DIAGRAM_MIN
    bad = []
    for name, (w, h) in sorted(DIAGRAM_MIN.items()):
        shapes, warns = _draw_diagram(name, w, h)
        assert len(shapes) > 3, f"{name}: drew {len(shapes)} shapes"
        out = [s for s in shapes
               if s[0] < 1.0 - 0.15 or s[1] < 1.5 - 0.15
               or s[0] + s[2] > 1.0 + w + 0.15 or s[1] + s[3] > 1.5 + h + 0.15]
        small = [r.font.size.pt for *_, sh in shapes if sh.has_text_frame
                 for p in sh.text_frame.paragraphs for r in p.runs
                 if r.font.size and r.text.strip() and r.font.size.pt < 9]
        if out or warns or small:
            bad.append(f"{name} {w}x{h}: {len(out)} outside, {len(warns)} warnings, "
                       f"fonts {sorted(set(small))}")
    assert not bad, "\n".join(bad)


def test_diagram_region_below_minimum_warns():
    from mckinsey_pptx.components import DIAGRAM_MIN
    w, h = DIAGRAM_MIN["radar"]
    _, warns = _draw_diagram("radar", w - 1.0, h)
    assert any("radar needs a region of at least" in x for x in warns), warns


def test_diagram_region_rejects_unsupported_or_misused_templates():
    b = PresentationBuilder()
    kw = _catalog_first_calls()
    for region in [
        {"type": "diagram", "template": "cycle", **kw["cycle"]},
        {"type": "diagram", "template": "radr", **kw["radar"]},
        {"type": "diagram", "template": "radar", "title": "x", **kw["radar"]},
        {"type": "diagram", "template": "radar", "bogus": 1, **kw["radar"]},
    ]:
        b.add("composite", title="t", columns=[region, {"type": "bullets", "items": ["a"]}])
    msg = ""
    try:
        with tempfile.TemporaryDirectory() as d:
            b.save(str(Path(d) / "x.pptx"))
    except Exception as e:      # noqa: BLE001
        msg = str(e)
    for expect in ("cannot be drawn into a region", "did you mean 'radar'",
                   "belong on the composite", "unknown argument 'bogus'"):
        assert expect in msg, (expect, msg)


def test_composite_with_diagram_is_held_to_the_label_floor():
    """Diagram labels are 10-11pt by design; the checker must not judge a
    composite slide that carries one by the 12pt body floor."""
    kw = _catalog_first_calls()
    b = PresentationBuilder(on_error="raise")
    b.add("cover_slide", title="Probe deck")
    b.add("composite", title="Capability gap sits in speed, not price",
          columns=[{"type": "diagram", "template": "heatmap", **kw["heatmap"]},
                   [{"type": "bullets", "items": ["Payments peaks at 47%"]}]],
          source="Source: probe")
    with tempfile.TemporaryDirectory() as d:
        deck = Path(d) / "d.pptx"
        b.save(str(deck))
        prs = Presentation(str(deck))
        slide = prs.slides[1]
        assert deck_check.diagrams_of(slide) == ["heatmap"]
        assert deck_check.template_of(slide) == "composite"
        assert deck_check.group_of(slide) is None
        _, report = _check(deck)
        assert not _section(report, "[7] Small text", "[8] Footer"), report



def test_catalog_entries_are_not_cut_at_code_comments():
    """`# comment` lines inside an example were read as headings, so catalog.py
    printed data_table, card_grid, swot, chart, tier_ladder, composite and
    logic_grid without their later examples (and with an unclosed code fence)."""
    import catalog
    text, templates, named = catalog.load()
    cut = [t["names"][0] for t in templates if "\n".join(t["body"]).count("```") % 2]
    assert not cut, f"entries with an unclosed code fence: {cut}"
    comment_heads = [k for k in named if k.startswith(("Recipes:", "Option deep dive",
                                                       "Finance —", "Targets as value"))]
    assert not comment_heads, comment_heads
    composite = next(t for t in templates if t["names"][0] == "composite")
    assert "Recipes: a diagram carries" in "\n".join(composite["body"])


def test_skill_md_is_slim_and_its_commands_exist():
    """SKILL.md is read in full on every run: it stays under 30 KB, and every
    catalog.py flag it names works."""
    import re
    import catalog
    skill = (ROOT / "SKILL.md").read_text(encoding="utf8")
    assert len(skill.encode("utf8")) <= 30000, len(skill.encode("utf8"))
    flags = set(re.findall(r"(?<![\w-])(--[a-z]+(?:-[a-z]+)*)", skill))
    known = {"--guide", "--options", "--focus", "--theme", "--spines", "--icons", "--plan",
             "--source", "--env", "--no-render", "--attachment"}          # catalog.py and run_deck.py flags
    assert flags <= known, f"unknown flags in SKILL.md: {sorted(flags - known)}"
    for flag in ("--guide", "--options", "--focus", "--theme", "--spines", "--icons"):
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            assert catalog.main([flag]) == 0
        assert len(out.getvalue()) > 200, flag


_PLAN = """| # | Source section | Required elements | Message (action title) | Relationship: main · secondary | Template(s) | Why not text cards | Source data used | Inferences added |
|---|---|---|---|---|---|---|---|---|
| 1 | cover | — | Probe | — | cover_slide | — | — | — |
| 2 | gap | — | Vendor B wins | profile · what follows | composite[diagram:radar + cards + callout] | — | scores | — |
| 3 | list | — | Four initiatives | list | card_grid | 4 parallel initiatives, no order | — | — |
"""


def test_sums_and_differences_of_source_numbers_are_derived():
    """Crest run: the agent deleted "16% + 5% = 21%" and "40% + 25% = 65% of the
    budget" to clear [1] — a real insight lost. A sum or difference of two source
    numbers on the same slide is now accepted and listed; small numbers and
    numbers with no such pair are still flagged."""
    assert deck_check.arithmetic("21", {"16", "5", "12"}) == "21 = 16 + 5"
    assert deck_check.arithmetic("65", {"40", "25", "20", "15"}) == "65 = 40 + 25"
    assert deck_check.arithmetic("4", {"16", "12"}) is None           # < 10: chance
    assert deck_check.arithmetic("37", {"16", "5", "12"}) is None
    assert deck_check.arithmetic("2.6", {"8.5", "5.9"}) is None
    assert deck_check.arithmetic("12.5", {"20.5", "8"}) == "12.5 = 20.5 - 8"
    b = PresentationBuilder(on_error="raise")
    b.add("cover_slide", title="Probe")
    b.add("card_grid", title="High-end tiers reach 21% of the market",
          cards=[{"title": "50-100 RMB", "body": "16% share"},
                 {"title": ">100 RMB", "body": "5% share"},
                 {"title": "Outlook", "body": "Leader holds 37% of high-end"}],
          source="", footnote="")
    with tempfile.TemporaryDirectory() as d:
        deck, src = Path(d) / "d.pptx", Path(d) / "s.md"
        b.save(str(deck))
        src.write_text("High-end 50-100 RMB: 16%. Ultra-high >100 RMB: 5%. 50 100", encoding="utf8")
        report = _check_report(deck, src)
        hits = _section(report, "[1] Numbers", "[2] Source")
        assert hits[0] == "slide 2: 37", hits
        assert "slide 2: 21 = 16 + 5" in report, report


def test_long_titles_warn_at_build():
    """Crest run: 14 of 20 titles were 79-97 characters and shrank to 18pt or
    wrapped; the rule lived only in SKILL.md prose. The build now warns."""
    long_t = "Four structural weaknesses constrain Crest: past penalties, complexity, dilution, and channel lag"
    ok_t = "Four structural weaknesses constrain Crest's growth"
    zh_t = "佳洁士市场份额从8.8%下滑至8.0%，而云南白药与舒客等本土品牌持续扩大领先优势并抢占高端市场"
    def title_warns(t):
        err = io.StringIO()
        with contextlib.redirect_stderr(err):
            PresentationBuilder().add("card_rows", title=t, rows=[{"title": "A", "body": "b"}])
        return "cut about" in err.getvalue()
    assert title_warns(long_t)
    assert not title_warns(ok_t)
    assert title_warns(zh_t)
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        PresentationBuilder().add("cover_slide", title=long_t)
    assert "cut about" not in err.getvalue()          # full-bleed pages are exempt


def test_two_period_table_gets_a_slopegraph_hint():
    """Crest slide 6 (price tiers, 2023 vs 2024 share + a text column) stayed a
    table in both runs; the generic [11] reminder never named it."""
    b = PresentationBuilder(on_error="raise")
    b.add("cover_slide", title="Probe")
    b.add("data_table", title="High-end tiers gained share as the middle shrank",
          columns=["Price tier", "2023 share", "2024 share", "Trend"],
          rows=[["<20 RMB", "34%", "32%", "Trade-up"], ["20-50 RMB", "50%", "47%", "Squeezed"],
                ["50-100 RMB", "12%", "16%", "Efficacy"], [">100 RMB", "4%", "5%", "Gifting"]],
          source="", footnote="")
    b.add("data_table", title="Owners and dates for each workstream",
          columns=["Workstream", "Owner", "Start", "End"],
          rows=[["Pricing", "Ann", "Q1", "Q2"], ["Channels", "Bo", "Q2", "Q3"],
                ["Brand", "Cy", "Q1", "Q4"]], source="", footnote="")
    with tempfile.TemporaryDirectory() as d:
        deck = Path(d) / "d.pptx"
        b.save(str(deck))
        _, report = _check(deck)
    assert "slide(s) 2 compare two periods" in report, report[-800:]
    assert "slide(s) 2, 3" not in report


def test_words_are_not_split_mid_word():
    """Crest R3: "Explosio-n" / "Emergin-g" in a narrow table column and
    "Ecosyste-m" in a hub circle. Tables widen such columns, fitted text steps
    down until the longest word fits, and the checker reports what is left."""
    b = PresentationBuilder(on_error="raise")
    b.add("cover_slide", title="Probe")
    b.add("data_table", title="Skinification shifts demand toward active ingredients",
          columns=["Functional Segment", "Maturity", "2022-2025 Trend", "Key Tech"],
          rows=[["Anti-sensitivity", "High Growth", "Surging due to spicy and cold diets", "Novamin, Potassium nitrate"],
                ["Oral Micro-ecology", "Explosion", "Targets bad breath and flora balance", "Probiotics, Prebiotics"],
                ["Enamel Repair", "Emerging", "High-end focus on remineralization", "Hydroxyapatite (HAP)"]],
          col_widths=[1.2, 0.5, 2.5, 2.0], insight="Actives lead growth", source="", footnote="")
    b.add("hub_spoke", title="Formats extend Crest into beauty routines", center="Oral Beauty Ecosystem",
          spokes=[{"title": "Flash-White Pen"}, {"title": "Night Repair Serum"},
                  {"title": "SPA-Grade Strips"}, {"title": "Lab Chic ID"}], source="", footnote="")
    with tempfile.TemporaryDirectory() as d:
        deck = Path(d) / "d.pptx"
        b.save(str(deck))
        prs = Presentation(str(deck))
        assert [deck_check.broken_words(s) for s in prs.slides] == [[], [], []]
        # the checker still sees a word that cannot fit
        from pptx.util import Inches, Pt
        s = prs.slides[1]
        tb = s.shapes.add_textbox(Inches(1), Inches(6), Inches(0.6), Inches(0.4))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text, r.font.size = "Explosion", Pt(14)
        assert deck_check.broken_words(s) == [("Explosion", 14.0)]


def test_dropped_source_table_text_is_reported():
    """Crest R3 drew the price-tier shares as a slopegraph and lost the
    "Trend Interpretation" column; the checker only counted numbers."""
    src = ("| Price Range | 2023 | 2024 | Trend Interpretation |\n"
           "| Low-end | 34% | 32% | Shrinking: Squeezed by inflation and upgrades. |\n"
           "| High-end | 12% | 16% | Exploding: occupies minds with medical efficacy focus. |\n")
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "s.md"
        p.write_text(src, encoding="utf8")
        cells = deck_check.source_table_cells(p)
        stems = {deck_check._stem(w) for w in deck_check._WORD.findall("Shrinking squeezed inflation upgrades")}
        lost = deck_check.missing_table_text(cells, stems)
        assert [t for _, t in lost] == ["Exploding: occupies minds with medical efficacy focus."], lost


def test_read_source_saves_long_text_and_finds_the_attachment():
    """Crest R3: read_source printed 8.2 KB and the host cut the first slides;
    the agent then read its own 22 KB transcript to find the attachment path."""
    import read_source
    import os
    with tempfile.TemporaryDirectory() as d:
        cwd = os.getcwd()
        os.chdir(d)
        try:
            up = Path(d) / "brain" / "abc" / ".user_uploaded"
            up.mkdir(parents=True)
            (up / "media_1.md").write_text("Slide 1: " + "word " * 2000, encoding="utf8")
            read_source.UPLOADS, saved = Path(d) / "brain", read_source.UPLOADS
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                assert read_source.main(["--attachment"]) == 0
            read_source.UPLOADS = saved
            text = out.getvalue()
            assert len(text) < 1000, len(text)
            assert (Path(d) / "output" / "source_media_1.md").read_text(encoding="utf8").startswith("Slide 1:")
            assert text.rstrip().splitlines()[-1].startswith("--source path: ") and "media_1.md" in text
        finally:
            os.chdir(cwd)


def test_catalog_output_fits_a_tail_window():
    """Antigravity shows only the last ~8,100 characters of a command's output.
    The index lost its first 34 templates, multi-template lookups lost their
    first entries, and `--plan` lost its problem list (printed first). Every
    catalog output now stays under catalog.BUDGET with the key lines last."""
    import catalog
    _, templates, _ = catalog.load()
    every = [t["names"][0] for t in templates]
    calls = [[], ["--guide"], ["--icons", "--theme"], ["--options"], ["--spines"], ["--focus"],
             ["composite"], ["composite", "--part", "2"], every[:12], every[40:]]
    for argv in calls:
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            catalog.main(list(argv))
        text = out.getvalue()
        assert len(text) <= catalog.BUDGET, (argv[:3], len(text))
    out = io.StringIO()
    with contextlib.redirect_stdout(out):
        catalog.main(every[40:])
    text = out.getvalue()
    left = text.split("NOT shown (output limit): ", 1)[1].split(" — run:", 1)[0].split(", ")
    shown = [n for n in every[40:] if f"(`{n}`" in text]
    assert sorted(shown + left) == sorted(every[40:]), (len(shown), len(left))
    assert "- `columns: list`" in text or "columns" in text      # arguments are kept


def test_plan_table_is_parsed_and_checked():
    import plan_check
    rows = plan_check.parse_plan(_PLAN)
    assert [r["template"] for r in rows] == ["cover_slide", "composite", "card_grid"]
    assert rows[1]["parts"] == ["diagram:radar", "cards", "callout"]
    assert rows[1]["relationship"].startswith("profile")
    assert plan_check.check_plan(rows) == []
    assert plan_check.entries_needed(rows) == ["cover_slide", "composite", "radar", "card_grid"]
    bad = plan_check.parse_plan(_PLAN
                                .replace("cover_slide", "cover_slid")
                                .replace("4 parallel initiatives, no order", "—")
                                .replace("diagram:radar", "diagram:venn"))
    msgs = [f"{lv} {row}: {m}" for lv, row, m in plan_check.check_plan(bad)]
    joined = "\n".join(msgs)
    assert "unknown template 'cover_slid'" in joined and "did you mean 'cover_slide'" in joined, joined
    assert "card_grid is a text layout" in joined, joined
    assert "diagram:venn cannot be drawn in a region" in joined, joined


def test_catalog_plan_checks_without_reprinting_entries():
    import catalog
    with tempfile.TemporaryDirectory() as d:
        plan = Path(d) / "p.md"
        plan.write_text(_PLAN, encoding="utf-8")
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            code = catalog.main(["--plan", str(plan)])
        text = out.getvalue()
        assert code == 0 and "  no problems — build it" in text, text[-400:]
        assert "## " not in text and len(text) < 1500, text      # no catalog entries
        assert "composite" in text and "radar" in text and "card_grid" in text
        assert "sankey" not in text
        plan.write_text(_PLAN.replace("card_grid", "card_gird"), encoding="utf-8")
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            code = catalog.main(["--plan", str(plan)])
        assert code == 1 and "did you mean 'card_grid'" in out.getvalue()


def test_checker_reports_where_the_deck_differs_from_the_plan():
    kw = _catalog_first_calls()
    b = PresentationBuilder(on_error="raise")
    b.add("cover_slide", title="Probe")
    b.add("composite", title="Vendor B wins",
          columns=[{"type": "diagram", "template": "radar", **kw["radar"]},
                   {"type": "bullets", "items": ["a"]}])
    b.add("card_rows", title="Three steps", rows=[{"title": "One", "body": "x"}])
    with tempfile.TemporaryDirectory() as d:
        deck, plan = Path(d) / "d.pptx", Path(d) / "plan.md"
        b.save(str(deck))
        plan.write_text(_PLAN, encoding="utf-8")
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            code = deck_check.main([str(deck), "--plan", str(plan)])
        report = out.getvalue()
        assert code == 1, report
        sec = _section(report, "[12] Plan vs deck", "\n[zz]")
        assert sec == ["slide 3: plan says card_grid, deck has card_rows"], sec
        assert deck_check.FINDINGS["[12]"] == [3]
        plan.write_text(_PLAN.replace("card_grid", "card_rows"), encoding="utf-8")
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            deck_check.main([str(deck), "--plan", str(plan)])
        assert "ok (3 slides match the plan)" in out.getvalue(), out.getvalue()[-300:]


if __name__ == "__main__":
    failed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"PASS  {name}")
            except Exception as e:      # noqa: BLE001 - a crash is a failure, not a silent exit
                failed += 1
                print(f"FAIL  {name}: {type(e).__name__}: {str(e)[:300]}")
    sys.exit(1 if failed else 0)
